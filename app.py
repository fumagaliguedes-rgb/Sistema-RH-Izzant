import os
import sqlite3
import shutil
import csv
import json
import re
import urllib.request
from datetime import datetime, date, timedelta
import calendar
import tempfile
import tkinter as tk
from tkinter import ttk, messagebox, filedialog
import webbrowser

try:
    from modulos.desenvolvedor import build_modo_desenvolvedor
except Exception:
    build_modo_desenvolvedor = None

from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
from reportlab.lib import colors

APP_NAME = 'Sistema Gestão Izzant'
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DATA_DIR = os.path.join(BASE_DIR, 'dados')
PDF_DIR = os.path.join(BASE_DIR, 'PDFs')
BACKUP_DIR = os.path.join(BASE_DIR, 'backups')
RELATORIO_DIR = os.path.join(BASE_DIR, 'relatorios')
MODELOS_DIR = os.path.join(BASE_DIR, 'modelos_documentos')
DOCS_GERADOS_DIR = os.path.join(BASE_DIR, 'documentos_gerados')
ASSETS_DIR = os.path.join(BASE_DIR, 'assets')
LOGO_APP = os.path.join(ASSETS_DIR, 'izzant_logo.png')

MODELO_MESTRE_WORD = os.path.join(MODELOS_DIR, '00_Modelo_Mestre_Izzant.docx')

# Campos que o sistema preenche automaticamente a partir do cadastro, empresa e cálculos internos.
CAMPOS_AUTOMATICOS_DOCUMENTOS = {
    'EMPRESA','CNPJ','ENDERECO','CIDADE','UF','FUNCIONARIO','NOME','CPF','CTPS','ADMISSAO','DATA_ADMISSAO',
    'FUNCAO','CARGO','SETOR','CBO','SALARIO','JORNADA','HORARIO_TRABALHO','DESCANSO','SABADO','DATA','DATA_DOCUMENTO',
    'TITULO','OBSERVACAO','OBSERVAÇÃO','MOTIVO','PERIODO','RESPONSAVEL','LOCAL_DATA',
    'PERIODO_AQUISITIVO','PERIODO_AQUISITIVO_INICIO','PERIODO_AQUISITIVO_FIM','PERIODO_AQUISITIVO_COMPLETO',
    'DATA_RETORNO','DIAS_FERIAS','DIAS_GOZADOS','DIAS_ABONO','DIAS_RESTANTES','INICIO','TERMINO','DESCRICAO','VALOR_RECEBIDO','ITEM','ITEM_UNIFORME',
    'SALARIO_BASE','MEDIA_VARIAVEIS','VALOR_FERIAS','VALOR_UM_TERCO','VALOR_ABONO','ADIANTAMENTO_13','VALOR_13','TOTAL_BRUTO','INSS_ESTIMADO','IRRF_ESTIMADO','LIQUIDO_ESTIMADO'
}

# Campos específicos por modelo de documento RH.
# A chave interna vira placeholder no Word, por exemplo: {{DATA_INICIO}}, {{MOTIVO}}, {{VALOR}}.
DOC_CAMPOS_MODELO = {
    'Contrato': [
        ('TIPO_CONTRATO','Tipo de contrato'), ('DATA_INICIO','Data de início'), ('DATA_FIM','Data final/experiência'),
        ('SALARIO_CONTRATUAL','Salário contratual'), ('FORMA_PAGAMENTO','Forma de pagamento'), ('LOCAL_TRABALHO','Local de trabalho'),
        ('HORARIO_CONTRATADO','Horário contratado'), ('CLAUSULAS_ESPECIAIS','Cláusulas especiais')
    ],
    'Advertência': [
        ('DATA_OCORRENCIA','Data da ocorrência'), ('MOTIVO','Motivo'), ('DESCRICAO_FATO','Descrição dos fatos'),
        ('TESTEMUNHA_1','Testemunha 1'), ('TESTEMUNHA_2','Testemunha 2')
    ],
    'Suspensão': [
        ('DATA_OCORRENCIA','Data da ocorrência'), ('MOTIVO','Motivo'), ('PERIODO','Período da suspensão'), ('DIAS_SUSPENSAO','Dias de suspensão'),
        ('DATA_INICIO','Início da suspensão'), ('DATA_FIM','Fim da suspensão'), ('DESCRICAO_FATO','Descrição dos fatos')
    ],
    'Declaração': [
        ('FINALIDADE','Finalidade'), ('DESTINATARIO','Destinatário'), ('TEXTO_DECLARACAO','Texto complementar')
    ],
    'Ficha de Registro': [
        ('RG','RG'), ('PIS','PIS'), ('DATA_NASCIMENTO','Data de nascimento'), ('ESTADO_CIVIL','Estado civil'),
        ('NATURALIDADE','Naturalidade'), ('NOME_MAE','Nome da mãe'), ('NOME_PAI','Nome do pai'), ('ENDERECO_FUNCIONARIO','Endereço do funcionário')
    ],
    'Termo': [
        ('ASSUNTO_TERMO','Assunto do termo'), ('DESCRICAO_TERMO','Descrição do termo'), ('VIGENCIA','Vigência')
    ],
    'Recibo': [
        ('DESCRICAO_RECIBO','Descrição do recibo'), ('VALOR','Valor'), ('REFERENCIA','Referência'), ('FORMA_PAGAMENTO','Forma de pagamento')
    ],
    'EPI': [
        ('EPI_ITEM','EPI entregue'), ('EPI_CA','CA'), ('EPI_QUANTIDADE','Quantidade'), ('EPI_TAMANHO','Tamanho'),
        ('DATA_ENTREGA','Data de entrega'), ('DATA_DEVOLUCAO','Data de devolução/troca'), ('MOTIVO_TROCA','Motivo da troca')
    ],
    'Uniforme': [
        ('UNIFORME_ITEM','Uniforme entregue'), ('UNIFORME_QUANTIDADE','Quantidade'), ('UNIFORME_TAMANHO','Tamanho'),
        ('DATA_ENTREGA','Data de entrega'), ('DATA_DEVOLUCAO','Data de devolução/troca'), ('OBS_UNIFORME','Observações do uniforme')
    ],
    'Banco de Horas': [
        ('PERIODO','Período'), ('SALDO_ANTERIOR','Saldo anterior'), ('HORAS_CREDITADAS','Horas creditadas'),
        ('HORAS_DEBITADAS','Horas debitadas'), ('SALDO_ATUAL','Saldo atual'), ('OBS_BANCO_HORAS','Observações do banco de horas')
    ]
}

def normalizar_placeholder(chave):
    return ''.join(c if c.isalnum() else '_' for c in str(chave).upper()).strip('_')

def _add_years_safe(d, years):
    try:
        return d.replace(year=d.year + years)
    except ValueError:
        return d.replace(year=d.year + years, day=28)

def calcular_periodo_aquisitivo(admissao_iso, referencia=None):
    """Calcula o período aquisitivo de férias a partir da admissão.
    Retorna início, fim e texto completo em DD/MM/AAAA.
    A referência pode ser a data de início das férias ou a data atual.
    """
    if not admissao_iso:
        return '', '', ''
    try:
        adm = datetime.strptime(str(admissao_iso)[:10], '%Y-%m-%d').date()
    except Exception:
        try:
            adm = parse_data_br(str(admissao_iso))
        except Exception:
            return '', '', ''
    ref = referencia or date.today()
    if isinstance(ref, str):
        try:
            ref = parse_data_br(ref)
        except Exception:
            try:
                ref = datetime.strptime(ref[:10], '%Y-%m-%d').date()
            except Exception:
                ref = date.today()
    anos = max(0, ref.year - adm.year)
    inicio = _add_years_safe(adm, anos)
    # Se a referência ainda não chegou no aniversário de admissão do ano, volta um período.
    if inicio > ref and anos > 0:
        anos -= 1
        inicio = _add_years_safe(adm, anos)
    fim = _add_years_safe(inicio, 1) - timedelta(days=1)
    return inicio.strftime('%d/%m/%Y'), fim.strftime('%d/%m/%Y'), f"{inicio.strftime('%d/%m/%Y')} a {fim.strftime('%d/%m/%Y')}"

def calcular_periodos_aquisitivos(admissao_iso, referencia=None, quantidade=8):
    """Retorna períodos aquisitivos calculados a partir da admissão.
    Cada item possui início/fim em ISO e BR, concessivo e status básico.
    Esta função centraliza a inteligência de férias para ser usada pelo módulo Férias
    e pelos documentos, evitando cálculos duplicados.
    """
    if not admissao_iso:
        return []
    try:
        adm = datetime.strptime(str(admissao_iso)[:10], '%Y-%m-%d').date()
    except Exception:
        try:
            adm = parse_data_br(str(admissao_iso))
        except Exception:
            return []
    ref = referencia or date.today()
    if isinstance(ref, str):
        try:
            ref = parse_data_br(ref)
        except Exception:
            try:
                ref = datetime.strptime(ref[:10], '%Y-%m-%d').date()
            except Exception:
                ref = date.today()
    anos = max(0, ref.year - adm.year)
    inicio_base = _add_years_safe(adm, anos)
    if inicio_base > ref and anos > 0:
        anos -= 1
    primeiro = max(0, anos - max(0, quantidade - 3))
    ultimo = anos + 2
    itens=[]
    for n in range(primeiro, ultimo + 1):
        ini = _add_years_safe(adm, n)
        fim = _add_years_safe(ini, 1) - timedelta(days=1)
        conc = _add_years_safe(fim, 1)
        if fim <= ref:
            status = 'Disponível'
        elif ini <= ref <= fim:
            status = 'Em aquisição'
        else:
            status = 'Futuro'
        itens.append({
            'ordem': n+1,
            'inicio': ini, 'fim': fim, 'concessivo_fim': conc,
            'inicio_iso': ini.isoformat(), 'fim_iso': fim.isoformat(), 'concessivo_iso': conc.isoformat(),
            'inicio_br': ini.strftime('%d/%m/%Y'), 'fim_br': fim.strftime('%d/%m/%Y'), 'concessivo_br': conc.strftime('%d/%m/%Y'),
            'periodo_br': f"{ini.strftime('%d/%m/%Y')} a {fim.strftime('%d/%m/%Y')}",
            'status': status
        })
    return itens

DB_PATH = os.path.join(DATA_DIR, 'ponto_saude.db')

MESES = ['Janeiro','Fevereiro','Março','Abril','Maio','Junho','Julho','Agosto','Setembro','Outubro','Novembro','Dezembro']
DIAS = ['Seg','Ter','Qua','Qui','Sex','Sáb','Dom']

DEFAULT_EMPRESA = {
    'nome': 'IZZANT SERVIÇOS LTDA',
    'cnpj': '44.177.413/0001-11',
    'endereco': 'R: ALMIRANTE TAMANDARÉ - ANDAR 1 - BOX 25',
    'numero': '114',
    'bairro': 'CENTRO',
    'cidade': 'ITAJAÍ',
    'uf': 'SC'
}

DEFAULT_FUNCIONARIOS = [
('ALEXANDRE NUNES BARBOSA','2023-12-05','MOTORISTA',2195.65),
('ARTHUR LOGETO SOLTMANN','2024-10-01','MOTORISTA',2195.65),
('CLAUDIO OSMAR ALEXANDRE','2024-07-11','MOTORISTA',2195.65),
('DARDO DA COSTA VALE','2023-03-13','MOTORISTA',2195.65),
('EDGARD CORDEIRO','2023-01-16','MOTORISTA',2195.65),
('ELTON SANTANNA','2023-01-16','MOTORISTA',2195.65),
('ELVIS PRINCES MATIAS CARPIS','2023-04-12','MOTORISTA',2195.65),
('JOEL JOSÉ NEGREIROS','2023-05-23','MOTORISTA',2195.65),
('JOSE ORIDES ALVES DE OLIVEIRA','2023-01-25','MOTORISTA',2195.65),
('LUIZ CARLOS GONÇALVES','2024-02-01','MOTORISTA',2195.65),
('MCGIVER ALEXANDRE LAMIM','2022-12-09','MOTORISTA',2195.65),
('SILVIO INACIO','2024-10-09','MOTORISTA',2195.65),
('DAIANE APARECIDA RODRIGUES','2025-01-03','MOTORISTA',2195.65),
('SALÉSIO STEFFENS','2025-02-05','MOTORISTA',2195.65),
('JOSE RENATO SOARES','2025-04-14','MOTORISTA',2195.65),
('ANA PAULA SEVERINO TALHEIMER','2025-09-11','MOTORISTA',2195.65),
]


def ensure_dirs():
    for d in (DATA_DIR, PDF_DIR, BACKUP_DIR, RELATORIO_DIR, MODELOS_DIR, DOCS_GERADOS_DIR):
        os.makedirs(d, exist_ok=True)


def con():
    ensure_dirs()
    return sqlite3.connect(DB_PATH)


def init_db():
    ensure_dirs()
    new = not os.path.exists(DB_PATH)
    with con() as db:
        db.execute('''CREATE TABLE IF NOT EXISTS empresa (
            id INTEGER PRIMARY KEY CHECK(id=1), nome TEXT, cnpj TEXT, endereco TEXT, numero TEXT,
            bairro TEXT, cidade TEXT, uf TEXT)''')
        db.execute('''CREATE TABLE IF NOT EXISTS funcionarios (
            id INTEGER PRIMARY KEY AUTOINCREMENT, nome TEXT NOT NULL, cpf TEXT, ctps TEXT,
            admissao TEXT, funcao TEXT, setor TEXT DEFAULT 'GERAL', cbo TEXT, salario REAL, jornada TEXT, horario_trabalho TEXT, descanso TEXT,
            sabado TEXT, ativo INTEGER DEFAULT 1)''')
        db.execute('''CREATE TABLE IF NOT EXISTS setores (
            id INTEGER PRIMARY KEY AUTOINCREMENT, nome TEXT UNIQUE NOT NULL, descricao TEXT, ativo INTEGER DEFAULT 1)''')
        db.execute('''CREATE TABLE IF NOT EXISTS jornadas (
            id INTEGER PRIMARY KEY AUTOINCREMENT, codigo TEXT UNIQUE, nome TEXT UNIQUE NOT NULL,
            entrada1 TEXT, saida1 TEXT, entrada2 TEXT, saida2 TEXT,
            horas_dia TEXT, horas_semana TEXT, intervalo TEXT, dias_semana TEXT, tipo TEXT,
            sabado_tratamento TEXT DEFAULT 'COMPENSADO', descanso_semanal TEXT DEFAULT 'DOMINGO',
            observacao TEXT, ativo INTEGER DEFAULT 1)''')
        db.execute('''CREATE TABLE IF NOT EXISTS config (
            chave TEXT PRIMARY KEY, valor TEXT)''')
        db.execute('''CREATE TABLE IF NOT EXISTS usuarios (
            id INTEGER PRIMARY KEY AUTOINCREMENT, usuario TEXT UNIQUE, senha TEXT, perfil TEXT, ativo INTEGER DEFAULT 1)''')
        db.execute('''CREATE TABLE IF NOT EXISTS logs (
            id INTEGER PRIMARY KEY AUTOINCREMENT, data_hora TEXT, usuario TEXT, acao TEXT, detalhes TEXT)''')
        db.execute('''CREATE TABLE IF NOT EXISTS pdf_historico (
            id INTEGER PRIMARY KEY AUTOINCREMENT, data_hora TEXT, usuario TEXT, tipo TEXT, setor TEXT, mes INTEGER, ano INTEGER, arquivo TEXT, quantidade INTEGER)''')
        db.execute('''CREATE TABLE IF NOT EXISTS ocorrencias (
            id INTEGER PRIMARY KEY AUTOINCREMENT, funcionario_id INTEGER NOT NULL, tipo TEXT NOT NULL,
            data_inicio TEXT NOT NULL, data_fim TEXT NOT NULL, observacao TEXT, abona INTEGER DEFAULT 1, ativo INTEGER DEFAULT 1,
            FOREIGN KEY(funcionario_id) REFERENCES funcionarios(id))''')
        db.execute('''CREATE TABLE IF NOT EXISTS feriados (
            id INTEGER PRIMARY KEY AUTOINCREMENT, nome TEXT NOT NULL, data TEXT NOT NULL,
            tipo TEXT DEFAULT 'Municipal', empresa TEXT DEFAULT 'TODAS', setor TEXT DEFAULT 'TODOS', observacao TEXT, ativo INTEGER DEFAULT 1)''')
        db.execute('''CREATE TABLE IF NOT EXISTS escalas (
            id INTEGER PRIMARY KEY AUTOINCREMENT, codigo TEXT UNIQUE, nome TEXT UNIQUE NOT NULL,
            tipo TEXT DEFAULT 'Personalizada', descricao TEXT, dias TEXT, ativo INTEGER DEFAULT 1)''')
        db.execute('''CREATE TABLE IF NOT EXISTS ferias_controle (
            id INTEGER PRIMARY KEY AUTOINCREMENT, funcionario_id INTEGER NOT NULL,
            aquisitivo_inicio TEXT, aquisitivo_fim TEXT, concessivo_fim TEXT,
            inicio TEXT NOT NULL, fim TEXT NOT NULL, retorno TEXT, dias INTEGER DEFAULT 0,
            observacao TEXT, status TEXT DEFAULT 'Programada', ativo INTEGER DEFAULT 1,
            FOREIGN KEY(funcionario_id) REFERENCES funcionarios(id))''')
        db.execute('''CREATE TABLE IF NOT EXISTS banco_horas (
            id INTEGER PRIMARY KEY AUTOINCREMENT, funcionario_id INTEGER NOT NULL,
            data TEXT NOT NULL, tipo TEXT NOT NULL, quantidade TEXT NOT NULL,
            motivo TEXT, usuario TEXT, criado_em TEXT, ativo INTEGER DEFAULT 1,
            FOREIGN KEY(funcionario_id) REFERENCES funcionarios(id))''')

        # Migração V1.3.7: campos financeiros e de saldo no módulo Férias.
        cols_ferias = [r[1] for r in db.execute('PRAGMA table_info(ferias_controle)').fetchall()]
        for _col, _sql in {
            'dias_abono': 'ALTER TABLE ferias_controle ADD COLUMN dias_abono INTEGER DEFAULT 0',
            'dias_restantes': 'ALTER TABLE ferias_controle ADD COLUMN dias_restantes INTEGER DEFAULT 0',
            'salario_base': 'ALTER TABLE ferias_controle ADD COLUMN salario_base REAL DEFAULT 0',
            'media_variaveis': 'ALTER TABLE ferias_controle ADD COLUMN media_variaveis REAL DEFAULT 0',
            'valor_ferias': 'ALTER TABLE ferias_controle ADD COLUMN valor_ferias REAL DEFAULT 0',
            'valor_um_terco': 'ALTER TABLE ferias_controle ADD COLUMN valor_um_terco REAL DEFAULT 0',
            'valor_abono': 'ALTER TABLE ferias_controle ADD COLUMN valor_abono REAL DEFAULT 0',
            'valor_13': 'ALTER TABLE ferias_controle ADD COLUMN valor_13 REAL DEFAULT 0',
            'total_bruto': 'ALTER TABLE ferias_controle ADD COLUMN total_bruto REAL DEFAULT 0',
            'inss_estimado': 'ALTER TABLE ferias_controle ADD COLUMN inss_estimado REAL DEFAULT 0',
            'irrf_estimado': 'ALTER TABLE ferias_controle ADD COLUMN irrf_estimado REAL DEFAULT 0',
            'liquido_estimado': 'ALTER TABLE ferias_controle ADD COLUMN liquido_estimado REAL DEFAULT 0'
        }.items():
            if _col not in cols_ferias:
                db.execute(_sql)

        db.execute('''CREATE TABLE IF NOT EXISTS documentos_rh (
            id INTEGER PRIMARY KEY AUTOINCREMENT, funcionario_id INTEGER, tipo TEXT, data TEXT,
            titulo TEXT, observacao TEXT, arquivo TEXT, ativo INTEGER DEFAULT 1)''')
        db.execute('''CREATE TABLE IF NOT EXISTS epis (
            id INTEGER PRIMARY KEY AUTOINCREMENT, funcionario_id INTEGER, item TEXT, ca TEXT,
            data_entrega TEXT, validade TEXT, data_devolucao TEXT, observacao TEXT, ativo INTEGER DEFAULT 1)''')
        db.execute('''CREATE TABLE IF NOT EXISTS exames (
            id INTEGER PRIMARY KEY AUTOINCREMENT, funcionario_id INTEGER, tipo TEXT,
            data_exame TEXT, validade TEXT, clinica TEXT, observacao TEXT, ativo INTEGER DEFAULT 1)''')
        db.execute('''CREATE TABLE IF NOT EXISTS agenda_rh (
            id INTEGER PRIMARY KEY AUTOINCREMENT, data TEXT NOT NULL, titulo TEXT NOT NULL,
            tipo TEXT DEFAULT 'Lembrete', funcionario_id INTEGER, observacao TEXT, concluido INTEGER DEFAULT 0)''')
        db.execute("INSERT OR IGNORE INTO escalas(codigo,nome,tipo,descricao,dias,ativo) VALUES('E001','SEGUNDA A SEXTA','5x2','Trabalho de segunda a sexta com sábado compensado.','Segunda,Terça,Quarta,Quinta,Sexta',1)")
        db.execute("INSERT OR IGNORE INTO escalas(codigo,nome,tipo,descricao,dias,ativo) VALUES('E002','12X36','12x36','Escala de revezamento 12 horas trabalhadas por 36 horas de descanso.','Conforme escala 12x36',1)")
        # Migração automática: adiciona o campo de abono nas ocorrências de bancos existentes.
        cols_occ = [r[1] for r in db.execute('PRAGMA table_info(ocorrencias)').fetchall()]
        if 'abona' not in cols_occ:
            db.execute("ALTER TABLE ocorrencias ADD COLUMN abona INTEGER DEFAULT 1")
        db.execute("UPDATE ocorrencias SET abona=1 WHERE abona IS NULL")
        # Migração automática: campos novos no cadastro de jornadas.
        cols_jorn = [r[1] for r in db.execute('PRAGMA table_info(jornadas)').fetchall()]
        if 'sabado_tratamento' not in cols_jorn:
            db.execute("ALTER TABLE jornadas ADD COLUMN sabado_tratamento TEXT DEFAULT 'COMPENSADO'")
        if 'descanso_semanal' not in cols_jorn:
            db.execute("ALTER TABLE jornadas ADD COLUMN descanso_semanal TEXT DEFAULT 'DOMINGO'")
        db.execute("UPDATE jornadas SET sabado_tratamento='COMPENSADO' WHERE sabado_tratamento IS NULL OR sabado_tratamento=''")
        db.execute("UPDATE jornadas SET descanso_semanal='DOMINGO' WHERE descanso_semanal IS NULL OR descanso_semanal=''")
        # Correção V10: jornadas 12x36 não têm sábado compensado/não aplicado; o sábado segue a escala.
        db.execute("UPDATE jornadas SET sabado_tratamento='CONFORME ESCALA', descanso_semanal='CONFORME ESCALA 12X36', dias_semana='CONFORME ESCALA 12X36' WHERE UPPER(COALESCE(tipo,'')) LIKE '%12X36%' OR UPPER(COALESCE(nome,'')) LIKE '%12X36%'")
        db.execute("INSERT OR IGNORE INTO usuarios(usuario,senha,perfil,ativo) VALUES('admin','admin','Administrador',1)")
        db.execute("INSERT OR IGNORE INTO setores(nome,descricao,ativo) VALUES('GERAL','Setor padrão',1)")
        # Carga inicial de feriados nacionais e municipais de Itajaí-SC (anos próximos).
        # Inserção feita na própria conexão para evitar bloqueios no SQLite.
        for _ano in range(2026, 2031):
            for _data, _nome in feriados_nacionais_padrao(_ano):
                if not db.execute('SELECT id FROM feriados WHERE data=? AND UPPER(nome)=?', (_data.isoformat(), _nome.upper())).fetchone():
                    db.execute('INSERT INTO feriados(nome,data,tipo,empresa,setor,observacao,ativo) VALUES(?,?,?,?,?,?,1)', (_nome.upper(), _data.isoformat(), 'Nacional', 'TODAS', 'TODOS', 'Carga inicial automática.'))
            for _data, _nome in feriados_itajai_padrao(_ano):
                if not db.execute('SELECT id FROM feriados WHERE data=? AND UPPER(nome)=?', (_data.isoformat(), _nome.upper())).fetchone():
                    db.execute('INSERT INTO feriados(nome,data,tipo,empresa,setor,observacao,ativo) VALUES(?,?,?,?,?,?,1)', (_nome.upper(), _data.isoformat(), 'Municipal', 'TODAS', 'TODOS', 'Carga inicial automática de feriados municipais de Itajaí-SC.'))
        db.execute('''INSERT OR IGNORE INTO jornadas(codigo,nome,entrada1,saida1,entrada2,saida2,horas_dia,horas_semana,intervalo,dias_semana,tipo,sabado_tratamento,descanso_semanal,observacao,ativo)
                      VALUES('J001','COMERCIAL 44H','08:00','12:00','13:30','17:30','08:00','44:00','90','Segunda,Terça,Quarta,Quinta,Sexta','Comercial','COMPENSADO','DOMINGO','Jornada padrão comercial',1)''')
        db.execute('''INSERT OR IGNORE INTO jornadas(codigo,nome,entrada1,saida1,entrada2,saida2,horas_dia,horas_semana,intervalo,dias_semana,tipo,sabado_tratamento,descanso_semanal,observacao,ativo)
                      VALUES('J002','12X36 DIURNO','07:00','19:00','','','12:00','36:00','0','CONFORME ESCALA 12X36','Escala 12x36','CONFORME ESCALA','CONFORME ESCALA 12X36','Plantão diurno 12x36, inclusive sábados quando recaírem na escala',1)''')
        # Migração automática: adiciona o campo setor em bancos já existentes, sem perder dados.
        cols = [r[1] for r in db.execute('PRAGMA table_info(funcionarios)').fetchall()]
        if 'setor' not in cols:
            db.execute("ALTER TABLE funcionarios ADD COLUMN setor TEXT DEFAULT 'GERAL'")
        if 'horario_trabalho' not in cols:
            db.execute("ALTER TABLE funcionarios ADD COLUMN horario_trabalho TEXT DEFAULT ''")
        # Atualização automática do banco para versões antigas
        cols = [r[1] for r in db.execute('PRAGMA table_info(funcionarios)').fetchall()]
        if 'setor' not in cols:
            db.execute("ALTER TABLE funcionarios ADD COLUMN setor TEXT DEFAULT 'GERAL'")
        if 'horario_trabalho' not in cols:
            db.execute("ALTER TABLE funcionarios ADD COLUMN horario_trabalho TEXT DEFAULT ''")
        db.execute('INSERT OR IGNORE INTO empresa(id,nome,cnpj,endereco,numero,bairro,cidade,uf) VALUES(1,?,?,?,?,?,?,?)',
                   (DEFAULT_EMPRESA['nome'], DEFAULT_EMPRESA['cnpj'], DEFAULT_EMPRESA['endereco'], DEFAULT_EMPRESA['numero'], DEFAULT_EMPRESA['bairro'], DEFAULT_EMPRESA['cidade'], DEFAULT_EMPRESA['uf']))
        cur = db.execute('SELECT COUNT(*) FROM funcionarios')
        if cur.fetchone()[0] == 0:
            for nome, adm, funcao, sal in DEFAULT_FUNCIONARIOS:
                db.execute('''INSERT INTO funcionarios(nome,admissao,funcao,setor,salario,jornada,horario_trabalho,descanso,sabado,ativo)
                              VALUES(?,?,?,?,?,?,?,?,?,1)''',
                           (nome, adm, funcao, 'GERAL', sal, 'HORÁRIO DE TRABALHO DE SEGUNDA A SEXTA-FEIRA', '07:42 - 12:00 / 13:30 - 18:00', 'REMUNERADO', 'COMPENSADO'))
        # Sincroniza setores existentes no cadastro de funcionários
        for (setor_nome,) in db.execute("SELECT DISTINCT UPPER(COALESCE(NULLIF(TRIM(setor),''),'GERAL')) FROM funcionarios").fetchall():
            db.execute("INSERT OR IGNORE INTO setores(nome,descricao,ativo) VALUES(?,?,1)", (setor_nome or 'GERAL', 'Importado dos funcionários'))
    if not new:
        backup_db()


def backup_db():
    if not os.path.exists(DB_PATH):
        return
    stamp = datetime.now().strftime('%Y%m%d_%H%M%S')
    shutil.copy2(DB_PATH, os.path.join(BACKUP_DIR, f'ponto_saude_{stamp}.db'))

def log_action(usuario, acao, detalhes=''):
    try:
        with con() as db:
            db.execute('INSERT INTO logs(data_hora,usuario,acao,detalhes) VALUES(?,?,?,?)',
                       (datetime.now().strftime('%d/%m/%Y %H:%M:%S'), usuario or 'sistema', acao, detalhes))
    except Exception:
        pass




def registrar_pdf(usuario, tipo, setor, mes, ano, arquivo, quantidade):
    try:
        with con() as db:
            db.execute('INSERT INTO pdf_historico(data_hora,usuario,tipo,setor,mes,ano,arquivo,quantidade) VALUES(?,?,?,?,?,?,?,?)',
                       (datetime.now().strftime('%d/%m/%Y %H:%M:%S'), usuario or 'sistema', tipo, setor or '', mes, ano, arquivo, quantidade))
    except Exception:
        pass

def autenticar(usuario, senha):
    with con() as db:
        row = db.execute('SELECT perfil FROM usuarios WHERE usuario=? AND senha=? AND ativo=1', (usuario, senha)).fetchone()
    return row[0] if row else None


class LoginDialog(tk.Tk):
    def __init__(self, usuario='admin', perfil='Administrador'):
        super().__init__()
        self.usuario = usuario
        self.perfil = perfil
        self.title(APP_NAME + ' - Login')
        self.geometry('430x430')
        self.resizable(False, False)
        self.configure(bg='#f8fafc')
        self.usuario = None
        self.perfil = None

        card = tk.Frame(self, bg='white', highlightbackground='#d1d5db', highlightthickness=1)
        card.pack(fill='both', expand=True, padx=28, pady=24)

        try:
            if os.path.exists(LOGO_APP):
                self.login_logo_img = tk.PhotoImage(file=LOGO_APP)
                tk.Label(card, image=self.login_logo_img, bg='white').pack(pady=(22,8))
        except Exception:
            pass

        tk.Label(card, text='Sistema Gestão Izzant', bg='white', fg='#111827', font=('Arial', 18, 'bold')).pack(pady=(2,2))
        tk.Label(card, text='Acesso restrito ao sistema', bg='white', fg='#6b7280', font=('Arial', 10)).pack(pady=(0,4))
        tk.Label(card, text='Enterprise v1.6.3 Interface', bg='white', fg='#9ca3af', font=('Arial', 9)).pack(pady=(0,14))

        frm = ttk.Frame(card, padding=(28, 4, 28, 18))
        frm.pack(fill='x')
        ttk.Label(frm, text='Usuário').pack(anchor='w')
        self.user_var = tk.StringVar(value='admin')
        ttk.Entry(frm, textvariable=self.user_var).pack(fill='x', pady=(0,10))
        ttk.Label(frm, text='Senha').pack(anchor='w')
        self.pass_var = tk.StringVar(value='admin')
        ttk.Entry(frm, textvariable=self.pass_var, show='*').pack(fill='x', pady=(0,16))
        ttk.Button(frm, text='Entrar no sistema', command=self.try_login).pack(fill='x')
        ttk.Label(frm, text='Senha inicial: admin / admin', font=('Arial', 8)).pack(anchor='w', pady=(10,0))
        self.bind('<Return>', lambda e: self.try_login())

    def try_login(self):
        perfil = autenticar(self.user_var.get().strip(), self.pass_var.get().strip())
        if not perfil:
            messagebox.showerror('Acesso negado', 'Usuário ou senha inválidos.')
            return
        self.usuario = self.user_var.get().strip()
        self.perfil = perfil
        log_action(self.usuario, 'LOGIN', 'Entrada no sistema')
        self.destroy()



def pascoa(ano):
    """Calcula a data da Páscoa pelo algoritmo de Meeus/Jones/Butcher."""
    a = ano % 19
    b = ano // 100
    c = ano % 100
    d = b // 4
    e = b % 4
    f = (b + 8) // 25
    g = (b - f + 1) // 3
    h = (19 * a + b - d - g + 15) % 30
    i = c // 4
    k = c % 4
    l = (32 + 2 * e + 2 * i - h - k) % 7
    m = (a + 11 * h + 22 * l) // 451
    mes = (h + l - 7 * m + 114) // 31
    dia = ((h + l - 7 * m + 114) % 31) + 1
    return date(ano, mes, dia)


def feriados_nacionais_padrao(ano):
    p = pascoa(ano)
    itens = [
        (date(ano, 1, 1), 'CONFRATERNIZAÇÃO UNIVERSAL'),
        (p - timedelta(days=48), 'CARNAVAL'),
        (p - timedelta(days=47), 'CARNAVAL'),
        (p - timedelta(days=46), 'QUARTA-FEIRA DE CINZAS'),
        (p - timedelta(days=2), 'SEXTA-FEIRA SANTA'),
        (p, 'PÁSCOA'),
        (date(ano, 4, 21), 'TIRADENTES'),
        (date(ano, 5, 1), 'DIA DO TRABALHO'),
        (date(ano, 9, 7), 'INDEPENDÊNCIA DO BRASIL'),
        (date(ano, 10, 12), 'NOSSA SENHORA APARECIDA'),
        (date(ano, 11, 2), 'FINADOS'),
        (date(ano, 11, 15), 'PROCLAMAÇÃO DA REPÚBLICA'),
        (date(ano, 11, 20), 'CONSCIÊNCIA NEGRA'),
        (date(ano, 12, 25), 'NATAL'),
    ]
    return itens


def feriados_itajai_padrao(ano):
    p = pascoa(ano)
    return [
        (p - timedelta(days=2), 'SEXTA-FEIRA SANTA'),
        (p + timedelta(days=60), 'CORPUS CHRISTI'),
        (date(ano, 6, 15), 'ANIVERSÁRIO DO MUNICÍPIO DE ITAJAÍ'),
        (date(ano, 12, 8), 'NOSSA SENHORA DA IMACULADA CONCEIÇÃO'),
    ]


def inserir_feriados_padrao(ano, usuario='sistema'):
    """Insere/atualiza feriados nacionais e municipais de Itajaí-SC no banco local."""
    total = 0
    obs_nac = 'Importado automaticamente. Fonte preferencial: BrasilAPI; fallback: calendário nacional calculado pelo sistema.'
    obs_mun = 'Feriados municipais de Itajaí-SC cadastrados no sistema.'
    with con() as db:
        for data, nome in feriados_nacionais_padrao(ano):
            existe = db.execute('SELECT id FROM feriados WHERE data=? AND UPPER(nome)=?', (data.isoformat(), nome.upper())).fetchone()
            if not existe:
                db.execute('INSERT INTO feriados(nome,data,tipo,empresa,setor,observacao,ativo) VALUES(?,?,?,?,?,?,1)', (nome.upper(), data.isoformat(), 'Nacional', 'TODAS', 'TODOS', obs_nac))
                total += 1
        for data, nome in feriados_itajai_padrao(ano):
            existe = db.execute('SELECT id FROM feriados WHERE data=? AND UPPER(nome)=?', (data.isoformat(), nome.upper())).fetchone()
            if not existe:
                db.execute('INSERT INTO feriados(nome,data,tipo,empresa,setor,observacao,ativo) VALUES(?,?,?,?,?,?,1)', (nome.upper(), data.isoformat(), 'Municipal', 'TODAS', 'TODOS', obs_mun))
                total += 1
    log_action(usuario, 'FERIADOS', f'Importados feriados padrão de {ano}: {total}')
    return total


def importar_nacionais_brasilapi(ano):
    """Tenta buscar feriados nacionais na BrasilAPI. Se falhar, usa fallback local."""
    url = f'https://brasilapi.com.br/api/feriados/v1/{ano}'
    try:
        with urllib.request.urlopen(url, timeout=8) as resp:
            dados = json.loads(resp.read().decode('utf-8'))
        itens=[]
        for item in dados:
            data_txt = item.get('date')
            nome = (item.get('name') or '').upper()
            tipo = (item.get('type') or '').upper()
            if data_txt and nome:
                itens.append((date.fromisoformat(data_txt), nome, tipo))
        return itens, 'BrasilAPI'
    except Exception:
        return [(d,n,'NATIONAL') for d,n in feriados_nacionais_padrao(ano)], 'Fallback local'


def importar_feriados_brasil_itajai(ano, usuario='sistema'):
    nacionais, fonte = importar_nacionais_brasilapi(ano)
    total=0
    with con() as db:
        for data, nome, _tipo in nacionais:
            existe = db.execute('SELECT id FROM feriados WHERE data=? AND UPPER(nome)=?', (data.isoformat(), nome.upper())).fetchone()
            if not existe:
                db.execute('INSERT INTO feriados(nome,data,tipo,empresa,setor,observacao,ativo) VALUES(?,?,?,?,?,?,1)', (nome.upper(), data.isoformat(), 'Nacional', 'TODAS', 'TODOS', f'Importado automaticamente. Fonte: {fonte}.'))
                total += 1
        for data, nome in feriados_itajai_padrao(ano):
            existe = db.execute('SELECT id FROM feriados WHERE data=? AND UPPER(nome)=?', (data.isoformat(), nome.upper())).fetchone()
            if not existe:
                db.execute('INSERT INTO feriados(nome,data,tipo,empresa,setor,observacao,ativo) VALUES(?,?,?,?,?,?,1)', (nome.upper(), data.isoformat(), 'Municipal', 'TODAS', 'TODOS', 'Feriado municipal de Itajaí-SC cadastrado automaticamente.'))
                total += 1
    log_action(usuario, 'FERIADOS', f'Importados feriados Brasil/Itajaí {ano}: {total} ({fonte})')
    return total, fonte



def garantir_feriados_padrao(anos=None, usuario='sistema'):
    """Garante que os feriados nacionais e municipais de Itajaí-SC estejam cadastrados.
    Aceita um ano único ou uma lista/range de anos. Usado pela consulta e pela geração da folha.
    """
    if anos is None:
        anos = range(datetime.now().year, datetime.now().year + 5)
    if isinstance(anos, int):
        anos = [anos]
    total = 0
    with con() as db:
        for ano in anos:
            for data, nome in feriados_nacionais_padrao(int(ano)):
                existe = db.execute('SELECT id FROM feriados WHERE data=? AND UPPER(nome)=?', (data.isoformat(), nome.upper())).fetchone()
                if not existe:
                    db.execute('INSERT INTO feriados(nome,data,tipo,empresa,setor,observacao,ativo) VALUES(?,?,?,?,?,?,1)',
                               (nome.upper(), data.isoformat(), 'Nacional', 'TODAS', 'TODOS', 'Carga automática Brasil.'))
                    total += 1
            for data, nome in feriados_itajai_padrao(int(ano)):
                existe = db.execute('SELECT id FROM feriados WHERE data=? AND UPPER(nome)=?', (data.isoformat(), nome.upper())).fetchone()
                if not existe:
                    db.execute('INSERT INTO feriados(nome,data,tipo,empresa,setor,observacao,ativo) VALUES(?,?,?,?,?,?,1)',
                               (nome.upper(), data.isoformat(), 'Municipal', 'TODAS', 'TODOS', 'Carga automática Itajaí-SC.'))
                    total += 1
    if total:
        log_action(usuario, 'FERIADOS', f'Garantida carga automática: {total} feriados')
    return total

def get_feriados_mes(mes, ano, setor='TODOS'):
    inicio = date(ano, mes, 1).isoformat()
    fim = date(ano, mes, calendar.monthrange(ano, mes)[1]).isoformat()
    mapa = {}
    try:
        with con() as db:
            rows = db.execute("""SELECT nome,data,tipo,setor FROM feriados
                                 WHERE ativo=1 AND date(data) BETWEEN date(?) AND date(?)
                                 AND (UPPER(COALESCE(setor,'TODOS'))='TODOS' OR UPPER(COALESCE(setor,''))=UPPER(?))
                                 ORDER BY data,tipo""", (inicio, fim, setor or 'TODOS')).fetchall()
        for nome, data_txt, tipo, _setor in rows:
            try:
                d = date.fromisoformat(data_txt)
            except Exception:
                continue
            mapa[d.day] = {'nome': nome or 'FERIADO', 'tipo': tipo or 'Feriado'}
    except Exception:
        pass
    return mapa


def fmt_data(iso):
    if not iso:
        return ''
    try:
        y,m,d = iso.split('-')
        return f'{d}/{m}/{y}'
    except Exception:
        return iso



def parse_data_br(valor):
    """Aceita DD/MM/AAAA, DDMMAAAA ou AAAA-MM-DD e retorna objeto date."""
    valor = (valor or '').strip()
    if not valor:
        raise ValueError('Data vazia')
    # mantém compatibilidade com datas antigas digitadas no padrão ISO
    if '-' in valor:
        return date.fromisoformat(valor)
    somente = ''.join(ch for ch in valor if ch.isdigit())
    if len(somente) == 8:
        dia = int(somente[0:2]); mes = int(somente[2:4]); ano = int(somente[4:8])
        return date(ano, mes, dia)
    if '/' in valor:
        partes = valor.split('/')
        if len(partes) == 3:
            dia, mes, ano = map(int, partes)
            return date(ano, mes, dia)
    raise ValueError('Formato inválido')


def norm_txt(valor):
    """Normaliza texto para regras internas, preservando o texto original na tela."""
    mapa = str.maketrans('ÁÀÂÃÄÉÈÊËÍÌÎÏÓÒÔÕÖÚÙÛÜÇáàâãäéèêëíìîïóòôõöúùûüç',
                         'AAAAAEEEEIIIIOOOOOUUUUCaaaaaeeeeiiiiooooouuuuc')
    return str(valor or '').translate(mapa).upper().strip()


def normalizar_datas_feriados():
    """Mantém datas dos feriados em ISO para filtro e ordenação, aceitando legados em DD/MM/AAAA."""
    try:
        with con() as db:
            rows = db.execute("SELECT id,data FROM feriados WHERE data IS NOT NULL AND data<>''").fetchall()
            for rid, data_txt in rows:
                txt = str(data_txt or '').strip()
                if '/' in txt:
                    try:
                        iso = parse_data_br(txt).isoformat()
                        db.execute('UPDATE feriados SET data=? WHERE id=?', (iso, rid))
                    except Exception:
                        pass
    except Exception:
        pass


def data_para_iso(valor):
    return parse_data_br(valor).isoformat()

def parse_moeda_br(valor):
    """Converte valores brasileiros como R$ 2.500,00 em float."""
    if valor is None:
        return 0.0
    if isinstance(valor, (int, float)):
        return float(valor)
    txt = str(valor).strip().replace('R$', '').replace(' ', '')
    if not txt:
        return 0.0
    if ',' in txt:
        txt = txt.replace('.', '').replace(',', '.')
    try:
        return float(txt)
    except Exception:
        return 0.0


def moeda_br(valor):
    try:
        v = float(valor or 0)
    except Exception:
        v = 0.0
    return ('R$ {:,.2f}'.format(v)).replace(',', 'X').replace('.', ',').replace('X', '.')


def adicionar_dias_corridos(inicio, dias):
    if not inicio or not dias:
        return None, None
    fim = inicio + timedelta(days=int(dias) - 1)
    retorno = fim + timedelta(days=1)
    return fim, retorno


def calcular_inss_estimado(base):
    """Estimativa simplificada e progressiva de INSS para conferência interna."""
    base = max(0.0, float(base or 0))
    faixas = [(1518.00, 0.075), (2793.88, 0.09), (4190.83, 0.12), (8157.41, 0.14)]
    anterior = 0.0
    imposto = 0.0
    for limite, aliquota in faixas:
        if base > anterior:
            tributavel = min(base, limite) - anterior
            imposto += max(0, tributavel) * aliquota
            anterior = limite
        if base <= limite:
            break
    return round(imposto, 2)


def calcular_irrf_estimado(base, inss=0.0, dependentes=0):
    """Estimativa simplificada de IRRF mensal para projeção. Pode ser ajustada conforme tabela vigente."""
    ded_dep = 189.59 * int(dependentes or 0)
    b = max(0.0, float(base or 0) - float(inss or 0) - ded_dep)
    tabela = [
        (2259.20, 0.0, 0.0),
        (2826.65, 0.075, 169.44),
        (3751.05, 0.15, 381.44),
        (4664.68, 0.225, 662.77),
        (10**12, 0.275, 896.00),
    ]
    for limite, aliq, ded in tabela:
        if b <= limite:
            return round(max(0.0, b * aliq - ded), 2)
    return 0.0


def calcular_ferias_valores(salario_base, dias_ferias, dias_abono=0, media_variaveis=0, adiantamento_13=False):
    """Calcula valores estimados de férias com base em dias corridos.
    Observação: os valores são estimativos para conferência e emissão interna.
    """
    salario = parse_moeda_br(salario_base)
    media = parse_moeda_br(media_variaveis)
    base_mensal = salario + media
    dias = max(0, int(dias_ferias or 0))
    abono = max(0, int(dias_abono or 0))
    valor_ferias = base_mensal / 30.0 * dias
    valor_um_terco = valor_ferias / 3.0
    valor_abono = base_mensal / 30.0 * abono
    # 1/3 sobre o abono pecuniário, quando houver
    valor_abono_terco = valor_abono / 3.0 if valor_abono else 0.0
    valor_13 = base_mensal / 2.0 if adiantamento_13 else 0.0
    total_bruto = valor_ferias + valor_um_terco + valor_abono + valor_abono_terco + valor_13
    inss = calcular_inss_estimado(total_bruto)
    irrf = calcular_irrf_estimado(total_bruto, inss)
    liquido = total_bruto - inss - irrf
    return {
        'salario_base': salario, 'media_variaveis': media,
        'valor_ferias': round(valor_ferias,2), 'valor_um_terco': round(valor_um_terco,2),
        'valor_abono': round(valor_abono + valor_abono_terco,2), 'valor_13': round(valor_13,2),
        'total_bruto': round(total_bruto,2), 'inss_estimado': round(inss,2),
        'irrf_estimado': round(irrf,2), 'liquido_estimado': round(liquido,2)
    }


def formatar_campo_data(var):
    """Formata StringVar visualmente para DD/MM/AAAA ao sair do campo."""
    try:
        var.set(parse_data_br(var.get()).strftime('%d/%m/%Y'))
    except Exception:
        pass


def safe_name(s):
    keep = ''.join(ch if ch.isalnum() or ch in ' _-' else '_' for ch in s)
    return '_'.join(keep.split())[:80]


def money(v):
    try:
        return f'{float(v):,.2f}'.replace(',', 'X').replace('.', ',').replace('X','.')
    except Exception:
        return ''


def get_empresa():
    with con() as db:
        row = db.execute('SELECT nome,cnpj,endereco,numero,bairro,cidade,uf FROM empresa WHERE id=1').fetchone()
    keys = ['nome','cnpj','endereco','numero','bairro','cidade','uf']
    return dict(zip(keys, row)) if row else DEFAULT_EMPRESA.copy()


def get_funcionarios(ativos=True):
    q = "SELECT id,nome,cpf,ctps,admissao,funcao,cbo,salario,setor,jornada,COALESCE(horario_trabalho,''),descanso,sabado,ativo FROM funcionarios"
    if ativos:
        q += ' WHERE ativo=1'
    q += " ORDER BY COALESCE(setor, ''), nome"
    with con() as db:
        rows = db.execute(q).fetchall()
    keys = ['id','nome','cpf','ctps','admissao','funcao','cbo','salario','setor','jornada','horario_trabalho','descanso','sabado','ativo']
    return [dict(zip(keys, r)) for r in rows]


def get_setores(ativos=True):
    q = 'SELECT nome FROM setores'
    if ativos:
        q += ' WHERE ativo=1'
    q += ' ORDER BY nome'
    try:
        with con() as db:
            rows = db.execute(q).fetchall()
        setores = [r[0] for r in rows if r and r[0]]
    except Exception:
        setores = []
    if not setores:
        funcs = get_funcionarios(ativos)
        setores = sorted({(f.get('setor') or 'GERAL').strip().upper() for f in funcs if (f.get('setor') or 'GERAL').strip()})
    return setores or ['GERAL']

def get_jornadas(ativas=True):
    q = 'SELECT codigo,nome,entrada1,saida1,entrada2,saida2,horas_dia,horas_semana,intervalo,dias_semana,tipo,sabado_tratamento,descanso_semanal,observacao,ativo FROM jornadas'
    if ativas:
        q += ' WHERE ativo=1'
    q += ' ORDER BY ativo DESC, nome'
    with con() as db:
        rows = db.execute(q).fetchall()
    keys = ['codigo','nome','entrada1','saida1','entrada2','saida2','horas_dia','horas_semana','intervalo','dias_semana','tipo','sabado_tratamento','descanso_semanal','observacao','ativo']
    return [dict(zip(keys,r)) for r in rows]

def get_jornada_por_nome(nome):
    if not nome:
        return None
    nome = str(nome).strip().upper()
    for j in get_jornadas(False):
        if str(j.get('nome') or '').strip().upper() == nome:
            return j
    return None

def texto_horario_jornada(j):
    if not j:
        return '07:42 - 12:00 / 13:30 - 18:00'
    e1, s1 = j.get('entrada1') or '', j.get('saida1') or ''
    e2, s2 = j.get('entrada2') or '', j.get('saida2') or ''
    partes = []
    if e1 or s1:
        partes.append(f"{e1} - {s1}".strip())
    if e2 or s2:
        partes.append(f"{e2} - {s2}".strip())
    return ' / '.join(partes) if partes else (j.get('nome') or '')


def normalizar_jornada_12x36(j):
    """Ajusta a exibição de jornadas 12x36: sábados e DSR devem seguir a escala."""
    if not j:
        return j
    tipo = str(j.get('tipo') or '').upper().replace(' ', '')
    nome = str(j.get('nome') or '').upper().replace(' ', '')
    if '12X36' in tipo or '12X36' in nome:
        j = dict(j)
        j['sabado_tratamento'] = 'CONFORME ESCALA'
        j['descanso_semanal'] = 'CONFORME ESCALA 12X36'
        if not j.get('dias_semana') or '12X36' not in str(j.get('dias_semana')).upper():
            j['dias_semana'] = 'CONFORME ESCALA 12X36'
    return j



def is_jornada_12x36(jornada_info):
    """Identifica jornada 12x36 para regras específicas de feriado/descanso."""
    if not jornada_info:
        return False
    tipo = str(jornada_info.get('tipo') or '').upper().replace(' ', '')
    nome = str(jornada_info.get('nome') or '').upper().replace(' ', '')
    dsr = str(jornada_info.get('descanso_semanal') or '').upper().replace(' ', '')
    dias = str(jornada_info.get('dias_semana') or '').upper().replace(' ', '')
    return ('12X36' in tipo) or ('12X36' in nome) or ('12X36' in dsr) or ('12X36' in dias)

def get_ocorrencias_mes(funcionario_id, mes, ano):
    """Retorna um dicionário {dia: ocorrencia} para preencher a folha sem mudar o layout."""
    inicio_mes = date(ano, mes, 1)
    fim_mes = date(ano, mes, calendar.monthrange(ano, mes)[1])
    mapa = {}
    try:
        with con() as db:
            rows = db.execute("""SELECT tipo, data_inicio, data_fim, COALESCE(observacao,''), COALESCE(abona,1)
                                 FROM ocorrencias
                                 WHERE funcionario_id=? AND ativo=1
                                 AND date(data_inicio) <= date(?) AND date(data_fim) >= date(?)
                                 ORDER BY data_inicio""",
                              (funcionario_id, fim_mes.isoformat(), inicio_mes.isoformat())).fetchall()
        for tipo, ini, fim, obs, abona in rows:
            try:
                d1 = max(date.fromisoformat(ini), inicio_mes)
                d2 = min(date.fromisoformat(fim), fim_mes)
            except Exception:
                continue
            d = d1
            while d <= d2:
                mapa[d.day] = {'tipo': (tipo or '').upper(), 'observacao': obs or '', 'abona': int(abona or 0)}
                d += timedelta(days=1)
    except Exception:
        pass
    return mapa

def setor_em_uso(nome):
    with con() as db:
        row = db.execute("SELECT COUNT(*) FROM funcionarios WHERE UPPER(COALESCE(setor,'GERAL'))=?", (nome.strip().upper(),)).fetchone()
    return bool(row and row[0])


def draw_cell(c, x, y, w, h, txt='', font='Helvetica', size=7, bold=False, align='left', valign='middle', fill=None, rotate=False):
    if fill:
        c.setFillColor(fill); c.rect(x,y,w,h,fill=1,stroke=0); c.setFillColor(colors.black)
    c.rect(x,y,w,h,fill=0,stroke=1)
    if txt is None: txt=''
    txt = str(txt)
    if not txt:
        return
    fn = 'Helvetica-Bold' if bold else font
    c.setFont(fn, size)
    if rotate:
        c.saveState(); c.translate(x+w/2,y+h/2); c.rotate(90); c.drawCentredString(0,-size/2,txt); c.restoreState(); return
    max_chars = max(1, int(w/(size*0.48)))
    lines=[]
    for part in txt.split('\n'):
        words=part.split(' '); line=''
        for word in words:
            if len((line+' '+word).strip()) <= max_chars:
                line=(line+' '+word).strip()
            else:
                lines.append(line); line=word
        lines.append(line)
    total_h=len(lines)*(size+1)
    ty = y + (h-total_h)/2 + total_h - size if valign=='middle' else y+h-size-2
    for line in lines[:max(1, int(h/(size+1))+1)]:
        if align=='center': c.drawCentredString(x+w/2, ty, line)
        elif align=='right': c.drawRightString(x+w-2, ty, line)
        else: c.drawString(x+2, ty, line)
        ty -= size+1



def status_dia_por_jornada(jornada_info, data_dia):
    """Retorna o bloqueio automático do dia conforme a jornada.

    Regra estabilizada v1.5:
    - Jornada 12x36: sábado/domingo não são bloqueados automaticamente; seguem a escala.
    - Domingo como descanso semanal aparece como DSR na folha, como nas versões aprovadas.
    - Sábado compensado aparece como COMP. / COMPENSADO.
    - O DSR/compensado tem prioridade sobre feriado na geração da folha.
    """
    if not jornada_info or not data_dia:
        return ''
    j = normalizar_jornada_12x36(dict(jornada_info))
    tipo = norm_txt(j.get('tipo'))
    nome = norm_txt(j.get('nome'))
    sab = norm_txt(j.get('sabado_tratamento'))
    dsr = norm_txt(j.get('descanso_semanal'))

    # Em 12x36, sábados e domingos devem seguir a escala. Não lançar DSR automático.
    if '12X36' in tipo or '12X36' in nome or 'ESCALA 12X36' in dsr or 'CONFORME ESCALA' in dsr:
        return ''

    # Compatibilidade com cadastros antigos: "REMUNERADO" no campo descanso
    # significa descanso semanal remunerado no domingo para jornadas comuns.
    if dsr in ('REMUNERADO', 'DESCANSO REMUNERADO'):
        dsr = 'DOMINGO'

    wd = data_dia.weekday()  # Segunda=0 ... Domingo=6
    nomes = ['SEGUNDA','TERCA','QUARTA','QUINTA','SEXTA','SABADO','DOMINGO']
    dia_nome = nomes[wd]

    # Descanso semanal configurado: no domingo, a folha deve mostrar apenas DSR.
    if wd == 6 and ('DOMINGO' in dsr or 'DSR' in dsr or not dsr):
        return 'DSR'
    if dsr == 'SABADO E DOMINGO' and wd in (5, 6):
        return 'DSR' if wd == 6 else 'COMPENSADO'
    if dsr == dia_nome:
        return 'DSR' if wd == 6 else 'DESCANSO SEMANAL'

    # Tratamento específico do sábado.
    if wd == 5:
        if 'COMPENSADO' in sab and 'NAO' not in sab:
            return 'COMPENSADO'
        if 'NAO SE APLICA' in sab:
            return 'NÃO SE APLICA'
        # NÃO COMPENSADO, TRABALHADO ou CONFORME ESCALA não bloqueiam a linha.
        return ''

    return ''

def abreviar_texto_pdf(valor, contexto='linha'):
    """Abrevia textos longos para caber nas células estreitas da folha de ponto."""
    v = (valor or '').strip().upper()
    if not v:
        return ''
    if v == 'DSR' or 'DESCANSO SEMANAL' in v:
        return 'DSR'
    if v == 'DOMINGO':
        return 'DOMINGO'
    if v == 'SÁBADO E DOMINGO' or v == 'SABADO E DOMINGO':
        return 'SÁB/DOM'
    if 'COMPENSADO' in v and 'NÃO' not in v:
        return 'COMP.' if contexto == 'linha' else 'COMPENSADO'
    if 'NÃO COMPENSADO' in v or 'NAO COMPENSADO' in v:
        return 'NÃO COMP.'
    if 'CONFORME ESCALA' in v:
        return 'ESCALA'
    if 'NÃO SE APLICA' in v or 'NAO SE APLICA' in v:
        return 'N/A'
    return valor

def gerar_pdf_funcionarios(funcionarios, mes, ano, destino):
    """Gera a folha de ponto no desenho da planilha enviada."""
    os.makedirs(os.path.dirname(destino), exist_ok=True)
    emp = get_empresa()
    c = canvas.Canvas(destino, pagesize=A4)
    W, H = A4

    def rect(x, y, w, h, lw=0.6):
        c.setLineWidth(lw)
        c.rect(x, y, w, h, fill=0, stroke=1)

    def txt(x, y, w, h, value='', size=7.0, bold=False, align='left', valign='middle'):
        value = '' if value is None else str(value)
        if not value:
            return
        font = 'Helvetica-Bold' if bold else 'Helvetica'
        c.setFont(font, size)
        lines = value.split('\n')
        line_h = size + 1.2
        if valign == 'top':
            ty = y + h - size - 2
        else:
            ty = y + (h + len(lines) * line_h) / 2 - size
        for line in lines:
            if align == 'center':
                c.drawCentredString(x + w / 2, ty, line)
            elif align == 'right':
                c.drawRightString(x + w - 3, ty, line)
            else:
                c.drawString(x + 2, ty, line)
            ty -= line_h

    def cell(x, y, w, h, value='', size=7.0, bold=False, align='left', valign='middle', lw=0.6):
        rect(x, y, w, h, lw)
        txt(x, y, w, h, value, size, bold, align, valign)

    def label_value(x, y, w, h, label, value, value_size=9.0, value_bold=True, align='left'):
        """Campo do cabeçalho com ajuste automático do texto.
        Usado especialmente para Sábado/Descanso semanal, para que expressões
        como 'CONFORME ESCALA 12X36' caibam sem invadir outros campos.
        """
        rect(x, y, w, h, 0.6)
        c.setFont('Helvetica', 6.4)
        c.drawString(x + 2, y + h - 7.5, label)

        text = str(value or '')
        font = 'Helvetica-Bold' if value_bold else 'Helvetica'
        size = float(value_size)
        max_w = max(8, w - 5)

        # Reduz a fonte automaticamente quando o conteúdo é maior que a célula.
        while size > 5.6 and c.stringWidth(text, font, size) > max_w:
            size -= 0.35

        c.setFont(font, size)
        if c.stringWidth(text, font, size) <= max_w:
            if align == 'center':
                c.drawCentredString(x + w/2, y + 5.0, text)
            else:
                c.drawString(x + 2, y + 5.0, text)
            return

        # Se ainda não couber, quebra em até 2 linhas mantendo dentro do campo.
        palavras = text.split()
        linhas = []
        atual = ''
        for palavra in palavras:
            teste = (atual + ' ' + palavra).strip()
            if atual and c.stringWidth(teste, font, size) > max_w:
                linhas.append(atual)
                atual = palavra
            else:
                atual = teste
        if atual:
            linhas.append(atual)
        if len(linhas) > 2:
            linhas = [linhas[0], ' '.join(linhas[1:])]

        line_h = size + 1
        ty = y + 10.6
        for linha in linhas[:2]:
            if c.stringWidth(linha, font, size) > max_w:
                # último recurso: fonte mínima na segunda linha
                s2 = size
                while s2 > 4.8 and c.stringWidth(linha, font, s2) > max_w:
                    s2 -= 0.25
                c.setFont(font, s2)
            else:
                c.setFont(font, size)
            if align == 'center':
                c.drawCentredString(x + w/2, ty, linha)
            else:
                c.drawString(x + 2, ty, linha)
            ty -= line_h

    for f in funcionarios:
        ocorrencias = get_ocorrencias_mes(f.get('id'), mes, ano)
        feriados_mes = get_feriados_mes(mes, ano, f.get('setor') or 'TODOS')
        c.setFont('Helvetica-Bold', 15)
        c.drawCentredString(W / 2, 814, 'FOLHA DE PONTO INDIVIDUAL DE TRABALHO')

        left = 43
        width = 509
        y_top = 790
        h = 26

        # Cabeçalho - medidas aproximadas do Excel original
        y = y_top - h
        label_value(left, y, 403, h, 'EMPREGADOR / NOME - EMPRESA:', emp.get('nome',''), 9.2)
        label_value(left + 403, y, width - 403, h, 'CEI/CNPJ:', emp.get('cnpj',''), 9.2, True, 'center')

        y -= h
        w_end, w_no, w_bairro, w_cidade, w_uf = 228, 43, 96, 113, 29
        x = left
        label_value(x, y, w_end, h, 'ENDEREÇO / LOGRADOURO:', emp.get('endereco',''), 8.8); x += w_end
        label_value(x, y, w_no, h, 'Nº', emp.get('numero',''), 8.8, True, 'center'); x += w_no
        label_value(x, y, w_bairro, h, 'BAIRRO / DISTRITO', emp.get('bairro',''), 8.8, True, 'center'); x += w_bairro
        label_value(x, y, w_cidade, h, 'CIDADE', emp.get('cidade',''), 8.8, True, 'center'); x += w_cidade
        label_value(x, y, w_uf, h, 'UF', emp.get('uf',''), 8.8, True, 'center')

        y -= h
        w_emp, w_ctps, w_adm = 289, 113, 107
        x = left
        label_value(x, y, w_emp, h, 'EMPREGADO(A):', f.get('nome',''), 8.8); x += w_emp
        label_value(x, y, w_ctps, h, 'CTPS / C.I. Nº e SÉRIE', f.get('ctps') or '0', 8.8, True, 'center'); x += w_ctps
        label_value(x, y, w_adm, h, 'DATA DE ADMISSÃO', fmt_data(f.get('admissao')), 8.8, True, 'center')

        jornada_info = normalizar_jornada_12x36(get_jornada_por_nome(f.get('jornada')))

        # V1.5.2 - Correção: quando o funcionário não tem uma jornada cadastrada vinculada
        # ou usa uma jornada digitada/manual, a folha deve continuar aplicando as regras
        # aprovadas anteriormente: sábado COMPENSADO e domingo DSR.
        if not jornada_info:
            # Mantém o comportamento estável aprovado anteriormente: quando o
            # funcionário não possui uma jornada cadastrada vinculada, a folha
            # assume sábado COMPENSADO e domingo DSR. O campo antigo
            # "REMUNERADO" não deve impedir o DSR no domingo.
            desc_func = (f.get('descanso') or '').strip().upper()
            if not desc_func or desc_func in ('REMUNERADO', 'DESCANSO REMUNERADO'):
                desc_func = 'DOMINGO'
            jornada_info = {
                'nome': f.get('jornada') or '',
                'tipo': f.get('jornada') or '',
                'entrada1': '', 'saida1': '', 'entrada2': '', 'saida2': '',
                'sabado_tratamento': f.get('sabado') or 'COMPENSADO',
                'descanso_semanal': desc_func,
                'dias_semana': ''
            }
        else:
            jornada_info = dict(jornada_info)
            if not jornada_info.get('sabado_tratamento'):
                jornada_info['sabado_tratamento'] = f.get('sabado') or 'COMPENSADO'
            if not jornada_info.get('descanso_semanal'):
                jornada_info['descanso_semanal'] = f.get('descanso') or 'DOMINGO'
            # Jornada comum com descanso antigo "REMUNERADO" deve gerar DSR aos domingos.
            if (str(jornada_info.get('descanso_semanal') or '').strip().upper() in ('REMUNERADO', 'DESCANSO REMUNERADO')) and not is_jornada_12x36(jornada_info):
                jornada_info['descanso_semanal'] = 'DOMINGO'

        # Regra aprovada: na escala 12x36 o feriado não deve sobrescrever a escala.
        # Como o sistema ainda não controla o ciclo individual de plantões, a folha 12x36
        # não exibe automaticamente "FERIADO"; mantém os dias disponíveis conforme escala.
        if is_jornada_12x36(jornada_info):
            feriados_mes = {}
        horario_jornada = (f.get('horario_trabalho') or '').strip() or texto_horario_jornada(jornada_info)
        sabado_info = jornada_info.get('sabado_tratamento') or f.get('sabado') or 'COMPENSADO'
        descanso_info = jornada_info.get('descanso_semanal') or f.get('descanso') or 'DOMINGO'

        y -= h
        w_func, w_sal, w_jornada = 236, 80, 193
        x = left
        label_value(x, y, w_func, h, 'FUNÇÃO', f.get('funcao') or '', 8.8); x += w_func
        label_value(x, y, w_sal, h, 'SALÁRIO BASE R$', money(f.get('salario')), 8.8, True, 'center'); x += w_sal
        label_value(x, y, w_jornada, h, 'HORÁRIO DE TRABALHO', horario_jornada, 8.2, True, 'center')

        y -= h
        w_sab, w_desc, w_mes, w_ano = 236, 97, 131, 45
        x = left
        label_value(x, y, w_sab, h, 'SÁBADOS', abreviar_texto_pdf(sabado_info, 'cabecalho'), 7.0, True, 'center'); x += w_sab
        label_value(x, y, w_desc, h, 'DSR', abreviar_texto_pdf(descanso_info, 'cabecalho'), 6.8, True, 'center'); x += w_desc
        label_value(x, y, w_mes, h, 'MÊS', MESES[mes-1].upper(), 8.8, True, 'center'); x += w_mes
        label_value(x, y, w_ano, h, 'ANO', str(ano), 8.8, True, 'center')

        # Tabela diária
        table_top = y - 10
        header_h1, header_h2 = 18, 18
        total_header_h = header_h1 + header_h2
        row_h = 12.4
        total_h = 17
        cols = [26, 32, 42, 42, 42, 42, 35, 42, 42, 36, 128]
        xs = [left]
        for cw in cols[:-1]:
            xs.append(xs[-1] + cw)
        ndays = calendar.monthrange(ano, mes)[1]

        # Cabeçalho da tabela
        y0 = table_top - total_header_h
        cell(xs[0], y0, cols[0], total_header_h, '', lw=0.8)
        c.saveState(); c.translate(xs[0] + cols[0]/2, y0 + total_header_h/2); c.rotate(90)
        c.setFont('Helvetica-Bold', 7.2); c.drawCentredString(0, -2, 'DIAS'); c.restoreState()
        cell(xs[1], y0, cols[1], total_header_h, '', lw=0.8)
        for i, lab in [(2,'ENTRADA'),(3,'SAÍDA'),(4,'ENTRADA'),(5,'SAÍDA')]:
            cell(xs[i], y0, cols[i], total_header_h, lab, 7.2, True, 'center', lw=0.8)
        cell(xs[6], y0, cols[6], total_header_h, 'TOTAL HS\nNORMAIS', 6.3, True, 'center', lw=0.8)
        cell(xs[7], table_top-header_h1, cols[7]+cols[8], header_h1, 'EXTRAS', 7.2, True, 'center', lw=0.8)
        cell(xs[7], y0, cols[7], header_h2, 'ENTRADA', 6.9, True, 'center', lw=0.8)
        cell(xs[8], y0, cols[8], header_h2, 'SAÍDA', 6.9, True, 'center', lw=0.8)
        cell(xs[9], y0, cols[9], total_header_h, 'TOTAL HS\nEXTRAS', 6.3, True, 'center', lw=0.8)
        cell(xs[10], table_top-header_h1, cols[10], header_h1, 'ASSINATURA OU VISTO', 7.2, True, 'center', lw=0.8)
        cell(xs[10], y0, cols[10], header_h2, 'DO(A) EMPREGADO(A)', 7.2, True, 'center', lw=0.8)

        y = y0
        dias_sem = ['Seg','Ter','Qua','Qui','Sex','Sáb','Dom']
        for d in range(1, 32):
            y -= row_h
            if d <= ndays:
                sem = dias_sem[date(ano, mes, d).weekday()]
                data_text = f'{d:02d}'  # mostra o dia do mês com dois dígitos
            else:
                sem = ''
                data_text = ''
            occ = ocorrencias.get(d) if d <= ndays else None
            feriado = feriados_mes.get(d) if d <= ndays else None
            jornada_status = status_dia_por_jornada(jornada_info, date(ano, mes, d)) if d <= ndays else ''
            for i, cw in enumerate(cols):
                value = data_text if i == 0 else (sem if i == 1 else '')
                bold = i == 1 and bool(sem)
                preenchido = False

                # Mantém as colunas originais da folha aprovada; apenas preenche a linha quando houver ocorrência.
                if occ and i >= 2:
                    if i == 2:
                        value = occ.get('tipo', 'OCORRÊNCIA')
                    elif i in (3,4,5,6,7,8,9):
                        value = '—'
                    elif i == 10:
                        value = 'DISPENSADO'
                    preenchido = True
                    bold = True

                # Aplica sábado/descanso semanal conforme a jornada antes do feriado.
                # Assim o DSR/compensado não é sobrescrito por feriado quando o dia já é descanso.
                elif jornada_status and i >= 2:
                    if i == 2:
                        value = abreviar_texto_pdf(jornada_status, 'linha')
                    elif i in (3,4,5,6,7,8,9):
                        value = '—'
                    elif i == 10:
                        value = 'DISPENSADO'
                    preenchido = True
                    bold = True

                # Feriados cadastrados no sistema. Só aparece se não houver descanso/compensação da jornada.
                elif feriado and i >= 2:
                    if i == 2:
                        value = 'FERIADO'
                    elif i in (3,4,5,6,7,8,9):
                        value = '—'
                    elif i == 10:
                        value = 'DISPENSADO'
                    preenchido = True
                    bold = True

                if preenchido:
                    c.setFillColor(colors.HexColor('#eeeeee'))
                    c.rect(xs[i], y, cw, row_h, fill=1, stroke=0)
                    c.setFillColor(colors.black)
                cell(xs[i], y, cw, row_h, value, 6.2 if jornada_status else (6.6 if (occ or feriado) else 7.0), bold, 'center', lw=0.55)
            # Linhas divisórias mais fortes, como no modelo impresso
            if d % 7 == 0:
                c.setLineWidth(0.95)
                c.line(left, y, left + width, y)
        y -= total_h
        cell(left, y, sum(cols[:6]), total_h, 'TOTAIS', 7.2, True, 'center', lw=0.8)
        x = left + sum(cols[:6])
        for cw in cols[6:]:
            cell(x, y, cw, total_h, '', lw=0.8); x += cw

        # Resumo geral
        resumo_top = y - 8
        visto_w = 132
        resumo_w = width - visto_w
        header_h = 15
        rr = 9.6
        cell(left, resumo_top-header_h, resumo_w, header_h, 'RESUMO GERAL', 8.5, True, 'center', lw=0.8)
        cell(left+resumo_w, resumo_top-header_h, visto_w, header_h, 'VISTO DA FISCALIZAÇÃO', 8.5, True, 'center', lw=0.8)

        resumo = [
            ('+', '', '', 'Dias ou Horas Normais'),
            ('+', '', '', 'Horas Extras a 50%'),
            ('+', '', '', 'Horas Extras a 100%'),
            ('+', '', '', 'Adicionais (Discriminar no Verso)'),
            ('+', '', '', 'Outros Proventos (Insalubridade)'),
            ('=', '', '', 'Sub Total / Base de Cálculo'),
            ('-', '', '%', 'INSS'),
            ('-', '', '', 'Dependentes do Imposto de Renda'),
            ('-', '', '%', 'IRRF'),
            ('-', '', '', 'Outros Descontos (Discriminar no Verso)'),
            ('+', '', '', 'Salário Família'),
            ('Total Líquido a Receber', '', '', ''),
        ]
        y = resumo_top - header_h
        for idx, (a,b,perc,desc) in enumerate(resumo):
            y -= rr
            if idx == len(resumo)-1:
                cell(left, y, 281, rr, 'Total Líquido a Receber', 7.5, True, 'left', lw=0.8)
                cell(left+281, y, 28, rr, 'R$', 7.5, True, 'center', lw=0.8)
                cell(left+309, y, resumo_w-309, rr, '', lw=0.8)
            else:
                x = left
                cell(x, y, 19, rr, a, 7.2, False, 'center'); x += 19
                cell(x, y, 24, rr, b, 7.2, False, 'center'); x += 24
                cell(x, y, 35, rr, perc, 7.2, False, 'center'); x += 35
                cell(x, y, 203, rr, desc, 7.2, False, 'left'); x += 203
                cell(x, y, 28, rr, 'R$', 7.2, False, 'center'); x += 28
                cell(x, y, resumo_w-(x-left), rr, '', lw=0.6)
            # Caixa da fiscalização fica em branco, sem observações nos fins de semana
            if idx == 0:
                rect(left+resumo_w, y - rr*(len(resumo)-1), visto_w, rr*len(resumo), 0.8)

        # Assinatura do funcionário: linha acima e nome abaixo, com espaço real para assinar
        # Mantida no rodapé para não interferir no quadro-resumo.
        sig_y = 9
        c.setLineWidth(0.75)
        c.line(W/2-135, sig_y+18, W/2+135, sig_y+18)
        c.setFont('Helvetica-Bold', 8.2)
        c.drawCentredString(W/2, sig_y+5, f.get('nome',''))
        c.setFont('Helvetica', 6.8)
        c.drawCentredString(W/2, sig_y-4, 'ASSINATURA DO(A) EMPREGADO(A)')
        c.showPage()
    c.save()
    return destino



class App(tk.Tk):
    def __init__(self, usuario='admin', perfil='Administrador'):
        super().__init__()
        self.usuario = usuario
        self.perfil = perfil
        self.title(APP_NAME + ' - Enterprise v1.6.3 Interface')
        self.geometry('1180x740')
        self.minsize(1040,680)
        self.configure(bg='#eef2f6')
        self.selected_id = None
        self.last_pdf = None
        self.build_ui()
        self.load_all()
        self.protocol('WM_DELETE_WINDOW', self.on_close)

    def build_ui(self):
        style=ttk.Style(self)
        try:
            style.theme_use('clam')
        except Exception:
            pass
        style.configure('TButton', padding=8, font=('Arial', 10))
        style.configure('Treeview', rowheight=28, font=('Arial', 10))
        style.configure('Treeview.Heading', font=('Arial', 10, 'bold'))
        style.configure('Title.TLabel', font=('Arial', 18, 'bold'))
        style.configure('Side.TButton', padding=(14, 10), font=('Arial', 10, 'bold'))
        # Oculta as abas superiores: a navegação principal passa a ser por grupos no menu lateral.
        try:
            style.layout('Hidden.TNotebook.Tab', [])
        except Exception:
            pass

        self.menu = tk.Menu(self)
        arquivo = tk.Menu(self.menu, tearoff=0)
        arquivo.add_command(label='Abrir pasta do programa', command=self.open_base)
        arquivo.add_command(label='Abrir pasta PDFs', command=self.open_pdfs)
        arquivo.add_command(label='Fazer backup agora', command=self.backup_now)
        arquivo.add_separator()
        arquivo.add_command(label='Sair', command=self.destroy)
        self.menu.add_cascade(label='Arquivo', menu=arquivo)

        ferramentas = tk.Menu(self.menu, tearoff=0)
        ferramentas.add_command(label='Modo Desenvolvedor', command=lambda: self.nb.select(self.tab_dev))
        ferramentas.add_command(label='Abrir pasta do projeto', command=self.open_base)
        ferramentas.add_command(label='Abrir pasta de logs', command=lambda: self._abrir_pasta_generica(os.path.join(BASE_DIR, 'logs')))
        self.menu.add_cascade(label='Ferramentas', menu=ferramentas)
        self.config(menu=self.menu)

        main=ttk.Frame(self)
        main.pack(fill='both', expand=True)

        sidebar=tk.Frame(main, bg='#111827', width=220)
        sidebar.pack(side='left', fill='y')
        sidebar.pack_propagate(False)

        logo_box=tk.Frame(sidebar, bg='#111827')
        logo_box.pack(fill='x', padx=16, pady=(18,12))
        try:
            if os.path.exists(LOGO_APP):
                self.logo_sidebar_img = tk.PhotoImage(file=LOGO_APP)
                tk.Label(logo_box, image=self.logo_sidebar_img, bg='#111827').pack(anchor='w', pady=(0,8))
        except Exception:
            pass
        tk.Label(logo_box, text='SISTEMA GESTÃO\nIZZANT', bg='#111827', fg='white',
                 font=('Arial',15,'bold'), justify='left').pack(anchor='w')
        tk.Label(logo_box, text='Enterprise v1.6.3', bg='#111827', fg='#9ca3af',
                 font=('Arial',9), justify='left').pack(anchor='w', pady=(4,0))

        self.nb=ttk.Notebook(main, style='Hidden.TNotebook')
        self.nb.pack(side='right', fill='both', expand=True, padx=10, pady=10)

        # Abas internas. A maioria é aberta pelos cartões de submenu, não pela barra superior.
        self.tab_inicio=ttk.Frame(self.nb); self.tab_grupos=ttk.Frame(self.nb)
        self.tab_empresa=ttk.Frame(self.nb); self.tab_func=ttk.Frame(self.nb); self.tab_setores=ttk.Frame(self.nb)
        self.tab_jornadas=ttk.Frame(self.nb); self.tab_escalas=ttk.Frame(self.nb); self.tab_feriados=ttk.Frame(self.nb)
        self.tab_ocorrencias=ttk.Frame(self.nb); self.tab_ferias=ttk.Frame(self.nb); self.tab_banco=ttk.Frame(self.nb)
        self.tab_documentos=ttk.Frame(self.nb); self.tab_epis=ttk.Frame(self.nb); self.tab_exames=ttk.Frame(self.nb)
        self.tab_agenda=ttk.Frame(self.nb); self.tab_central_pdfs=ttk.Frame(self.nb); self.tab_assistente=ttk.Frame(self.nb)
        self.tab_pdf=ttk.Frame(self.nb); self.tab_backup=ttk.Frame(self.nb); self.tab_rel=ttk.Frame(self.nb)
        self.tab_logs=ttk.Frame(self.nb); self.tab_usuarios=ttk.Frame(self.nb); self.tab_atualizador=ttk.Frame(self.nb); self.tab_dev=ttk.Frame(self.nb)

        for frame, titulo in [
            (self.tab_inicio,'Painel'), (self.tab_grupos,'Módulos'), (self.tab_empresa,'Empresa'),
            (self.tab_func,'Funcionários'), (self.tab_setores,'Setores'), (self.tab_jornadas,'Jornadas'),
            (self.tab_escalas,'Escalas'), (self.tab_feriados,'Feriados'), (self.tab_ocorrencias,'Ocorrências'),
            (self.tab_ferias,'Férias'), (self.tab_banco,'Banco de Horas'), (self.tab_documentos,'Documentos'),
            (self.tab_epis,'EPIs'), (self.tab_exames,'Exames'), (self.tab_agenda,'Agenda'),
            (self.tab_central_pdfs,'Central PDFs'), (self.tab_assistente,'Assistente'), (self.tab_pdf,'Folha de Ponto'),
            (self.tab_backup,'Backup'), (self.tab_rel,'Relatórios'), (self.tab_logs,'Auditoria'),
            (self.tab_usuarios,'Usuários'), (self.tab_atualizador,'Atualizador'), (self.tab_dev,'Modo Desenvolvedor')
        ]:
            self.nb.add(frame, text=titulo)

        # Mapa de grupos: clique no grupo lateral para exibir os submódulos na área principal.
        self.grupos_menu = {
            'Painel': [('Abrir painel inicial', 'Resumo geral do sistema', self.tab_inicio, '🏠')],
            'Cadastros': [
                ('Empresa', 'Dados da empresa e identificação', self.tab_empresa, '🏢'),
                ('Funcionários', 'Cadastro, importação, status e consulta', self.tab_func, '👥'),
                ('Setores', 'Criação e manutenção dos setores', self.tab_setores, '📂'),
                ('Jornadas', 'Horários, sábado compensado, DSR e 12x36', self.tab_jornadas, '🕒'),
                ('Escalas', 'Escalas de trabalho e revezamento', self.tab_escalas, '📅'),
            ],
            'Gestão de Pessoas': [
                ('Férias', 'Programação, cálculos, documentos e histórico', self.tab_ferias, '🏖️'),
                ('Ocorrências', 'Férias, atestados, faltas, licenças e afastamentos', self.tab_ocorrencias, '📌'),
                ('Banco de Horas', 'Créditos, débitos e saldos', self.tab_banco, '⏱️'),
                ('EPIs', 'Controle de entrega, CA e validade', self.tab_epis, '🦺'),
                ('Exames', 'ASO, periódico, retorno e demissional', self.tab_exames, '🩺'),
                ('Agenda RH', 'Compromissos, alertas e lembretes', self.tab_agenda, '🗓️'),
            ],
            'Controle de Jornada': [
                ('Gerar Folhas de Ponto', 'PDF por funcionário, setor ou todos', self.tab_pdf, '🧾'),
                ('Feriados', 'Consulta, importação e manutenção dos feriados', self.tab_feriados, '📆'),
                ('Jornadas', 'Cadastro e regras de jornada', self.tab_jornadas, '🕒'),
                ('Escalas', 'Escalas e revezamentos', self.tab_escalas, '📅'),
            ],
            'Documentos': [
                ('Documentos Inteligentes', 'Modelos Word, campos automáticos e geração', self.tab_documentos, '📄'),
                ('Central de PDFs', 'Histórico e arquivos emitidos', self.tab_central_pdfs, '🗂️'),
                ('Assistente', 'Consultas rápidas e comandos internos', self.tab_assistente, '🤖'),
            ],
            'Relatórios': [
                ('Relatórios', 'Relatórios gerenciais e exportações', self.tab_rel, '📊'),
                ('Auditoria', 'Logs, ações e rastreabilidade', self.tab_logs, '🔎'),
            ],
            'Administração': [
                ('Usuários', 'Login, perfis e permissões', self.tab_usuarios, '🔐'),
                ('Backup', 'Cópias de segurança e restauração', self.tab_backup, '💾'),
                ('Atualizador', 'Controle de versão e atualizações futuras', self.tab_atualizador, '⬆️'),
            ],
            'Ferramentas': [
                ('Modo Desenvolvedor', 'Diagnóstico técnico e testes do sistema', self.tab_dev, '🛠️'),
                ('Abrir pasta do projeto', 'Acessar diretório local do sistema', None, '📁', self.open_base),
                ('Abrir pasta PDFs', 'Acessar folhas e documentos gerados', None, '📎', self.open_pdfs),
            ]
        }

        sidebar_buttons=[
            ('🏠 Painel', lambda: self.nb.select(self.tab_inicio)),
            ('📁 Cadastros', lambda: self.show_group('Cadastros')),
            ('👥 Gestão de Pessoas', lambda: self.show_group('Gestão de Pessoas')),
            ('⏰ Controle de Jornada', lambda: self.show_group('Controle de Jornada')),
            ('📄 Documentos', lambda: self.show_group('Documentos')),
            ('📊 Relatórios', lambda: self.show_group('Relatórios')),
            ('⚙ Administração', lambda: self.show_group('Administração')),
            ('🛠 Ferramentas', lambda: self.show_group('Ferramentas')),
        ]
        for label, cmd in sidebar_buttons:
            tk.Button(sidebar, text=label, anchor='w', bg='#1f2937', fg='white', activebackground='#374151',
                      activeforeground='white', bd=0, padx=16, pady=12, font=('Arial',10,'bold'), command=cmd).pack(fill='x', padx=10, pady=3)

        tk.Label(sidebar, text='Folha aprovada preservada\nJornadas/DSR estáveis', bg='#111827', fg='#d1d5db',
                 font=('Arial',9), justify='left').pack(side='bottom', anchor='w', padx=18, pady=18)

        self.build_inicio(); self.build_grupo_placeholder(); self.build_empresa(); self.build_func(); self.build_setores(); self.build_jornadas(); self.build_escalas(); self.build_feriados(); self.build_ocorrencias(); self.build_ferias(); self.build_banco_horas(); self.build_documentos(); self.build_epis(); self.build_exames(); self.build_agenda(); self.build_central_pdfs(); self.build_assistente(); self.build_pdf(); self.build_backup(); self.build_relatorios(); self.build_logs(); self.build_usuarios(); self.build_atualizador(); self.build_dev()
        self.nb.bind('<<NotebookTabChanged>>', self.on_tab_changed)

    def build_grupo_placeholder(self):
        self.grupo_container = ttk.Frame(self.tab_grupos)
        self.grupo_container.pack(fill='both', expand=True)

    def _tema_grupo(self, nome_grupo):
        temas = {
            'Cadastros': ('#2563eb', '#eff6ff', '#1d4ed8'),
            'Gestão de Pessoas': ('#16a34a', '#f0fdf4', '#15803d'),
            'Controle de Jornada': ('#f97316', '#fff7ed', '#c2410c'),
            'Documentos': ('#7c3aed', '#f5f3ff', '#6d28d9'),
            'Relatórios': ('#ca8a04', '#fefce8', '#a16207'),
            'Administração': ('#dc2626', '#fef2f2', '#b91c1c'),
            'Ferramentas': ('#475569', '#f8fafc', '#334155'),
        }
        return temas.get(nome_grupo, ('#111827', '#f9fafb', '#111827'))

    def show_group(self, nome_grupo):
        for child in self.grupo_container.winfo_children():
            child.destroy()
        cor, fundo, cor_escura = self._tema_grupo(nome_grupo)

        # Fundo geral mais claro para dar aparência de sistema comercial.
        try:
            self.grupo_container.configure(style='TFrame')
        except Exception:
            pass

        header = tk.Frame(self.grupo_container, bg=fundo, highlightbackground='#dbe4ef', highlightthickness=1)
        header.pack(fill='x', padx=20, pady=(18, 10))
        tk.Frame(header, bg=cor, width=7, height=86).pack(side='left', fill='y', padx=(0,16), pady=14)
        header_text = tk.Frame(header, bg=fundo)
        header_text.pack(side='left', fill='both', expand=True, pady=14)
        tk.Label(header_text, text=nome_grupo, bg=fundo, fg='#0f172a', font=('Arial', 23, 'bold')).pack(anchor='w')
        tk.Label(header_text, text='Selecione um módulo. Os submenus aparecem como cartões executivos para reduzir botões e facilitar a navegação.',
                 bg=fundo, fg='#475569', font=('Arial', 10), wraplength=820, justify='left').pack(anchor='w', pady=(4,0))

        resumo = tk.Frame(header, bg=fundo)
        resumo.pack(side='right', padx=16, pady=14)
        try:
            qtd = len(self.grupos_menu.get(nome_grupo, []))
            with con() as db:
                ativos = db.execute('SELECT COUNT(*) FROM funcionarios WHERE ativo=1').fetchone()[0]
            itens_resumo = [('Módulos', qtd), ('Ativos', ativos)]
        except Exception:
            itens_resumo = [('Módulos', len(self.grupos_menu.get(nome_grupo, [])))]
        for rotulo, valor in itens_resumo:
            mini = tk.Frame(resumo, bg='white', highlightbackground='#e2e8f0', highlightthickness=1)
            mini.pack(side='left', padx=5)
            tk.Label(mini, text=str(valor), bg='white', fg=cor_escura, font=('Arial', 15, 'bold')).pack(padx=12, pady=(7,0))
            tk.Label(mini, text=rotulo, bg='white', fg='#64748b', font=('Arial', 8, 'bold')).pack(padx=12, pady=(0,7))

        grid_wrap = tk.Frame(self.grupo_container, bg='#f1f5f9')
        grid_wrap.pack(fill='both', expand=True, padx=14, pady=8)
        grid = tk.Frame(grid_wrap, bg='#f1f5f9')
        grid.pack(fill='both', expand=True, padx=4, pady=4)
        itens = self.grupos_menu.get(nome_grupo, [])
        for idx, item in enumerate(itens):
            titulo, desc, tab, icone = item[0], item[1], item[2], item[3]
            comando_extra = item[4] if len(item) > 4 else None
            self.make_submenu_card(grid, idx, titulo, desc, tab, icone, comando_extra, cor, cor_escura)
        for c in range(3):
            grid.columnconfigure(c, weight=1, uniform='cards')
        for r in range((len(itens)+2)//3):
            grid.rowconfigure(r, weight=1)
        self.nb.select(self.tab_grupos)

    def make_submenu_card(self, parent, idx, titulo, desc, tab, icone='•', comando_extra=None, cor='#2563eb', cor_escura='#1d4ed8'):
        r, ccol = divmod(idx, 3)
        # Moldura externa cria um efeito de sombra sutil sem depender de bibliotecas externas.
        shadow = tk.Frame(parent, bg='#cbd5e1')
        shadow.grid(row=r, column=ccol, sticky='nsew', padx=12, pady=12)
        shadow.grid_propagate(False)
        try:
            shadow.configure(width=330, height=178)
        except Exception:
            pass

        card = tk.Frame(shadow, bg='white', highlightbackground='#dbe4ef', highlightthickness=1, cursor='hand2')
        card.place(x=0, y=0, relwidth=0.985, relheight=0.97)

        accent = tk.Frame(card, bg=cor, width=7)
        accent.pack(side='left', fill='y')

        conteudo = tk.Frame(card, bg='white')
        conteudo.pack(side='left', fill='both', expand=True)

        top = tk.Frame(conteudo, bg='white')
        top.pack(fill='x', padx=16, pady=(15,6))
        icon_box = tk.Label(top, text=icone, bg=cor, fg='white', font=('Arial', 20), width=3, height=1)
        icon_box.pack(side='left')
        text_box = tk.Frame(top, bg='white')
        text_box.pack(side='left', fill='x', expand=True, padx=12)
        tk.Label(text_box, text=titulo, bg='white', fg='#0f172a', font=('Arial',14,'bold')).pack(anchor='w')
        tk.Label(text_box, text='Clique para acessar', bg='white', fg=cor_escura, font=('Arial',8,'bold')).pack(anchor='w', pady=(2,0))

        tk.Label(conteudo, text=desc, bg='white', fg='#475569', font=('Arial',9), wraplength=250, justify='left').pack(anchor='w', padx=17, pady=(4,10))

        footer = tk.Frame(conteudo, bg='white')
        footer.pack(fill='x', side='bottom', padx=16, pady=(0,13))
        linha = tk.Frame(footer, bg='#e2e8f0', height=1)
        linha.pack(fill='x', pady=(0,8))
        abrir_lbl = tk.Label(footer, text='Abrir módulo  →', bg='white', fg=cor_escura, font=('Arial',9,'bold'), cursor='hand2')
        abrir_lbl.pack(side='right')
        status_lbl = tk.Label(footer, text='Disponível', bg='white', fg='#64748b', font=('Arial',8))
        status_lbl.pack(side='left')

        def abrir(_event=None):
            if comando_extra:
                comando_extra()
            elif tab is not None:
                self.nb.select(tab)

        def hover_on(_event=None):
            try:
                shadow.configure(bg=cor)
                card.configure(highlightbackground=cor, highlightthickness=2)
                abrir_lbl.configure(fg='#0f172a')
                status_lbl.configure(text='Pronto para abrir', fg=cor_escura)
            except Exception:
                pass
        def hover_off(_event=None):
            try:
                shadow.configure(bg='#cbd5e1')
                card.configure(highlightbackground='#dbe4ef', highlightthickness=1)
                abrir_lbl.configure(fg=cor_escura)
                status_lbl.configure(text='Disponível', fg='#64748b')
            except Exception:
                pass

        for w in (shadow, card, conteudo, top, icon_box, text_box, abrir_lbl, footer, status_lbl):
            w.bind('<Button-1>', abrir)
            w.bind('<Double-Button-1>', abrir)
            w.bind('<Enter>', hover_on)
            w.bind('<Leave>', hover_off)

    def _abrir_pasta_generica(self, caminho):
        try:
            os.makedirs(caminho, exist_ok=True)
            webbrowser.open(caminho)
        except Exception as e:
            messagebox.showerror('Abrir pasta', str(e))

    def build_dev(self):
        contexto = {
            'APP_NAME': APP_NAME,
            'BASE_DIR': BASE_DIR,
            'DATA_DIR': DATA_DIR,
            'PDF_DIR': PDF_DIR,
            'BACKUP_DIR': BACKUP_DIR,
            'RELATORIO_DIR': RELATORIO_DIR,
            'MODELOS_DIR': MODELOS_DIR,
            'DOCS_GERADOS_DIR': DOCS_GERADOS_DIR,
            'ASSETS_DIR': ASSETS_DIR,
            'DB_PATH': DB_PATH,
            'LOGO_APP': LOGO_APP,
        }
        if build_modo_desenvolvedor:
            build_modo_desenvolvedor(self, self.tab_dev, contexto)
        else:
            ttk.Label(self.tab_dev, text='Modo Desenvolvedor indisponível.', style='Title.TLabel').pack(padx=20, pady=20)

    def on_tab_changed(self, event=None):
        try:
            if self.nb.select() == str(self.tab_feriados):
                self.populate_feriados_tree()
        except Exception:
            pass

    def card(self, parent, title, value, col):
        frame=tk.Frame(parent, bg='white', highlightbackground='#d1d5db', highlightthickness=1)
        frame.grid(row=1, column=col, sticky='nsew', padx=8, pady=8)
        tk.Label(frame, text=title, bg='white', fg='#4b5563', font=('Arial',10)).pack(anchor='w', padx=14, pady=(12,4))
        lab=tk.Label(frame, text=value, bg='white', fg='#111827', font=('Arial',20,'bold'))
        lab.pack(anchor='w', padx=14, pady=(0,12))
        return lab

    def build_inicio(self):
        f=self.tab_inicio
        header = ttk.Frame(f)
        header.grid(row=0,column=0,columnspan=4,sticky='ew',padx=20,pady=(18,8))
        try:
            if os.path.exists(LOGO_APP):
                self.logo_inicio_img = tk.PhotoImage(file=LOGO_APP)
                ttk.Label(header, image=self.logo_inicio_img).pack(side='left', padx=(0,16))
        except Exception:
            pass
        ttk.Label(header, text='Sistema Gestão Izzant', style='Title.TLabel').pack(side='left', anchor='center')
        self.card_total=self.card(f,'Funcionários ativos','0',0)
        self.card_mes=self.card(f,'Mês padrão',MESES[datetime.now().month-1],1)
        self.card_pdf=self.card(f,'Último PDF','Nenhum',2)
        self.card_backup=self.card(f,'Backups','0',3)
        self.card_setores=self.card(f,'Setores','0',0)
        self.card_jornadas=self.card(f,'Jornadas','0',1)
        self.card_ocorrencias=self.card(f,'Ocorrências mês','0',2)
        self.card_feriados=self.card(f,'Feriados','0',3)
        actions=ttk.LabelFrame(f, text='Atalhos rápidos')
        actions.grid(row=2,column=0,columnspan=4,sticky='ew',padx=20,pady=18)
        ttk.Button(actions, text='Cadastrar funcionário', command=lambda:(self.nb.select(self.tab_func), self.clear_func())).pack(side='left', padx=8, pady=12)
        ttk.Button(actions, text='Gerar PDF de todos', command=lambda:(self.nb.select(self.tab_pdf), self.gerar_pdf(True))).pack(side='left', padx=8, pady=12)
        ttk.Button(actions, text='Importar Excel', command=self.importar_excel).pack(side='left', padx=8, pady=12)
        ttk.Button(actions, text='Baixar Modelo', command=self.baixar_modelo_importacao).pack(side='left', padx=8, pady=12)
        ttk.Button(actions, text='Backup agora', command=self.backup_now).pack(side='left', padx=8, pady=12)
        info=('Sistema local com salvamento automático em banco SQLite.\n'
              'Os PDFs são salvos automaticamente por ano e mês. A folha de ponto aprovada foi preservada. Férias são calculadas em dias corridos, sem prorrogação por feriados.')
        ttk.Label(f, text=info, font=('Arial',11), justify='left').grid(row=3,column=0,columnspan=4,sticky='w',padx=20,pady=8)
        for c in range(4): f.columnconfigure(c, weight=1)

    def build_empresa(self):
        f=self.tab_empresa
        ttk.Label(f, text='Cadastro da Empresa', style='Title.TLabel').grid(row=0,column=0,columnspan=4,sticky='w',padx=16,pady=16)
        self.emp_vars={}
        campos=[('nome','Empresa'),('cnpj','CNPJ/CEI'),('endereco','Endereço'),('numero','Nº'),('bairro','Bairro'),('cidade','Cidade'),('uf','UF')]
        for i,(key,label) in enumerate(campos):
            r=1+i//2; c=i%2*2
            ttk.Label(f,text=label).grid(row=r,column=c,sticky='w',padx=16,pady=6)
            v=tk.StringVar(); self.emp_vars[key]=v
            ttk.Entry(f,textvariable=v,width=48).grid(row=r,column=c+1,sticky='ew',padx=8,pady=6)
        ttk.Button(f,text='Salvar empresa',command=self.save_empresa).grid(row=6,column=1,sticky='w',padx=8,pady=20)
        ttk.Button(f,text='Abrir pasta de dados',command=self.open_data).grid(row=6,column=2,sticky='w',padx=8,pady=20)
        for c in range(4): f.columnconfigure(c, weight=1 if c%2 else 0)

    def build_func(self):
        f=self.tab_func
        ttk.Label(f,text='Cadastro de Funcionários',style='Title.TLabel').grid(row=0,column=0,columnspan=7,sticky='w',padx=12,pady=12)
        self.fvars={}
        fields=[('nome','Nome'),('cpf','CPF'),('ctps','CTPS / C.I.'),('admissao','Admissão DD/MM/AAAA'),('funcao','Função'),('setor','Setor'),('cbo','CBO'),('salario','Salário'),('jornada','Jornada'),('horario_trabalho','Horário de trabalho'),('descanso','Descanso semanal'),('sabado','Sábado')]
        for i,(key,label) in enumerate(fields):
            r=1+i//2; c=(i%2)*3
            ttk.Label(f,text=label).grid(row=r,column=c,sticky='w',padx=12,pady=4)
            v=tk.StringVar(); self.fvars[key]=v
            if key == 'setor':
                self.combo_setor_func = ttk.Combobox(f, textvariable=v, width=35, state='readonly')
                self.combo_setor_func.grid(row=r,column=c+1,columnspan=2,sticky='ew',padx=6,pady=4)
            elif key == 'jornada':
                self.combo_jornada_func = ttk.Combobox(f, textvariable=v, width=35, state='readonly')
                self.combo_jornada_func.grid(row=r,column=c+1,columnspan=2,sticky='ew',padx=6,pady=4)
                self.combo_jornada_func.bind('<<ComboboxSelected>>', lambda e: self.aplicar_jornada_funcionario())
            elif key == 'horario_trabalho':
                self.combo_horario_func = ttk.Combobox(f, textvariable=v, width=35, state='normal')
                self.combo_horario_func.grid(row=r,column=c+1,columnspan=2,sticky='ew',padx=6,pady=4)
            else:
                ent = ttk.Entry(f,textvariable=v,width=35)
                ent.grid(row=r,column=c+1,columnspan=2,sticky='ew',padx=6,pady=4)
                if key == 'admissao':
                    ent.bind('<FocusOut>', lambda e, var=v: formatar_campo_data(var))
        row=7
        ttk.Button(f,text='Novo',command=self.clear_func).grid(row=row,column=0,padx=6,pady=10,sticky='ew')
        ttk.Button(f,text='Salvar / Atualizar',command=self.save_func).grid(row=row,column=1,padx=6,pady=10,sticky='ew')
        ttk.Button(f,text='Duplicar',command=self.duplicar_func).grid(row=row,column=2,padx=6,pady=10,sticky='ew')
        ttk.Button(f,text='Inativar',command=self.delete_func).grid(row=row,column=3,padx=6,pady=10,sticky='ew')
        ttk.Button(f,text='Importar Excel',command=self.importar_excel).grid(row=row,column=4,padx=6,pady=10,sticky='ew')
        ttk.Button(f,text='Baixar Modelo',command=self.baixar_modelo_importacao).grid(row=row,column=5,padx=6,pady=10,sticky='ew')
        ttk.Button(f,text='PDFs individuais',command=self.gerar_pdfs_individuais_todos).grid(row=row,column=6,padx=6,pady=10,sticky='ew')
        ttk.Button(f,text='Exportar Excel',command=self.exportar_excel).grid(row=row+1,column=5,padx=6,pady=2,sticky='ew')
        ttk.Button(f,text='Atualizar jornadas dos funcionários',command=self.atualizar_funcionarios_conforme_jornada).grid(row=row+1,column=6,padx=6,pady=2,sticky='ew')
        ttk.Label(f,text='Pesquisar').grid(row=8,column=0,sticky='w',padx=12)
        self.search_var=tk.StringVar(); self.search_var.trace_add('write', lambda *_: self.populate_tree())
        ttk.Entry(f,textvariable=self.search_var,width=40).grid(row=8,column=1,columnspan=3,sticky='ew',padx=6,pady=4)
        self.show_inativos=tk.BooleanVar(value=False)
        ttk.Checkbutton(f,text='Mostrar inativos', variable=self.show_inativos, command=self.populate_tree).grid(row=8,column=4,sticky='w')
        ttk.Button(f,text='Reativar selecionado',command=self.reativar_func).grid(row=8,column=5,sticky='ew',padx=6)
        self.tree=ttk.Treeview(f, columns=('nome','setor','funcao','admissao','salario','status'), show='headings')
        for col,txt_h,w in [('nome','Nome',300),('setor','Setor',120),('funcao','Função',120),('admissao','Admissão',95),('salario','Salário',85),('status','Status',75)]:
            self.tree.heading(col,text=txt_h); self.tree.column(col,width=w)
        self.tree.grid(row=9,column=0,columnspan=7,sticky='nsew',padx=12,pady=8)
        self.tree.bind('<<TreeviewSelect>>', self.on_select)
        f.rowconfigure(9,weight=1)
        for c in range(7): f.columnconfigure(c,weight=1)

    def build_setores(self):
        f=self.tab_setores
        ttk.Label(f,text='Cadastro de Setores',style='Title.TLabel').grid(row=0,column=0,columnspan=5,sticky='w',padx=16,pady=16)
        ttk.Label(f,text='Nome do setor').grid(row=1,column=0,sticky='w',padx=12,pady=6)
        self.setor_nome_var=tk.StringVar()
        ttk.Entry(f,textvariable=self.setor_nome_var,width=35).grid(row=1,column=1,sticky='ew',padx=8,pady=6)
        ttk.Label(f,text='Descrição').grid(row=2,column=0,sticky='w',padx=12,pady=6)
        self.setor_desc_var=tk.StringVar()
        ttk.Entry(f,textvariable=self.setor_desc_var,width=55).grid(row=2,column=1,columnspan=3,sticky='ew',padx=8,pady=6)
        ttk.Button(f,text='Novo',command=self.clear_setor).grid(row=3,column=0,sticky='ew',padx=8,pady=10)
        ttk.Button(f,text='Salvar setor',command=self.save_setor).grid(row=3,column=1,sticky='ew',padx=8,pady=10)
        ttk.Button(f,text='Inativar setor',command=self.inativar_setor).grid(row=3,column=2,sticky='ew',padx=8,pady=10)
        ttk.Button(f,text='Reativar setor',command=self.reativar_setor).grid(row=3,column=3,sticky='ew',padx=8,pady=10)
        ttk.Label(f,text='Setores inativos não aparecem no cadastro do funcionário nem na geração por setor. Funcionários já vinculados continuam salvos.',font=('Arial',10)).grid(row=4,column=0,columnspan=5,sticky='w',padx=12,pady=4)
        self.setores_tree=ttk.Treeview(f, columns=('nome','descricao','ativo','qtd'), show='headings')
        for col,txt_h,w in [('nome','Setor',180),('descricao','Descrição',350),('ativo','Ativo',80),('qtd','Funcionários',100)]:
            self.setores_tree.heading(col,text=txt_h); self.setores_tree.column(col,width=w)
        self.setores_tree.grid(row=5,column=0,columnspan=5,sticky='nsew',padx=12,pady=12)
        self.setores_tree.bind('<<TreeviewSelect>>', self.on_select_setor)
        f.rowconfigure(5,weight=1)
        for c in range(5): f.columnconfigure(c,weight=1)

    def clear_setor(self):
        self.setor_nome_var.set('')
        self.setor_desc_var.set('')
        if hasattr(self,'setores_tree'):
            self.setores_tree.selection_remove(self.setores_tree.selection())

    def populate_setores_tree(self):
        if not hasattr(self,'setores_tree'):
            return
        for i in self.setores_tree.get_children(): self.setores_tree.delete(i)
        with con() as db:
            rows=db.execute("""SELECT s.nome, COALESCE(s.descricao,''), s.ativo,
                               (SELECT COUNT(*) FROM funcionarios f WHERE UPPER(COALESCE(f.setor,'GERAL'))=UPPER(s.nome))
                               FROM setores s ORDER BY s.ativo DESC, s.nome""").fetchall()
        for nome, desc, ativo, qtd in rows:
            self.setores_tree.insert('', 'end', iid=nome, values=(nome, desc, 'Sim' if ativo else 'Não', qtd))

    def on_select_setor(self, _=None):
        sel=self.setores_tree.selection()
        if not sel: return
        nome=sel[0]
        with con() as db:
            row=db.execute('SELECT nome, COALESCE(descricao,'') FROM setores WHERE nome=?',(nome,)).fetchone()
        if row:
            self.setor_nome_var.set(row[0]); self.setor_desc_var.set(row[1])

    def save_setor(self):
        nome=self.setor_nome_var.get().strip().upper()
        desc=self.setor_desc_var.get().strip()
        if not nome:
            messagebox.showwarning('Atenção','Informe o nome do setor.'); return
        with con() as db:
            db.execute('INSERT INTO setores(nome,descricao,ativo) VALUES(?,?,1) ON CONFLICT(nome) DO UPDATE SET descricao=excluded.descricao, ativo=1', (nome,desc))
        log_action(self.usuario,'SETOR','Setor salvo/atualizado: '+nome)
        self.clear_setor(); self.populate_setores_tree(); self.populate_setor_combo()
        messagebox.showinfo('Setor','Setor salvo e liberado para vincular aos funcionários.')

    def inativar_setor(self):
        nome=self.setor_nome_var.get().strip().upper()
        if not nome:
            messagebox.showwarning('Atenção','Selecione um setor.'); return
        if nome=='GERAL':
            messagebox.showwarning('Atenção','O setor GERAL não pode ser inativado.'); return
        if setor_em_uso(nome) and not messagebox.askyesno('Confirmar','Este setor possui funcionários vinculados. Deseja inativar mesmo assim?'):
            return
        with con() as db:
            db.execute('UPDATE setores SET ativo=0 WHERE nome=?',(nome,))
        log_action(self.usuario,'SETOR','Setor inativado: '+nome)
        self.populate_setores_tree(); self.populate_setor_combo()
        messagebox.showinfo('Setor','Setor inativado. Ele não aparecerá nas listas para novos cadastros.')

    def reativar_setor(self):
        nome=self.setor_nome_var.get().strip().upper()
        if not nome:
            messagebox.showwarning('Atenção','Selecione um setor.'); return
        with con() as db:
            db.execute('UPDATE setores SET ativo=1 WHERE nome=?',(nome,))
        log_action(self.usuario,'SETOR','Setor reativado: '+nome)
        self.populate_setores_tree(); self.populate_setor_combo()
        messagebox.showinfo('Setor','Setor reativado.')

    def build_jornadas(self):
        f=self.tab_jornadas
        ttk.Label(f,text='Cadastro de Jornadas',style='Title.TLabel').grid(row=0,column=0,columnspan=8,sticky='w',padx=16,pady=16)
        self.jvars={}
        campos=[('codigo','Código'),('nome','Nome da jornada'),('entrada1','Entrada 1'),('saida1','Saída 1'),('entrada2','Entrada 2'),('saida2','Saída 2'),('horas_dia','Horas diárias'),('horas_semana','Horas semanais'),('intervalo','Intervalo min.'),('tipo','Tipo'),('sabado_tratamento','Sábado'),('descanso_semanal','Descanso semanal')]
        for i,(key,label) in enumerate(campos):
            r=1+i//2; c=(i%2)*4
            ttk.Label(f,text=label).grid(row=r,column=c,sticky='w',padx=10,pady=4)
            v=tk.StringVar(); self.jvars[key]=v
            if key == 'tipo':
                ttk.Combobox(f,textvariable=v,values=['Comercial','Escala 12x36','Escala 6x1','Escala 5x2','Plantão','Personalizada'],state='readonly',width=28).grid(row=r,column=c+1,columnspan=3,sticky='ew',padx=6,pady=4)
            elif key == 'sabado_tratamento':
                ttk.Combobox(f,textvariable=v,values=['COMPENSADO','NÃO COMPENSADO','TRABALHADO','CONFORME ESCALA','NÃO SE APLICA'],state='readonly',width=28).grid(row=r,column=c+1,columnspan=3,sticky='ew',padx=6,pady=4)
            elif key == 'descanso_semanal':
                ttk.Combobox(f,textvariable=v,values=['DOMINGO','SÁBADO E DOMINGO','SEGUNDA','TERÇA','QUARTA','QUINTA','SEXTA','SÁBADO','ESCALA 12X36','CONFORME ESCALA','CONFORME ESCALA 12X36'],state='readonly',width=28).grid(row=r,column=c+1,columnspan=3,sticky='ew',padx=6,pady=4)
            else:
                ttk.Entry(f,textvariable=v,width=28).grid(row=r,column=c+1,columnspan=3,sticky='ew',padx=6,pady=4)
        ttk.Label(f,text='Dias trabalhados').grid(row=7,column=0,sticky='nw',padx=10,pady=6)
        self.j_dias={}
        dias=['Segunda','Terça','Quarta','Quinta','Sexta','Sábado','Domingo']
        frame_dias=ttk.Frame(f); frame_dias.grid(row=7,column=1,columnspan=7,sticky='w',padx=6,pady=4)
        for d in dias:
            var=tk.BooleanVar(value=d not in ['Sábado','Domingo']); self.j_dias[d]=var
            ttk.Checkbutton(frame_dias,text=d,variable=var).pack(side='left',padx=6)
        ttk.Label(f,text='Observações').grid(row=8,column=0,sticky='nw',padx=10,pady=6)
        self.j_obs=tk.Text(f,height=4,width=80,wrap='word')
        self.j_obs.grid(row=8,column=1,columnspan=7,sticky='ew',padx=6,pady=6)
        row=9
        ttk.Button(f,text='Nova',command=self.clear_jornada).grid(row=row,column=0,sticky='ew',padx=6,pady=10)
        ttk.Button(f,text='Salvar jornada',command=self.save_jornada).grid(row=row,column=1,sticky='ew',padx=6,pady=10)
        ttk.Button(f,text='Duplicar jornada',command=self.duplicar_jornada).grid(row=row,column=2,sticky='ew',padx=6,pady=10)
        ttk.Button(f,text='Inativar jornada',command=self.inativar_jornada).grid(row=row,column=3,sticky='ew',padx=6,pady=10)
        ttk.Button(f,text='Reativar jornada',command=self.reativar_jornada).grid(row=row,column=4,sticky='ew',padx=6,pady=10)
        ttk.Label(f,text='A jornada cadastrada será selecionada no funcionário e aparecerá na folha de ponto aprovada, sem alterar o layout.',font=('Arial',10)).grid(row=10,column=0,columnspan=8,sticky='w',padx=12,pady=4)
        self.j_tree=ttk.Treeview(f, columns=('codigo','nome','horario','horas','dias','sabado','descanso','tipo','ativo'), show='headings')
        for col,txt_h,w in [('codigo','Código',75),('nome','Jornada',165),('horario','Horário',190),('horas','Carga',95),('dias','Dias',190),('sabado','Sábado',120),('descanso','DSR',120),('tipo','Tipo',100),('ativo','Ativa',60)]:
            self.j_tree.heading(col,text=txt_h); self.j_tree.column(col,width=w)
        self.j_tree.grid(row=11,column=0,columnspan=8,sticky='nsew',padx=12,pady=10)
        self.j_tree.bind('<<TreeviewSelect>>', self.on_select_jornada)
        f.rowconfigure(11,weight=1)
        for c in range(8): f.columnconfigure(c,weight=1)

    def clear_jornada(self):
        for v in getattr(self,'jvars',{}).values(): v.set('')
        if hasattr(self,'j_obs'): self.j_obs.delete('1.0','end')
        for d,var in getattr(self,'j_dias',{}).items(): var.set(d not in ['Sábado','Domingo'])
        if hasattr(self,'jvars'):
            self.jvars.get('sabado_tratamento', tk.StringVar()).set('COMPENSADO')
            self.jvars.get('descanso_semanal', tk.StringVar()).set('DOMINGO')
        if hasattr(self,'j_tree'): self.j_tree.selection_remove(self.j_tree.selection())

    def populate_jornadas_tree(self):
        if not hasattr(self,'j_tree'): return
        for i in self.j_tree.get_children(): self.j_tree.delete(i)
        for j in get_jornadas(False):
            j = normalizar_jornada_12x36(j)
            horario = f"{j.get('entrada1') or ''} - {j.get('saida1') or ''}"
            if j.get('entrada2') or j.get('saida2'):
                horario += f" / {j.get('entrada2') or ''} - {j.get('saida2') or ''}"
            carga = f"{j.get('horas_dia') or ''} / {j.get('horas_semana') or ''}"
            self.j_tree.insert('', 'end', iid=j['nome'], values=(j.get('codigo') or '',j['nome'],horario,carga,j.get('dias_semana') or '',j.get('sabado_tratamento') or 'COMPENSADO',j.get('descanso_semanal') or 'DOMINGO',j.get('tipo') or '','Sim' if j.get('ativo') else 'Não'))

    def on_select_jornada(self, _=None):
        sel=self.j_tree.selection()
        if not sel: return
        nome=sel[0]
        j=next((x for x in get_jornadas(False) if x['nome']==nome), None)
        j=normalizar_jornada_12x36(j)
        if not j: return
        for key,var in self.jvars.items(): var.set(str(j.get(key) or ''))
        dias_set={x.strip() for x in (j.get('dias_semana') or '').split(',') if x.strip()}
        for d,var in self.j_dias.items(): var.set(d in dias_set)
        self.j_obs.delete('1.0','end'); self.j_obs.insert('1.0', j.get('observacao') or '')

    def save_jornada(self):
        vals={k:v.get().strip() for k,v in self.jvars.items()}
        nome=(vals.get('nome') or '').upper()
        if not nome:
            messagebox.showwarning('Atenção','Informe o nome da jornada.'); return
        codigo=(vals.get('codigo') or '').upper()
        dias=','.join([d for d,var in self.j_dias.items() if var.get()])
        obs=self.j_obs.get('1.0','end').strip()
        tipo_norm=(vals.get('tipo') or '').upper().replace(' ','')
        nome_norm=nome.upper().replace(' ','')
        if '12X36' in tipo_norm or '12X36' in nome_norm:
            vals['sabado_tratamento']='CONFORME ESCALA'
            vals['descanso_semanal']='CONFORME ESCALA 12X36'
            dias='CONFORME ESCALA 12X36'
            if not obs:
                obs='Sábados, domingos e descanso semanal seguem a escala 12x36.'
        with con() as db:
            db.execute("""INSERT INTO jornadas(codigo,nome,entrada1,saida1,entrada2,saida2,horas_dia,horas_semana,intervalo,dias_semana,tipo,sabado_tratamento,descanso_semanal,observacao,ativo)
                          VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,1)
                          ON CONFLICT(nome) DO UPDATE SET codigo=excluded.codigo, entrada1=excluded.entrada1, saida1=excluded.saida1,
                          entrada2=excluded.entrada2, saida2=excluded.saida2, horas_dia=excluded.horas_dia, horas_semana=excluded.horas_semana,
                          intervalo=excluded.intervalo, dias_semana=excluded.dias_semana, tipo=excluded.tipo,
                          sabado_tratamento=excluded.sabado_tratamento, descanso_semanal=excluded.descanso_semanal,
                          observacao=excluded.observacao, ativo=1""",
                       (codigo,nome,vals.get('entrada1'),vals.get('saida1'),vals.get('entrada2'),vals.get('saida2'),vals.get('horas_dia'),vals.get('horas_semana'),vals.get('intervalo'),dias,vals.get('tipo'),(vals.get('sabado_tratamento') or 'COMPENSADO').upper(),(vals.get('descanso_semanal') or 'DOMINGO').upper(),obs))
        self.clear_jornada(); self.populate_jornadas_tree(); self.populate_jornada_combo()
        log_action(self.usuario,'JORNADA','Jornada salva/atualizada: '+nome)
        messagebox.showinfo('Jornada','Jornada salva e liberada para o cadastro de funcionários.')

    def duplicar_jornada(self):
        if not hasattr(self,'jvars') or not self.jvars.get('nome') or not self.jvars['nome'].get():
            messagebox.showwarning('Atenção','Selecione uma jornada para duplicar.'); return
        self.jvars['codigo'].set('')
        self.jvars['nome'].set(self.jvars['nome'].get() + ' - CÓPIA')
        messagebox.showinfo('Duplicada','Revise o nome/código e clique em Salvar jornada.')

    def inativar_jornada(self):
        nome=self.jvars.get('nome').get().strip().upper() if hasattr(self,'jvars') else ''
        if not nome:
            messagebox.showwarning('Atenção','Selecione uma jornada.'); return
        with con() as db: db.execute('UPDATE jornadas SET ativo=0 WHERE nome=?',(nome,))
        self.populate_jornadas_tree(); self.populate_jornada_combo()
        log_action(self.usuario,'JORNADA','Jornada inativada: '+nome)
        messagebox.showinfo('Jornada','Jornada inativada. Ela não aparecerá para novos cadastros.')

    def reativar_jornada(self):
        nome=self.jvars.get('nome').get().strip().upper() if hasattr(self,'jvars') else ''
        if not nome:
            messagebox.showwarning('Atenção','Selecione uma jornada.'); return
        with con() as db: db.execute('UPDATE jornadas SET ativo=1 WHERE nome=?',(nome,))
        self.populate_jornadas_tree(); self.populate_jornada_combo()
        log_action(self.usuario,'JORNADA','Jornada reativada: '+nome)
        messagebox.showinfo('Jornada','Jornada reativada.')


    def build_escalas(self):
        f=self.tab_escalas
        ttk.Label(f,text='Cadastro de Escalas',style='Title.TLabel').grid(row=0,column=0,columnspan=6,sticky='w',padx=16,pady=16)
        self.esc_id=None
        self.esc_codigo=tk.StringVar(); self.esc_nome=tk.StringVar(); self.esc_tipo=tk.StringVar(value='Personalizada')
        campos=[('Código',self.esc_codigo),('Nome da escala',self.esc_nome),('Tipo',self.esc_tipo)]
        for i,(lab,var) in enumerate(campos):
            ttk.Label(f,text=lab).grid(row=1+i,column=0,sticky='w',padx=12,pady=6)
            if lab=='Tipo':
                ttk.Combobox(f,textvariable=var,values=['5x2','6x1','12x36','Plantão','Revezamento','Personalizada'],state='readonly').grid(row=1+i,column=1,columnspan=2,sticky='ew',padx=8,pady=6)
            else:
                ttk.Entry(f,textvariable=var,width=38).grid(row=1+i,column=1,columnspan=2,sticky='ew',padx=8,pady=6)
        ttk.Label(f,text='Dias / regra da escala').grid(row=4,column=0,sticky='nw',padx=12,pady=6)
        self.esc_dias=tk.Text(f,height=3,width=70,wrap='word'); self.esc_dias.grid(row=4,column=1,columnspan=4,sticky='ew',padx=8,pady=6)
        ttk.Label(f,text='Descrição').grid(row=5,column=0,sticky='nw',padx=12,pady=6)
        self.esc_desc=tk.Text(f,height=4,width=70,wrap='word'); self.esc_desc.grid(row=5,column=1,columnspan=4,sticky='ew',padx=8,pady=6)
        ttk.Button(f,text='Nova',command=self.clear_escala).grid(row=6,column=0,sticky='ew',padx=8,pady=10)
        ttk.Button(f,text='Salvar escala',command=self.save_escala).grid(row=6,column=1,sticky='ew',padx=8,pady=10)
        ttk.Button(f,text='Inativar',command=lambda:self.set_escala_ativo(0)).grid(row=6,column=2,sticky='ew',padx=8,pady=10)
        ttk.Button(f,text='Reativar',command=lambda:self.set_escala_ativo(1)).grid(row=6,column=3,sticky='ew',padx=8,pady=10)
        self.esc_tree=ttk.Treeview(f,columns=('codigo','nome','tipo','dias','ativo'),show='headings')
        for col,txt,w in [('codigo','Código',80),('nome','Escala',200),('tipo','Tipo',120),('dias','Dias/Regra',360),('ativo','Ativa',70)]:
            self.esc_tree.heading(col,text=txt); self.esc_tree.column(col,width=w)
        self.esc_tree.grid(row=7,column=0,columnspan=6,sticky='nsew',padx=12,pady=12)
        self.esc_tree.bind('<<TreeviewSelect>>', self.on_select_escala)
        f.rowconfigure(7,weight=1)
        for c in range(6): f.columnconfigure(c,weight=1)

    def clear_escala(self):
        self.esc_id=None; self.esc_codigo.set(''); self.esc_nome.set(''); self.esc_tipo.set('Personalizada')
        self.esc_dias.delete('1.0','end'); self.esc_desc.delete('1.0','end')

    def populate_escalas_tree(self):
        if not hasattr(self,'esc_tree'): return
        for i in self.esc_tree.get_children(): self.esc_tree.delete(i)
        with con() as db:
            rows=db.execute('SELECT id,codigo,nome,tipo,dias,ativo FROM escalas ORDER BY ativo DESC,nome').fetchall()
        for id_,codigo,nome,tipo,dias,ativo in rows:
            self.esc_tree.insert('', 'end', iid=str(id_), values=(codigo or '',nome or '',tipo or '',dias or '','Sim' if ativo else 'Não'))

    def on_select_escala(self,_=None):
        sel=self.esc_tree.selection()
        if not sel: return
        self.esc_id=int(sel[0])
        with con() as db:
            row=db.execute('SELECT codigo,nome,tipo,dias,descricao FROM escalas WHERE id=?',(self.esc_id,)).fetchone()
        if row:
            self.esc_codigo.set(row[0] or ''); self.esc_nome.set(row[1] or ''); self.esc_tipo.set(row[2] or 'Personalizada')
            self.esc_dias.delete('1.0','end'); self.esc_dias.insert('1.0', row[3] or '')
            self.esc_desc.delete('1.0','end'); self.esc_desc.insert('1.0', row[4] or '')

    def save_escala(self):
        nome=self.esc_nome.get().strip().upper()
        if not nome:
            messagebox.showwarning('Atenção','Informe o nome da escala.'); return
        vals=(self.esc_codigo.get().strip().upper(), nome, self.esc_tipo.get().strip(), self.esc_desc.get('1.0','end').strip(), self.esc_dias.get('1.0','end').strip())
        with con() as db:
            if self.esc_id:
                db.execute('UPDATE escalas SET codigo=?,nome=?,tipo=?,descricao=?,dias=?,ativo=1 WHERE id=?', (*vals,self.esc_id))
            else:
                db.execute('INSERT INTO escalas(codigo,nome,tipo,descricao,dias,ativo) VALUES(?,?,?,?,?,1)', vals)
        log_action(self.usuario,'ESCALA','Escala salva: '+nome)
        self.clear_escala(); self.populate_escalas_tree(); self.refresh_dashboard(); messagebox.showinfo('Escalas','Escala salva.')

    def set_escala_ativo(self, ativo):
        if not self.esc_id:
            messagebox.showwarning('Atenção','Selecione uma escala.'); return
        with con() as db: db.execute('UPDATE escalas SET ativo=? WHERE id=?',(ativo,self.esc_id))
        log_action(self.usuario,'ESCALA',('Reativada' if ativo else 'Inativada'))
        self.populate_escalas_tree(); self.refresh_dashboard()

    def build_feriados(self):
        f=self.tab_feriados
        # Visualização ampliada do módulo Feriados
        style = ttk.Style()
        style.configure('Feriados.Treeview', font=('Arial', 11), rowheight=30)
        style.configure('Feriados.Treeview.Heading', font=('Arial', 11, 'bold'))
        ttk.Label(f,text='Cadastro e Consulta de Feriados',style='Title.TLabel').grid(row=0,column=0,columnspan=6,sticky='w',padx=16,pady=16)
        self.fer_id=None
        self.fer_nome=tk.StringVar(); self.fer_data=tk.StringVar(); self.fer_tipo=tk.StringVar(value='Municipal'); self.fer_setor=tk.StringVar(value='TODOS')
        campos=[('Nome',self.fer_nome),('Data DD/MM/AAAA',self.fer_data),('Tipo',self.fer_tipo),('Setor',self.fer_setor)]
        for i,(lab,var) in enumerate(campos):
            r=1+i//2; c=(i%2)*3
            ttk.Label(f,text=lab).grid(row=r,column=c,sticky='w',padx=12,pady=6)
            if lab=='Tipo':
                ttk.Combobox(f,textvariable=var,values=['Nacional','Estadual','Municipal','Interno'],state='readonly').grid(row=r,column=c+1,columnspan=2,sticky='ew',padx=8,pady=6)
            elif lab=='Setor':
                self.combo_fer_setor=ttk.Combobox(f,textvariable=var,values=['TODOS'],state='readonly'); self.combo_fer_setor.grid(row=r,column=c+1,columnspan=2,sticky='ew',padx=8,pady=6)
            else:
                ttk.Entry(f,textvariable=var,width=35).grid(row=r,column=c+1,columnspan=2,sticky='ew',padx=8,pady=6)
        ttk.Label(f,text='Observação').grid(row=3,column=0,sticky='nw',padx=12,pady=6)
        self.fer_obs=tk.Text(f,height=4,width=80,wrap='word'); self.fer_obs.grid(row=3,column=1,columnspan=5,sticky='ew',padx=8,pady=6)
        ttk.Button(f,text='Novo',command=self.clear_feriado).grid(row=4,column=0,sticky='ew',padx=8,pady=10)
        ttk.Button(f,text='Salvar feriado',command=self.save_feriado).grid(row=4,column=1,sticky='ew',padx=8,pady=10)
        ttk.Button(f,text='Inativar',command=lambda:self.set_feriado_ativo(0)).grid(row=4,column=2,sticky='ew',padx=8,pady=10)
        ttk.Button(f,text='Reativar',command=lambda:self.set_feriado_ativo(1)).grid(row=4,column=3,sticky='ew',padx=8,pady=10)
        ttk.Button(f,text='Importar Brasil + Itajaí-SC',command=self.importar_feriados_brasil_itajai).grid(row=4,column=4,sticky='ew',padx=8,pady=10)
        ttk.Button(f,text='Importar 2026-2030',command=self.importar_feriados_2026_2030).grid(row=4,column=5,sticky='ew',padx=8,pady=10)

        # Área de consulta dos feriados já cadastrados.
        # Esta área é apenas para visualização/filtro e não altera a regra de geração da folha.
        box=ttk.LabelFrame(f,text='Consulta ampliada de feriados cadastrados')
        box.grid(row=5,column=0,columnspan=6,sticky='ew',padx=12,pady=(4,8))
        self.fer_filtro_ano=tk.StringVar(value='Todos')
        self.fer_filtro_tipo=tk.StringVar(value='Todos')
        self.fer_total_var=tk.StringVar(value='Feriados cadastrados: 0')
        self.fer_info_var=tk.StringVar(value='Use os filtros e clique em Consultar. Os feriados nacionais e municipais de Itajaí-SC são carregados automaticamente.')
        ttk.Label(box,text='Ano').grid(row=0,column=0,sticky='w',padx=8,pady=6)
        self.combo_fer_ano=ttk.Combobox(box,textvariable=self.fer_filtro_ano,values=['Todos'],state='readonly',width=12)
        self.combo_fer_ano.grid(row=0,column=1,sticky='w',padx=8,pady=6)
        ttk.Label(box,text='Tipo').grid(row=0,column=2,sticky='w',padx=8,pady=6)
        self.combo_fer_tipo=ttk.Combobox(box,textvariable=self.fer_filtro_tipo,values=['Todos','Nacional','Estadual','Municipal','Interno'],state='readonly',width=16)
        self.combo_fer_tipo.grid(row=0,column=3,sticky='w',padx=8,pady=6)
        ttk.Button(box,text='Consultar',command=self.populate_feriados_tree).grid(row=0,column=4,sticky='ew',padx=8,pady=6)
        ttk.Button(box,text='Limpar filtros',command=self.limpar_filtros_feriados).grid(row=0,column=5,sticky='ew',padx=8,pady=6)
        ttk.Label(box,textvariable=self.fer_total_var,font=('Arial',12,'bold')).grid(row=0,column=6,sticky='e',padx=12,pady=6)
        ttk.Label(box,textvariable=self.fer_info_var,font=('Arial',10)).grid(row=1,column=0,columnspan=7,sticky='w',padx=8,pady=(0,8))
        box.columnconfigure(6,weight=1)

        tree_frame = ttk.Frame(f)
        tree_frame.grid(row=6,column=0,columnspan=6,sticky='nsew',padx=12,pady=12)
        self.fer_tree=ttk.Treeview(tree_frame,columns=('data','nome','tipo','setor','ativo'),show='headings',height=22,style='Feriados.Treeview')
        for col,txt,w in [('data','Data',130),('nome','Nome do feriado',470),('tipo','Tipo',150),('setor','Setor',190),('ativo','Ativo',90)]:
            self.fer_tree.heading(col,text=txt); self.fer_tree.column(col,width=w, minwidth=w, anchor='center' if col in ('data','ativo','tipo') else 'w')
        self.fer_tree.tag_configure('nacional', background='#eef5ff')
        self.fer_tree.tag_configure('municipal', background='#f1fff0')
        self.fer_tree.tag_configure('estadual', background='#fff8e8')
        self.fer_tree.tag_configure('interno', background='#f6f0ff')
        self.fer_tree.tag_configure('inativo', background='#eeeeee', foreground='#777777')
        fer_scroll = ttk.Scrollbar(tree_frame, orient='vertical', command=self.fer_tree.yview)
        self.fer_tree.configure(yscrollcommand=fer_scroll.set)
        self.fer_tree.grid(row=0,column=0,sticky='nsew')
        fer_scroll.grid(row=0,column=1,sticky='ns')
        tree_frame.rowconfigure(0,weight=1); tree_frame.columnconfigure(0,weight=1)
        self.fer_tree.bind('<<TreeviewSelect>>', self.on_select_feriado)

        # Consulta textual de apoio: garante que os feriados fiquem visíveis mesmo em computadores
        # onde o Treeview do Windows/Tk apresente problema de renderização.
        consulta_text_frame = ttk.LabelFrame(f, text='Lista detalhada para conferência')
        consulta_text_frame.grid(row=7,column=0,columnspan=6,sticky='ew',padx=12,pady=(0,10))
        self.fer_text = tk.Text(consulta_text_frame, height=10, wrap='none', font=('Consolas', 11))
        self.fer_text.pack(fill='both', expand=True, padx=8, pady=8)
        self.fer_text.configure(state='disabled')

        f.rowconfigure(6,weight=1)
        for c in range(6): f.columnconfigure(c,weight=1)
        # Carrega imediatamente a consulta ao abrir a tela.
        self.after(250, self.populate_feriados_tree)

    def clear_feriado(self):
        self.fer_id=None; self.fer_nome.set(''); self.fer_data.set(''); self.fer_tipo.set('Municipal'); self.fer_setor.set('TODOS')
        self.fer_obs.delete('1.0','end')

    def limpar_filtros_feriados(self):
        if hasattr(self,'fer_filtro_ano'):
            self.fer_filtro_ano.set('Todos')
        if hasattr(self,'fer_filtro_tipo'):
            self.fer_filtro_tipo.set('Todos')
        self.populate_feriados_tree()

    def populate_feriados_tree(self):
        if not hasattr(self,'fer_tree'):
            return
        # Garante a carga local Brasil + Itajaí-SC antes de consultar.
        # Isso evita tela vazia mesmo em bancos novos ou migrados de versões anteriores.
        try:
            normalizar_datas_feriados()
            garantir_feriados_padrao(range(2026, 2031), getattr(self, 'usuario', 'sistema'))
            normalizar_datas_feriados()
        except Exception:
            pass
        for i in self.fer_tree.get_children():
            self.fer_tree.delete(i)
        ano = self.fer_filtro_ano.get() if hasattr(self,'fer_filtro_ano') else 'Todos'
        tipo_filtro = self.fer_filtro_tipo.get() if hasattr(self,'fer_filtro_tipo') else 'Todos'
        sql='SELECT id,nome,data,tipo,setor,ativo FROM feriados WHERE 1=1'
        params=[]
        if ano and ano!='Todos':
            # Compatível com datas em ISO (AAAA-MM-DD) e legados em DD/MM/AAAA.
            sql += ' AND (substr(data,1,4)=? OR substr(data,7,4)=?)'
            params.extend([str(ano), str(ano)])
        if tipo_filtro and tipo_filtro!='Todos':
            sql += ' AND tipo=?'
            params.append(tipo_filtro)
        sql += " ORDER BY CASE WHEN substr(data,3,1)='/' THEN substr(data,7,4)||'-'||substr(data,4,2)||'-'||substr(data,1,2) ELSE data END ASC, nome"
        with con() as db:
            rows=db.execute(sql, params).fetchall()
            anos=[]
            for (dt,) in db.execute("SELECT DISTINCT data FROM feriados WHERE data IS NOT NULL AND data<>''").fetchall():
                txt=str(dt or '')
                ano_val = txt[:4] if '-' in txt else (txt[6:10] if len(txt)>=10 else '')
                if ano_val and ano_val not in anos:
                    anos.append(ano_val)
            anos=sorted(anos)
        linhas_texto = []
        for id_,nome,data,tipo,setor,ativo in rows:
            data_br = fmt_data(data)
            ativo_txt = 'Sim' if ativo else 'Não'
            tag = 'inativo' if not ativo else (tipo or '').lower()
            if tag not in ('nacional','municipal','estadual','interno','inativo'):
                tag = ''
            self.fer_tree.insert('', 'end', iid=str(id_), values=(data_br,nome,tipo,setor,ativo_txt), tags=(tag,))
            linhas_texto.append(f'{data_br:<12} | {tipo:<10} | {ativo_txt:<3} | {setor:<15} | {nome}')
        if hasattr(self,'fer_text'):
            self.fer_text.configure(state='normal')
            self.fer_text.delete('1.0','end')
            if linhas_texto:
                self.fer_text.insert('end', 'DATA         | TIPO       | AT. | SETOR           | FERIADO\n')
                self.fer_text.insert('end', '-'*120 + '\n')
                self.fer_text.insert('end', '\n'.join(linhas_texto))
            else:
                self.fer_text.insert('end', 'Nenhum feriado encontrado para os filtros selecionados.')
            self.fer_text.configure(state='disabled')
        if hasattr(self,'fer_total_var'):
            ativos=sum(1 for r in rows if r[5])
            self.fer_total_var.set(f'Total: {len(rows)} feriados | Ativos: {ativos}')
            if hasattr(self,'fer_info_var'):
                self.fer_info_var.set('Consulta atualizada. Clique em um feriado da tabela para carregar os dados no cadastro acima.')
        if hasattr(self,'combo_fer_ano'):
            vals=['Todos']+[str(a) for a in anos if a]
            atual=self.fer_filtro_ano.get() if hasattr(self,'fer_filtro_ano') else 'Todos'
            self.combo_fer_ano['values']=vals
            if atual not in vals:
                self.fer_filtro_ano.set('Todos')
        if hasattr(self,'combo_fer_setor'):
            vals=['TODOS']+get_setores(True)
            self.combo_fer_setor['values']=vals

    def on_select_feriado(self,_=None):
        sel=self.fer_tree.selection()
        if not sel: return
        self.fer_id=int(sel[0])
        with con() as db:
            row=db.execute('SELECT nome,data,tipo,setor,observacao FROM feriados WHERE id=?',(self.fer_id,)).fetchone()
        if row:
            self.fer_nome.set(row[0] or ''); self.fer_data.set(fmt_data(row[1] or '')); self.fer_tipo.set(row[2] or 'Municipal'); self.fer_setor.set(row[3] or 'TODOS')
            self.fer_obs.delete('1.0','end'); self.fer_obs.insert('1.0', row[4] or '')

    def save_feriado(self):
        nome=self.fer_nome.get().strip().upper(); data_txt=self.fer_data.get().strip()
        if not nome or not data_txt:
            messagebox.showwarning('Atenção','Informe nome e data do feriado.'); return
        try: iso=parse_data_br(data_txt).isoformat()
        except Exception:
            messagebox.showerror('Data inválida','Use DD/MM/AAAA.'); return
        vals=(nome,iso,self.fer_tipo.get().strip(), 'TODAS', self.fer_setor.get().strip() or 'TODOS', self.fer_obs.get('1.0','end').strip())
        with con() as db:
            if self.fer_id:
                db.execute('UPDATE feriados SET nome=?,data=?,tipo=?,empresa=?,setor=?,observacao=?,ativo=1 WHERE id=?', (*vals,self.fer_id))
            else:
                db.execute('INSERT INTO feriados(nome,data,tipo,empresa,setor,observacao,ativo) VALUES(?,?,?,?,?,?,1)', vals)
        log_action(self.usuario,'FERIADO','Feriado salvo: '+nome)
        self.clear_feriado(); self.populate_feriados_tree(); self.refresh_dashboard(); messagebox.showinfo('Feriados','Feriado salvo.')

    def set_feriado_ativo(self, ativo):
        if not self.fer_id:
            messagebox.showwarning('Atenção','Selecione um feriado.'); return
        with con() as db: db.execute('UPDATE feriados SET ativo=? WHERE id=?',(ativo,self.fer_id))
        log_action(self.usuario,'FERIADO',('Reativado' if ativo else 'Inativado'))
        self.populate_feriados_tree(); self.refresh_dashboard()

    def importar_feriados_brasil_itajai(self):
        ano = datetime.now().year
        # Se o usuário já digitou uma data no cadastro, usa o ano dela como referência.
        try:
            if self.fer_data.get().strip():
                ano = parse_data_br(self.fer_data.get().strip()).year
        except Exception:
            pass
        total, fonte = importar_feriados_brasil_itajai(ano, self.usuario)
        self.populate_feriados_tree(); self.refresh_dashboard()
        messagebox.showinfo('Feriados importados', f'Importação concluída para {ano}.\nFonte nacionais: {fonte}.\nNovos feriados adicionados: {total}.')

    def importar_feriados_2026_2030(self):
        total = 0; fontes = []
        for ano in range(2026, 2031):
            t, fonte = importar_feriados_brasil_itajai(ano, self.usuario)
            total += t
            if fonte not in fontes:
                fontes.append(fonte)
        self.populate_feriados_tree(); self.refresh_dashboard()
        messagebox.showinfo('Feriados importados', f'Importação concluída de 2026 a 2030.\nFontes nacionais: {", ".join(fontes)}.\nNovos feriados adicionados: {total}.')

    def build_ocorrencias(self):
        f=self.tab_ocorrencias
        ttk.Label(f,text='Férias / Afastamentos / Ocorrências',style='Title.TLabel').grid(row=0,column=0,columnspan=6,sticky='w',padx=16,pady=16)
        ttk.Label(f,text='Funcionário').grid(row=1,column=0,sticky='w',padx=12,pady=6)
        self.occ_func_var=tk.StringVar()
        self.combo_occ_func=ttk.Combobox(f,textvariable=self.occ_func_var,width=65,state='readonly')
        self.combo_occ_func.grid(row=1,column=1,columnspan=5,sticky='ew',padx=8,pady=6)
        ttk.Label(f,text='Tipo').grid(row=2,column=0,sticky='w',padx=12,pady=6)
        self.occ_tipo_var=tk.StringVar(value='FÉRIAS')
        self.combo_occ_tipo=ttk.Combobox(f,textvariable=self.occ_tipo_var,values=['FÉRIAS','ATESTADO','LICENÇA','FOLGA','SUSPENSÃO','FALTA JUSTIFICADA','FALTA INJUSTIFICADA','LICENÇA MATERNIDADE','LICENÇA PATERNIDADE','LICENÇA REMUNERADA','CURSO/TREINAMENTO','TRABALHO EXTERNO'],state='readonly',width=25)
        self.combo_occ_tipo.grid(row=2,column=1,sticky='ew',padx=8,pady=6)
        self.combo_occ_tipo.bind('<<ComboboxSelected>>', lambda e: self.atualizar_abono_padrao())
        ttk.Label(f,text='Início DD/MM/AAAA').grid(row=2,column=2,sticky='w',padx=12,pady=6)
        self.occ_ini_var=tk.StringVar()
        self.occ_ini_entry=ttk.Entry(f,textvariable=self.occ_ini_var,width=18)
        self.occ_ini_entry.grid(row=2,column=3,sticky='ew',padx=8,pady=6)
        self.occ_ini_entry.bind('<FocusOut>', lambda e: formatar_campo_data(self.occ_ini_var))
        ttk.Label(f,text='Fim DD/MM/AAAA').grid(row=2,column=4,sticky='w',padx=12,pady=6)
        self.occ_fim_var=tk.StringVar()
        self.occ_fim_entry=ttk.Entry(f,textvariable=self.occ_fim_var,width=18)
        self.occ_fim_entry.grid(row=2,column=5,sticky='ew',padx=8,pady=6)
        self.occ_fim_entry.bind('<FocusOut>', lambda e: formatar_campo_data(self.occ_fim_var))
        self.occ_abona_var=tk.BooleanVar(value=True)
        ttk.Checkbutton(f,text='Abona o dia',variable=self.occ_abona_var).grid(row=3,column=1,sticky='w',padx=8,pady=6)
        ttk.Label(f,text='Observação / justificativa').grid(row=4,column=0,sticky='nw',padx=12,pady=6)
        self.occ_obs_text=tk.Text(f,height=4,width=70,wrap='word')
        self.occ_obs_text.grid(row=4,column=1,columnspan=5,sticky='ew',padx=8,pady=6)
        ttk.Button(f,text='Lançar ocorrência',command=self.salvar_ocorrencia).grid(row=5,column=1,sticky='ew',padx=8,pady=10)
        ttk.Button(f,text='Cancelar ocorrência selecionada',command=self.cancelar_ocorrencia).grid(row=5,column=2,sticky='ew',padx=8,pady=10)
        ttk.Button(f,text='Atualizar lista',command=self.carregar_ocorrencias).grid(row=5,column=3,sticky='ew',padx=8,pady=10)
        ttk.Label(f,text='Digite as datas no padrão brasileiro DD/MM/AAAA. O campo Abona o dia permite separar atestado/férias/licenças de faltas não abonadas. Na folha aprovada, os dias lançados serão preenchidos automaticamente, sem alterar o layout.',font=('Arial',10)).grid(row=6,column=0,columnspan=6,sticky='w',padx=12,pady=6)
        self.occ_tree=ttk.Treeview(f, columns=('func','tipo','ini','fim','abona','obs','status'), show='headings')
        for col,txt_h,w in [('func','Funcionário',240),('tipo','Tipo',130),('ini','Início',90),('fim','Fim',90),('abona','Abona',70),('obs','Observação',260),('status','Status',75)]:
            self.occ_tree.heading(col,text=txt_h); self.occ_tree.column(col,width=w)
        self.occ_tree.grid(row=7,column=0,columnspan=6,sticky='nsew',padx=12,pady=10)
        f.rowconfigure(7,weight=1)
        for c in range(6): f.columnconfigure(c,weight=1)

    def atualizar_abono_padrao(self):
        tipo = (self.occ_tipo_var.get() or '').strip().upper()
        # Padrão sugerido: faltas injustificadas e suspensão não abonam; demais ocorrências abonam por padrão.
        if tipo in ('FALTA INJUSTIFICADA','SUSPENSÃO'):
            self.occ_abona_var.set(False)
        else:
            self.occ_abona_var.set(True)

    def populate_ocorrencias_combo(self):
        if not hasattr(self,'combo_occ_func'): return
        vals=[f"{f['id']} - {f['nome']}" for f in get_funcionarios(True)]
        self.combo_occ_func['values']=vals
        if vals and self.occ_func_var.get() not in vals:
            self.occ_func_var.set(vals[0])

    def carregar_ocorrencias(self):
        if not hasattr(self,'occ_tree'): return
        for i in self.occ_tree.get_children(): self.occ_tree.delete(i)
        with con() as db:
            rows=db.execute("""SELECT o.id, f.nome, o.tipo, o.data_inicio, o.data_fim, COALESCE(o.observacao,''), COALESCE(o.abona,1), o.ativo
                               FROM ocorrencias o JOIN funcionarios f ON f.id=o.funcionario_id
                               ORDER BY date(o.data_inicio) DESC, f.nome""").fetchall()
        for oid,nome,tipo,ini,fim,obs,abona,ativo in rows:
            self.occ_tree.insert('', 'end', iid=str(oid), values=(nome,tipo,fmt_data(ini),fmt_data(fim),'Sim' if int(abona or 0) else 'Não',obs,'Ativa' if ativo else 'Cancelada'))

    def salvar_ocorrencia(self):
        val=self.occ_func_var.get().strip()
        if not val:
            messagebox.showwarning('Atenção','Selecione o funcionário.'); return
        try:
            fid=int(val.split(' - ')[0])
            data_ini = parse_data_br(self.occ_ini_var.get())
            data_fim = parse_data_br(self.occ_fim_var.get())
        except Exception:
            messagebox.showwarning('Atenção','Informe as datas no formato DD/MM/AAAA. Exemplo: 01/07/2026'); return
        if data_fim < data_ini:
            messagebox.showwarning('Atenção','A data final não pode ser anterior à inicial.'); return
        self.occ_ini_var.set(data_ini.strftime('%d/%m/%Y'))
        self.occ_fim_var.set(data_fim.strftime('%d/%m/%Y'))
        with con() as db:
            obs = self.occ_obs_text.get('1.0','end').strip() if hasattr(self,'occ_obs_text') else ''
            db.execute("""INSERT INTO ocorrencias(funcionario_id,tipo,data_inicio,data_fim,observacao,abona,ativo)
                          VALUES(?,?,?,?,?,?,1)""", (fid,self.occ_tipo_var.get().strip().upper(),data_ini.isoformat(),data_fim.isoformat(),obs,1 if self.occ_abona_var.get() else 0))
        log_action(self.usuario,'OCORRÊNCIA',f"{self.occ_tipo_var.get()} lançada para {val}")
        self.carregar_ocorrencias()
        if hasattr(self,'occ_obs_text'):
            self.occ_obs_text.delete('1.0','end')
        messagebox.showinfo('Ocorrência','Ocorrência lançada. Ao gerar/visualizar o PDF, os dias serão preenchidos automaticamente.')

    def cancelar_ocorrencia(self):
        sel=self.occ_tree.selection() if hasattr(self,'occ_tree') else []
        if not sel:
            messagebox.showwarning('Atenção','Selecione uma ocorrência.'); return
        if not messagebox.askyesno('Confirmar','Cancelar esta ocorrência?'):
            return
        with con() as db:
            db.execute('UPDATE ocorrencias SET ativo=0 WHERE id=?',(int(sel[0]),))
        log_action(self.usuario,'OCORRÊNCIA','Ocorrência cancelada')
        self.carregar_ocorrencias()


    def build_ferias(self):
        f=self.tab_ferias
        ttk.Label(f,text='Controle Completo de Férias - Cálculos, Histórico e Documentos',style='Title.TLabel').grid(row=0,column=0,columnspan=8,sticky='w',padx=16,pady=16)
        ttk.Label(f,text='Funcionário').grid(row=1,column=0,sticky='w',padx=12,pady=6)
        self.ferias_func_var=tk.StringVar()
        self.combo_ferias_func=ttk.Combobox(f,textvariable=self.ferias_func_var,width=65,state='readonly')
        self.combo_ferias_func.grid(row=1,column=1,columnspan=5,sticky='ew',padx=8,pady=6)
        self.combo_ferias_func.bind('<<ComboboxSelected>>', lambda e: (self.atualizar_periodos_ferias_funcionario(), self.carregar_dados_financeiros_ferias()))

        ttk.Label(f,text='Período aquisitivo calculado pela admissão').grid(row=2,column=0,sticky='w',padx=12,pady=6)
        self.ferias_periodo_var=tk.StringVar()
        self.combo_ferias_periodo=ttk.Combobox(f,textvariable=self.ferias_periodo_var,width=65,state='readonly')
        self.combo_ferias_periodo.grid(row=2,column=1,columnspan=4,sticky='ew',padx=8,pady=6)
        self.combo_ferias_periodo.bind('<<ComboboxSelected>>', lambda e: self.aplicar_periodo_aquisitivo_ferias())
        ttk.Button(f,text='Atualizar períodos',command=self.atualizar_periodos_ferias_funcionario).grid(row=2,column=5,sticky='ew',padx=8,pady=6)
        ttk.Button(f,text='Calcular férias',command=self.calcular_ferias_tela).grid(row=2,column=6,sticky='ew',padx=8,pady=6)

        campos=[('fer_aq_ini','Aquisitivo início'),('fer_aq_fim','Aquisitivo fim'),('fer_conc','Limite concessivo'),('fer_ini','Início férias'),('fer_fim','Fim férias'),('fer_ret','Retorno')]
        self.ferias_vars={}
        for i,(key,label) in enumerate(campos):
            r=3+i//3; c=(i%3)*2
            ttk.Label(f,text=label+' DD/MM/AAAA').grid(row=r,column=c,sticky='w',padx=12,pady=6)
            v=tk.StringVar(); self.ferias_vars[key]=v
            e=ttk.Entry(f,textvariable=v,width=18); e.grid(row=r,column=c+1,sticky='ew',padx=8,pady=6)
            e.bind('<FocusOut>', lambda ev, var=v: (formatar_campo_data(var), self.calcular_ferias_tela()))

        ttk.Label(f,text='Dias a gozar').grid(row=5,column=0,sticky='w',padx=12,pady=6)
        self.fer_dias_gozar=tk.StringVar(value='30'); ttk.Entry(f,textvariable=self.fer_dias_gozar,width=10).grid(row=5,column=1,sticky='ew',padx=8,pady=6)
        ttk.Label(f,text='Dias abono').grid(row=5,column=2,sticky='w',padx=12,pady=6)
        self.fer_dias_abono=tk.StringVar(value='0'); ttk.Entry(f,textvariable=self.fer_dias_abono,width=10).grid(row=5,column=3,sticky='ew',padx=8,pady=6)
        self.fer_adianta_13=tk.IntVar(value=0); ttk.Checkbutton(f,text='Adiantamento 13º',variable=self.fer_adianta_13,command=self.calcular_ferias_tela).grid(row=5,column=4,sticky='w',padx=8,pady=6)
        ttk.Label(f,text='Status').grid(row=5,column=5,sticky='w',padx=12,pady=6)
        self.fer_status=tk.StringVar(value='Programada')
        ttk.Combobox(f,textvariable=self.fer_status,values=['Programada','Em gozo','Concluída','Cancelada'],state='readonly').grid(row=5,column=6,sticky='ew',padx=8,pady=6)

        ttk.Label(f,text='Salário-base').grid(row=6,column=0,sticky='w',padx=12,pady=6)
        self.fer_salario=tk.StringVar(value='0,00'); ttk.Entry(f,textvariable=self.fer_salario,width=14).grid(row=6,column=1,sticky='ew',padx=8,pady=6)
        ttk.Label(f,text='Médias/variáveis').grid(row=6,column=2,sticky='w',padx=12,pady=6)
        self.fer_media=tk.StringVar(value='0,00'); ttk.Entry(f,textvariable=self.fer_media,width=14).grid(row=6,column=3,sticky='ew',padx=8,pady=6)
        ttk.Label(f,text='Dias restantes').grid(row=6,column=4,sticky='w',padx=12,pady=6)
        self.fer_dias_restantes=tk.StringVar(value='0'); ttk.Entry(f,textvariable=self.fer_dias_restantes,width=10,state='readonly').grid(row=6,column=5,sticky='ew',padx=8,pady=6)

        self.fer_calc_vars={}
        calc_campos=[('valor_ferias','Valor férias'),('valor_um_terco','1/3 constitucional'),('valor_abono','Abono + 1/3'),('valor_13','Adiant. 13º'),('total_bruto','Total bruto'),('inss_estimado','INSS estimado'),('irrf_estimado','IRRF estimado'),('liquido_estimado','Líquido estimado')]
        for i,(key,label) in enumerate(calc_campos):
            r=7+i//4; c=(i%4)*2
            ttk.Label(f,text=label).grid(row=r,column=c,sticky='w',padx=12,pady=4)
            v=tk.StringVar(value='R$ 0,00'); self.fer_calc_vars[key]=v
            ttk.Entry(f,textvariable=v,width=16,state='readonly').grid(row=r,column=c+1,sticky='ew',padx=8,pady=4)

        ttk.Label(f,text='Observação').grid(row=9,column=0,sticky='nw',padx=12,pady=6)
        self.fer_obs=tk.Text(f,height=3,wrap='word'); self.fer_obs.grid(row=9,column=1,columnspan=7,sticky='ew',padx=8,pady=6)
        ttk.Button(f,text='Salvar férias',command=self.salvar_ferias).grid(row=10,column=1,sticky='ew',padx=8,pady=10)
        ttk.Button(f,text='Cancelar férias selecionada',command=self.cancelar_ferias).grid(row=10,column=2,sticky='ew',padx=8,pady=10)
        ttk.Button(f,text='Gerar ocorrência FÉRIAS na folha',command=self.gerar_ocorrencia_ferias).grid(row=10,column=3,sticky='ew',padx=8,pady=10)
        ttk.Button(f,text='Gerar Aviso de Férias',command=lambda:self.gerar_documento_ferias('Aviso de Férias')).grid(row=10,column=4,sticky='ew',padx=8,pady=10)
        ttk.Button(f,text='Gerar Recibo de Férias',command=lambda:self.gerar_documento_ferias('Recibo de Férias')).grid(row=10,column=5,sticky='ew',padx=8,pady=10)
        ttk.Button(f,text='Gerar Comunicação/Termo',command=lambda:self.gerar_documento_ferias('Comunicação de Férias')).grid(row=10,column=6,sticky='ew',padx=8,pady=10)
        ttk.Button(f,text='Concluir Férias',command=self.concluir_ferias).grid(row=10,column=7,sticky='ew',padx=8,pady=10)
        ttk.Button(f,text='Atualizar cálculo',command=self.calcular_ferias_tela).grid(row=12,column=6,sticky='ew',padx=8,pady=4)
        ttk.Button(f,text='Abrir documentos de férias',command=self.abrir_documentos_ferias).grid(row=12,column=7,sticky='ew',padx=8,pady=4)
        self.fer_tree=ttk.Treeview(f,columns=('func','aq','conc','periodo','ret','dias','rest','total','status'),show='headings')
        for col,txt,w in [('func','Funcionário',210),('aq','Aquisitivo',160),('conc','Concessivo até',100),('periodo','Férias',165),('ret','Retorno',85),('dias','Dias',50),('rest','Rest.',50),('total','Total bruto',95),('status','Status',90)]:
            self.fer_tree.heading(col,text=txt); self.fer_tree.column(col,width=w)
        self.fer_tree.grid(row=13,column=0,columnspan=8,sticky='nsew',padx=12,pady=8)
        f.rowconfigure(13,weight=1)
        for c in range(8): f.columnconfigure(c,weight=1)

    def _ferias_func_id(self):
        val = getattr(self, 'ferias_func_var', tk.StringVar()).get().strip()
        try:
            return int(val.split(' - ')[0])
        except Exception:
            return None

    def atualizar_periodos_ferias_funcionario(self):
        fid = self._ferias_func_id()
        if not fid or not hasattr(self, 'combo_ferias_periodo'):
            return
        with con() as db:
            row = db.execute('SELECT admissao FROM funcionarios WHERE id=?', (fid,)).fetchone()
            usados = db.execute("""SELECT aquisitivo_inicio, aquisitivo_fim, status FROM ferias_controle
                                WHERE funcionario_id=? AND ativo=1 AND COALESCE(status,'')<>'Cancelada'""", (fid,)).fetchall()
        admissao = row[0] if row else None
        usados_set = {(u[0] or '', u[1] or '') for u in usados}
        periodos = calcular_periodos_aquisitivos(admissao, date.today(), quantidade=8)
        self.ferias_periodos_map = {}
        valores=[]
        for p in periodos:
            status = 'Já lançado' if (p['inicio_iso'], p['fim_iso']) in usados_set else p['status']
            label = f"{p['periodo_br']} | Concessivo até {p['concessivo_br']} | {status}"
            self.ferias_periodos_map[label] = p
            valores.append(label)
        self.combo_ferias_periodo['values'] = valores
        if valores:
            # Prioriza o primeiro período disponível ainda não lançado.
            escolha = next((v for v in valores if v.endswith('| Disponível')), valores[0])
            self.ferias_periodo_var.set(escolha)
            self.aplicar_periodo_aquisitivo_ferias()

    def aplicar_periodo_aquisitivo_ferias(self):
        label = getattr(self, 'ferias_periodo_var', tk.StringVar()).get()
        p = getattr(self, 'ferias_periodos_map', {}).get(label)
        if not p or not hasattr(self, 'ferias_vars'):
            return
        self.ferias_vars['fer_aq_ini'].set(p['inicio_br'])
        self.ferias_vars['fer_aq_fim'].set(p['fim_br'])
        self.ferias_vars['fer_conc'].set(p['concessivo_br'])

    def atualizar_periodo_por_inicio_ferias(self):
        # Se o usuário informar o início das férias antes de escolher o período, recalcula o período correto.
        if not hasattr(self, 'ferias_vars'):
            return
        fid = self._ferias_func_id()
        if not fid:
            return
        inicio_gozo = self.ferias_vars.get('fer_ini').get().strip() if self.ferias_vars.get('fer_ini') else ''
        if not inicio_gozo:
            return
        with con() as db:
            row = db.execute('SELECT admissao FROM funcionarios WHERE id=?', (fid,)).fetchone()
        admissao = row[0] if row else None
        aq_ini, aq_fim, _ = calcular_periodo_aquisitivo(admissao, inicio_gozo)
        if aq_ini and aq_fim:
            self.ferias_vars['fer_aq_ini'].set(aq_ini)
            self.ferias_vars['fer_aq_fim'].set(aq_fim)
            try:
                conc = _add_years_safe(parse_data_br(aq_fim), 1)
                self.ferias_vars['fer_conc'].set(conc.strftime('%d/%m/%Y'))
            except Exception:
                pass

    def carregar_dados_financeiros_ferias(self):
        fid=self._ferias_func_id()
        if not fid: return
        try:
            with con() as db:
                row=db.execute('SELECT salario FROM funcionarios WHERE id=?',(fid,)).fetchone()
            sal = row[0] if row else 0
            if hasattr(self,'fer_salario'):
                self.fer_salario.set(moeda_br(sal).replace('R$','').strip())
            self.calcular_ferias_tela()
        except Exception:
            pass

    def calcular_ferias_tela(self):
        if not hasattr(self,'ferias_vars'):
            return
        try:
            dias=int(str(getattr(self,'fer_dias_gozar',tk.StringVar(value='0')).get() or '0').strip() or 0)
        except Exception:
            dias=0
        try:
            dias_abono=int(str(getattr(self,'fer_dias_abono',tk.StringVar(value='0')).get() or '0').strip() or 0)
        except Exception:
            dias_abono=0
        # Calcula fim e retorno automaticamente a partir do início + dias a gozar.
        try:
            ini_txt=self.ferias_vars.get('fer_ini').get().strip()
            if ini_txt and dias>0:
                ini=parse_data_br(ini_txt)
                fim, ret=adicionar_dias_corridos(ini,dias)
                if fim and ret:
                    self.ferias_vars['fer_fim'].set(fim.strftime('%d/%m/%Y'))
                    self.ferias_vars['fer_ret'].set(ret.strftime('%d/%m/%Y'))
                    self.atualizar_periodo_por_inicio_ferias()
        except Exception:
            pass
        try:
            usados=0
            fid=self._ferias_func_id()
            aqi=data_para_iso(self.ferias_vars['fer_aq_ini'].get()) if self.ferias_vars.get('fer_aq_ini') and self.ferias_vars['fer_aq_ini'].get().strip() else ''
            aqf=data_para_iso(self.ferias_vars['fer_aq_fim'].get()) if self.ferias_vars.get('fer_aq_fim') and self.ferias_vars['fer_aq_fim'].get().strip() else ''
            if fid and aqi and aqf:
                with con() as db:
                    r=db.execute("""SELECT COALESCE(SUM(dias),0) FROM ferias_controle
                                    WHERE funcionario_id=? AND aquisitivo_inicio=? AND aquisitivo_fim=?
                                    AND ativo=1 AND COALESCE(status,'')<>'Cancelada'""",(fid,aqi,aqf)).fetchone()
                    usados=int(r[0] or 0)
            restantes=max(0,30-usados-dias)
            if hasattr(self,'fer_dias_restantes'):
                self.fer_dias_restantes.set(str(restantes))
        except Exception:
            restantes=max(0,30-dias)
            if hasattr(self,'fer_dias_restantes'): self.fer_dias_restantes.set(str(restantes))
        vals=calcular_ferias_valores(getattr(self,'fer_salario',tk.StringVar(value='0')).get(), dias, dias_abono, getattr(self,'fer_media',tk.StringVar(value='0')).get(), bool(getattr(self,'fer_adianta_13',tk.IntVar(value=0)).get()))
        if hasattr(self,'fer_calc_vars'):
            for k,v in vals.items():
                if k in self.fer_calc_vars:
                    self.fer_calc_vars[k].set(moeda_br(v))
        return vals

    def _ferias_contexto_mais_recente(self, fid):
        with con() as db:
            row = db.execute("""SELECT aquisitivo_inicio,aquisitivo_fim,concessivo_fim,inicio,fim,retorno,dias,observacao,status,
                                      dias_abono,dias_restantes,salario_base,media_variaveis,valor_ferias,valor_um_terco,valor_abono,valor_13,total_bruto,inss_estimado,irrf_estimado,liquido_estimado
                                FROM ferias_controle
                                WHERE funcionario_id=? AND ativo=1 AND COALESCE(status,'')<>'Cancelada'
                                ORDER BY date(inicio) DESC, id DESC LIMIT 1""", (fid,)).fetchone()
        if not row:
            return {}
        aqi, aqf, conc, ini, fim, ret, dias, obs, status, dias_abono, dias_restantes, salario_base, media_variaveis, valor_ferias, valor_um_terco, valor_abono, valor_13, total_bruto, inss_estimado, irrf_estimado, liquido_estimado = row
        return {
            'PERIODO_AQUISITIVO': f'{fmt_data(aqi)} a {fmt_data(aqf)}',
            'PERIODO_AQUISITIVO_INICIO': fmt_data(aqi),
            'PERIODO_AQUISITIVO_FIM': fmt_data(aqf),
            'PERIODO_AQUISITIVO_COMPLETO': f'{fmt_data(aqi)} a {fmt_data(aqf)}',
            'PERIODO_CONCESSIVO': f'até {fmt_data(conc)}' if conc else '',
            'PERIODO_CONCESSIVO_FIM': fmt_data(conc),
            'DATA_INICIO': fmt_data(ini), 'INICIO': fmt_data(ini),
            'DATA_FIM': fmt_data(fim), 'TERMINO': fmt_data(fim),
            'DATA_RETORNO': fmt_data(ret),
            'DIAS_FERIAS': str(dias or ''), 'DIAS_GOZADOS': str(dias or ''),
            'DIAS_ABONO': str(dias_abono or 0), 'DIAS_RESTANTES': str(dias_restantes or 0),
            'PERIODO': f'{fmt_data(ini)} a {fmt_data(fim)}',
            'SALARIO_BASE': moeda_br(salario_base), 'MEDIA_VARIAVEIS': moeda_br(media_variaveis),
            'VALOR_FERIAS': moeda_br(valor_ferias), 'VALOR_UM_TERCO': moeda_br(valor_um_terco),
            'VALOR_ABONO': moeda_br(valor_abono), 'ADIANTAMENTO_13': 'SIM' if (valor_13 or 0)>0 else 'NÃO',
            'VALOR_13': moeda_br(valor_13), 'TOTAL_BRUTO': moeda_br(total_bruto),
            'INSS_ESTIMADO': moeda_br(inss_estimado), 'IRRF_ESTIMADO': moeda_br(irrf_estimado),
            'LIQUIDO_ESTIMADO': moeda_br(liquido_estimado),
            'OBS_FERIAS': obs or '', 'STATUS_FERIAS': status or ''
        }

    def _ferias_registro_selecionado(self):
        sel=self.fer_tree.selection() if hasattr(self,'fer_tree') else []
        if not sel:
            return None
        with con() as db:
            row=db.execute('SELECT id, funcionario_id FROM ferias_controle WHERE id=?',(int(sel[0]),)).fetchone()
        return row

    def _dados_documento_ferias(self, ferias_id=None, funcionario_id=None):
        if ferias_id:
            with con() as db:
                r=db.execute('SELECT funcionario_id FROM ferias_controle WHERE id=?',(int(ferias_id),)).fetchone()
            if r:
                funcionario_id=r[0]
        if not funcionario_id:
            funcionario_id=self._ferias_func_id()
        if not funcionario_id:
            return None
        dados=self._dados_documento_funcionario(funcionario_id)
        if not dados:
            return None
        contexto=self._ferias_contexto_mais_recente(funcionario_id)
        dados.update(contexto)
        # Se houver dados na tela ainda não salvos, também leva para o documento.
        try:
            tela = {
                'PERIODO_AQUISITIVO_INICIO': self.ferias_vars['fer_aq_ini'].get(),
                'PERIODO_AQUISITIVO_FIM': self.ferias_vars['fer_aq_fim'].get(),
                'PERIODO_AQUISITIVO': f"{self.ferias_vars['fer_aq_ini'].get()} a {self.ferias_vars['fer_aq_fim'].get()}",
                'PERIODO_AQUISITIVO_COMPLETO': f"{self.ferias_vars['fer_aq_ini'].get()} a {self.ferias_vars['fer_aq_fim'].get()}",
                'PERIODO_CONCESSIVO_FIM': self.ferias_vars['fer_conc'].get(),
                'PERIODO_CONCESSIVO': 'até ' + self.ferias_vars['fer_conc'].get() if self.ferias_vars['fer_conc'].get() else '',
                'DATA_INICIO': self.ferias_vars['fer_ini'].get(), 'INICIO': self.ferias_vars['fer_ini'].get(),
                'DATA_FIM': self.ferias_vars['fer_fim'].get(), 'TERMINO': self.ferias_vars['fer_fim'].get(),
                'DATA_RETORNO': self.ferias_vars['fer_ret'].get(),
                'DIAS_FERIAS': self.fer_dias_gozar.get(), 'DIAS_GOZADOS': self.fer_dias_gozar.get(),
                'DIAS_ABONO': self.fer_dias_abono.get(), 'DIAS_RESTANTES': self.fer_dias_restantes.get(),
                'PERIODO': f"{self.ferias_vars['fer_ini'].get()} a {self.ferias_vars['fer_fim'].get()}",
                'SALARIO_BASE': 'R$ ' + self.fer_salario.get().strip().replace('R$','').strip(),
                'MEDIA_VARIAVEIS': 'R$ ' + self.fer_media.get().strip().replace('R$','').strip(),
                'OBS_FERIAS': self.fer_obs.get('1.0','end').strip(),
                'STATUS_FERIAS': self.fer_status.get()
            }
            for k,v in getattr(self,'fer_calc_vars',{}).items():
                tela[normalizar_placeholder(k)] = v.get()
            for k,v in tela.items():
                if v and str(v).strip() not in ('R$','R$ 0,00','0'):
                    dados[k]=v
        except Exception:
            pass
        return dados

    def _modelo_documento_ferias(self, tipo):
        mapa = {
            'Aviso de Férias': '08_Aviso_de_Ferias.docx',
            'Recibo de Férias': '09_Recibo_de_Ferias.docx',
            'Recibo de Abono': '09_Recibo_de_Ferias.docx',
            'Comunicação de Férias': '13_Comunicacao_de_Ferias.docx',
            'Termo de Ciência': '14_Termo_Ciencia_Ferias.docx',
            'Recibo de Pagamento das Férias': '09_Recibo_de_Ferias.docx'
        }
        return mapa.get(tipo, '08_Aviso_de_Ferias.docx')

    def gerar_documento_ferias(self, tipo='Aviso de Férias'):
        reg=self._ferias_registro_selecionado()
        ferias_id=reg[0] if reg else None
        funcionario_id=reg[1] if reg else self._ferias_func_id()
        dados=self._dados_documento_ferias(ferias_id, funcionario_id)
        if not dados:
            messagebox.showwarning('Férias','Selecione um funcionário ou um registro de férias.'); return
        modelo=os.path.join(MODELOS_DIR, self._modelo_documento_ferias(tipo))
        if not os.path.exists(modelo):
            messagebox.showerror('Férias','Modelo de férias não encontrado na pasta modelos_documentos.'); return
        pasta=os.path.join(DOCS_GERADOS_DIR, 'ferias')
        os.makedirs(pasta, exist_ok=True)
        seguro=self._safe_filename(dados.get('FUNCIONARIO'), 60)
        tipo_seguro=self._safe_filename(tipo, 40)
        data_arq=datetime.now().strftime('%Y%m%d_%H%M%S')
        destino=os.path.join(pasta, f'{tipo_seguro}_{seguro}_{data_arq}.docx')
        shutil.copy2(modelo, destino)
        try:
            self._substituir_placeholders_docx(destino, dados)
            with con() as db:
                db.execute("""INSERT INTO documentos_rh(funcionario_id,tipo,data,titulo,observacao,arquivo,ativo)
                              VALUES(?,?,?,?,?,?,1)""", (funcionario_id, tipo, date.today().isoformat(), tipo, 'Gerado diretamente no módulo Férias', destino))
            log_action(self.usuario,'FÉRIAS',f'Documento de férias gerado: {os.path.basename(destino)}')
            messagebox.showinfo('Férias', f'{tipo} gerado em documentos_gerados/ferias:\n{os.path.basename(destino)}')
            try:
                if os.name=='nt': os.startfile(destino)
            except Exception:
                pass
            self.carregar_documentos()
        except Exception as exc:
            messagebox.showerror('Erro ao gerar documento de férias', str(exc))
            try:
                if os.path.exists(destino): os.remove(destino)
            except Exception:
                pass

    def abrir_documentos_ferias(self):
        pasta=os.path.join(DOCS_GERADOS_DIR, 'ferias')
        os.makedirs(pasta, exist_ok=True)
        self._open(pasta)

    def concluir_ferias(self):
        sel=self.fer_tree.selection() if hasattr(self,'fer_tree') else []
        if not sel:
            messagebox.showwarning('Férias','Selecione um registro de férias.'); return
        if not messagebox.askyesno('Concluir Férias','Concluir férias, gerar ocorrência na folha e marcar como concluída?'):
            return
        with con() as db:
            row=db.execute('SELECT funcionario_id,inicio,fim,observacao FROM ferias_controle WHERE id=?',(int(sel[0]),)).fetchone()
            if row:
                db.execute('UPDATE ferias_controle SET status=? WHERE id=?',('Concluída',int(sel[0])))
                existe=db.execute('SELECT id FROM ocorrencias WHERE funcionario_id=? AND tipo=? AND data_inicio=? AND data_fim=? AND ativo=1',(row[0],'FÉRIAS',row[1],row[2])).fetchone()
                if not existe:
                    db.execute('INSERT INTO ocorrencias(funcionario_id,tipo,data_inicio,data_fim,observacao,abona,ativo) VALUES(?,?,?,?,?,1,1)',(row[0],'FÉRIAS',row[1],row[2],row[3] or 'Gerado ao concluir férias'))
        self.carregar_ferias(); self.carregar_ocorrencias(); self.refresh_dashboard()
        log_action(self.usuario,'FÉRIAS','Férias concluídas e ocorrência gerada')
        messagebox.showinfo('Férias','Férias concluídas. A ocorrência FÉRIAS foi criada para refletir na folha de ponto.')

    def enviar_ferias_para_documentos(self):
        # Mantido apenas por compatibilidade com versões antigas: agora férias gera documentos no próprio módulo Férias.
        self.gerar_documento_ferias('Aviso de Férias')

    def carregar_ferias(self):
        if hasattr(self,'combo_ferias_func'):
            vals=[f"{x['id']} - {x['nome']}" for x in get_funcionarios(True)]
            self.combo_ferias_func['values']=vals
            if vals and self.ferias_func_var.get() not in vals: self.ferias_func_var.set(vals[0])
            self.atualizar_periodos_ferias_funcionario()
        if not hasattr(self,'fer_tree'): return
        for i in self.fer_tree.get_children(): self.fer_tree.delete(i)
        with con() as db:
            rows=db.execute("""SELECT fc.id, f.nome, fc.aquisitivo_inicio, fc.aquisitivo_fim, fc.concessivo_fim, fc.inicio, fc.fim, fc.retorno, fc.dias, fc.dias_restantes, fc.total_bruto, fc.status
                               FROM ferias_controle fc JOIN funcionarios f ON f.id=fc.funcionario_id
                               WHERE fc.ativo=1 ORDER BY date(fc.inicio) DESC, f.nome""").fetchall()
        for rid,nome,aqi,aqf,conc,ini,fim,ret,dias,rest,total,status in rows:
            self.fer_tree.insert('', 'end', iid=str(rid), values=(nome, f'{fmt_data(aqi)} a {fmt_data(aqf)}', fmt_data(conc), f'{fmt_data(ini)} a {fmt_data(fim)}', fmt_data(ret), dias or '', rest or '', moeda_br(total or 0), status))

    def salvar_ferias(self):
        val=self.ferias_func_var.get().strip()
        if not val: messagebox.showwarning('Atenção','Selecione o funcionário.'); return
        try:
            fid=int(val.split(' - ')[0])
            dias=int(str(self.fer_dias_gozar.get() or '0').strip() or 0)
            dias_abono=int(str(self.fer_dias_abono.get() or '0').strip() or 0)
            if dias <= 0:
                messagebox.showwarning('Atenção','Informe a quantidade de dias de férias a gozar.'); return
            ini=parse_data_br(self.ferias_vars['fer_ini'].get())
            fim, ret_auto = adicionar_dias_corridos(ini, dias)
            self.ferias_vars['fer_fim'].set(fim.strftime('%d/%m/%Y'))
            self.ferias_vars['fer_ret'].set(ret_auto.strftime('%d/%m/%Y'))
            aqi=parse_data_br(self.ferias_vars['fer_aq_ini'].get()) if self.ferias_vars['fer_aq_ini'].get().strip() else None
            aqf=parse_data_br(self.ferias_vars['fer_aq_fim'].get()) if self.ferias_vars['fer_aq_fim'].get().strip() else None
            if not aqi or not aqf:
                with con() as db:
                    adm_row = db.execute('SELECT admissao FROM funcionarios WHERE id=?', (fid,)).fetchone()
                aq_ini_br, aq_fim_br, _ = calcular_periodo_aquisitivo(adm_row[0] if adm_row else None, ini)
                aqi = parse_data_br(aq_ini_br) if aq_ini_br else None
                aqf = parse_data_br(aq_fim_br) if aq_fim_br else None
            conc=parse_data_br(self.ferias_vars['fer_conc'].get()) if self.ferias_vars['fer_conc'].get().strip() else (_add_years_safe(aqf, 1) if aqf else None)
            ret=parse_data_br(self.ferias_vars['fer_ret'].get()) if self.ferias_vars['fer_ret'].get().strip() else ret_auto
            vals=self.calcular_ferias_tela() or calcular_ferias_valores(self.fer_salario.get(), dias, dias_abono, self.fer_media.get(), bool(self.fer_adianta_13.get()))
            dias_rest=int(self.fer_dias_restantes.get() or 0) if hasattr(self,'fer_dias_restantes') else max(0,30-dias)
        except Exception as exc:
            messagebox.showwarning('Atenção','Confira datas, dias e valores informados.\n'+str(exc)); return
        with con() as db:
            db.execute("""INSERT INTO ferias_controle(funcionario_id,aquisitivo_inicio,aquisitivo_fim,concessivo_fim,inicio,fim,retorno,dias,observacao,status,ativo,
                          dias_abono,dias_restantes,salario_base,media_variaveis,valor_ferias,valor_um_terco,valor_abono,valor_13,total_bruto,inss_estimado,irrf_estimado,liquido_estimado)
                          VALUES(?,?,?,?,?,?,?,?,?,?,1,?,?,?,?,?,?,?,?,?,?,?,?)""",(
                          fid, aqi.isoformat() if aqi else '', aqf.isoformat() if aqf else '', conc.isoformat() if conc else '', ini.isoformat(), fim.isoformat(), ret.isoformat(), dias,
                          self.fer_obs.get('1.0','end').strip(), self.fer_status.get(), dias_abono, dias_rest, vals.get('salario_base',0), vals.get('media_variaveis',0), vals.get('valor_ferias',0), vals.get('valor_um_terco',0), vals.get('valor_abono',0), vals.get('valor_13',0), vals.get('total_bruto',0), vals.get('inss_estimado',0), vals.get('irrf_estimado',0), vals.get('liquido_estimado',0)))
        log_action(self.usuario,'FÉRIAS',f'Férias lançadas para {val}: {ini.strftime("%d/%m/%Y")} a {fim.strftime("%d/%m/%Y")} - total {moeda_br(vals.get("total_bruto",0))}')
        self.carregar_ferias(); self.atualizar_periodos_ferias_funcionario(); self.refresh_dashboard(); messagebox.showinfo('Férias','Registro de férias salvo com cálculo de datas e valores.')

    def cancelar_ferias(self):
        sel=self.fer_tree.selection() if hasattr(self,'fer_tree') else []
        if not sel: messagebox.showwarning('Atenção','Selecione um registro.'); return
        if not messagebox.askyesno('Confirmar','Cancelar este registro de férias?'): return
        with con() as db: db.execute('UPDATE ferias_controle SET ativo=0,status=? WHERE id=?',('Cancelada',int(sel[0])))
        log_action(self.usuario,'FÉRIAS','Férias canceladas'); self.carregar_ferias(); self.refresh_dashboard()

    def gerar_ocorrencia_ferias(self):
        sel=self.fer_tree.selection() if hasattr(self,'fer_tree') else []
        if not sel: messagebox.showwarning('Atenção','Selecione um registro de férias.'); return
        with con() as db:
            row=db.execute('SELECT funcionario_id,inicio,fim,observacao FROM ferias_controle WHERE id=?',(int(sel[0]),)).fetchone()
            if row:
                db.execute('INSERT INTO ocorrencias(funcionario_id,tipo,data_inicio,data_fim,observacao,abona,ativo) VALUES(?,?,?,?,?,1,1)',(row[0],'FÉRIAS',row[1],row[2],row[3] or 'Gerado pelo controle de férias'))
        self.carregar_ocorrencias(); log_action(self.usuario,'FÉRIAS','Ocorrência de férias gerada'); messagebox.showinfo('Férias','Ocorrência FÉRIAS criada para aparecer na folha.')

    def build_banco_horas(self):
        f=self.tab_banco
        ttk.Label(f,text='Banco de Horas',style='Title.TLabel').grid(row=0,column=0,columnspan=6,sticky='w',padx=16,pady=16)
        ttk.Label(f,text='Funcionário').grid(row=1,column=0,sticky='w',padx=12,pady=6)
        self.bh_func_var=tk.StringVar(); self.combo_bh_func=ttk.Combobox(f,textvariable=self.bh_func_var,width=65,state='readonly')
        self.combo_bh_func.grid(row=1,column=1,columnspan=5,sticky='ew',padx=8,pady=6)
        ttk.Label(f,text='Data DD/MM/AAAA').grid(row=2,column=0,sticky='w',padx=12,pady=6)
        self.bh_data=tk.StringVar(value=datetime.now().strftime('%d/%m/%Y')); e=ttk.Entry(f,textvariable=self.bh_data,width=18); e.grid(row=2,column=1,sticky='ew',padx=8,pady=6); e.bind('<FocusOut>', lambda ev: formatar_campo_data(self.bh_data))
        ttk.Label(f,text='Tipo').grid(row=2,column=2,sticky='w',padx=12,pady=6)
        self.bh_tipo=tk.StringVar(value='Crédito'); ttk.Combobox(f,textvariable=self.bh_tipo,values=['Crédito','Débito','Ajuste'],state='readonly').grid(row=2,column=3,sticky='ew',padx=8,pady=6)
        ttk.Label(f,text='Quantidade HH:MM').grid(row=2,column=4,sticky='w',padx=12,pady=6)
        self.bh_qtd=tk.StringVar(value='00:00'); ttk.Entry(f,textvariable=self.bh_qtd,width=10).grid(row=2,column=5,sticky='ew',padx=8,pady=6)
        ttk.Label(f,text='Motivo').grid(row=3,column=0,sticky='w',padx=12,pady=6)
        self.bh_motivo=tk.StringVar(); ttk.Entry(f,textvariable=self.bh_motivo).grid(row=3,column=1,columnspan=5,sticky='ew',padx=8,pady=6)
        ttk.Button(f,text='Lançar',command=self.salvar_banco_horas).grid(row=4,column=1,sticky='ew',padx=8,pady=10)
        ttk.Button(f,text='Cancelar lançamento',command=self.cancelar_banco_horas).grid(row=4,column=2,sticky='ew',padx=8,pady=10)
        ttk.Button(f,text='Exportar extrato CSV',command=self.exportar_banco_horas).grid(row=4,column=3,sticky='ew',padx=8,pady=10)
        self.bh_tree=ttk.Treeview(f,columns=('func','data','tipo','qtd','motivo','usuario'),show='headings')
        for col,txt,w in [('func','Funcionário',230),('data','Data',90),('tipo','Tipo',90),('qtd','Horas',80),('motivo','Motivo',280),('usuario','Usuário',100)]:
            self.bh_tree.heading(col,text=txt); self.bh_tree.column(col,width=w)
        self.bh_tree.grid(row=5,column=0,columnspan=6,sticky='nsew',padx=12,pady=8)
        f.rowconfigure(5,weight=1)
        for c in range(6): f.columnconfigure(c,weight=1)

    def carregar_banco_horas(self):
        if hasattr(self,'combo_bh_func'):
            vals=[f"{x['id']} - {x['nome']}" for x in get_funcionarios(True)]
            self.combo_bh_func['values']=vals
            if vals and self.bh_func_var.get() not in vals: self.bh_func_var.set(vals[0])
        if not hasattr(self,'bh_tree'): return
        for i in self.bh_tree.get_children(): self.bh_tree.delete(i)
        with con() as db:
            rows=db.execute("""SELECT bh.id,f.nome,bh.data,bh.tipo,bh.quantidade,bh.motivo,bh.usuario FROM banco_horas bh JOIN funcionarios f ON f.id=bh.funcionario_id WHERE bh.ativo=1 ORDER BY date(bh.data) DESC, f.nome""").fetchall()
        for rid,nome,data,tipo,qtd,motivo,usuario in rows:
            self.bh_tree.insert('', 'end', iid=str(rid), values=(nome,fmt_data(data),tipo,qtd,motivo or '',usuario or ''))

    def salvar_banco_horas(self):
        val=self.bh_func_var.get().strip()
        if not val: messagebox.showwarning('Atenção','Selecione o funcionário.'); return
        try:
            fid=int(val.split(' - ')[0]); data=parse_data_br(self.bh_data.get())
        except Exception:
            messagebox.showwarning('Atenção','Confira a data.'); return
        qtd=self.bh_qtd.get().strip()
        if ':' not in qtd: messagebox.showwarning('Atenção','Informe a quantidade em HH:MM.'); return
        with con() as db:
            db.execute('INSERT INTO banco_horas(funcionario_id,data,tipo,quantidade,motivo,usuario,criado_em,ativo) VALUES(?,?,?,?,?,?,?,1)',(fid,data.isoformat(),self.bh_tipo.get(),qtd,self.bh_motivo.get().strip(),self.usuario,datetime.now().strftime('%d/%m/%Y %H:%M:%S')))
        log_action(self.usuario,'BANCO DE HORAS',f'{self.bh_tipo.get()} {qtd} para {val}')
        self.carregar_banco_horas(); messagebox.showinfo('Banco de Horas','Lançamento salvo.')

    def cancelar_banco_horas(self):
        sel=self.bh_tree.selection() if hasattr(self,'bh_tree') else []
        if not sel: messagebox.showwarning('Atenção','Selecione um lançamento.'); return
        if not messagebox.askyesno('Confirmar','Cancelar este lançamento?'): return
        with con() as db: db.execute('UPDATE banco_horas SET ativo=0 WHERE id=?',(int(sel[0]),))
        log_action(self.usuario,'BANCO DE HORAS','Lançamento cancelado'); self.carregar_banco_horas()

    def exportar_banco_horas(self):
        os.makedirs(RELATORIO_DIR, exist_ok=True); path=os.path.join(RELATORIO_DIR,'extrato_banco_horas.csv')
        with con() as db:
            rows=db.execute("""SELECT f.nome,bh.data,bh.tipo,bh.quantidade,bh.motivo,bh.usuario,bh.criado_em FROM banco_horas bh JOIN funcionarios f ON f.id=bh.funcionario_id WHERE bh.ativo=1 ORDER BY f.nome,date(bh.data)""").fetchall()
        with open(path,'w',newline='',encoding='utf-8-sig') as fp:
            wr=csv.writer(fp,delimiter=';'); wr.writerow(['Funcionário','Data','Tipo','Horas','Motivo','Usuário','Criado em'])
            for nome,data,tipo,qtd,motivo,usuario,criado in rows: wr.writerow([nome,fmt_data(data),tipo,qtd,motivo,usuario,criado])
        log_action(self.usuario,'RELATÓRIO','Extrato banco de horas'); messagebox.showinfo('Banco de Horas','Extrato salvo em:\n'+path)

    def build_atualizador(self):
        f=self.tab_atualizador
        ttk.Label(f,text='Atualizador Automático',style='Title.TLabel').grid(row=0,column=0,columnspan=4,sticky='w',padx=16,pady=16)
        texto=('Este módulo prepara o Sistema Gestão Izzant para atualizações futuras sem perder o banco de dados.\n\n'
               'Fluxo recomendado:\n1. Fazer backup antes de atualizar.\n2. Substituir somente os arquivos do programa.\n3. Manter a pasta dados intacta.\n4. Executar o sistema para aplicar migrações automáticas do banco.')
        ttk.Label(f,text=texto,justify='left',font=('Arial',11)).grid(row=1,column=0,columnspan=4,sticky='w',padx=16,pady=8)
        ttk.Button(f,text='Fazer backup antes de atualizar',command=self.backup_now).grid(row=2,column=0,sticky='ew',padx=12,pady=10)
        ttk.Button(f,text='Abrir pasta do programa',command=self.open_base).grid(row=2,column=1,sticky='ew',padx=12,pady=10)
        ttk.Button(f,text='Abrir pasta de backups',command=self.open_backups).grid(row=2,column=2,sticky='ew',padx=12,pady=10)
        ttk.Button(f,text='Gerar instruções de instalação',command=self.gerar_instrucoes_instalacao).grid(row=2,column=3,sticky='ew',padx=12,pady=10)
        for c in range(4): f.columnconfigure(c,weight=1)

    def gerar_instrucoes_instalacao(self):
        path=os.path.join(BASE_DIR,'INSTALACAO_WINDOWS_EXE.txt')
        with open(path,'w',encoding='utf-8') as fp:
            fp.write('Sistema Gestão Izzant - Instalação Windows\n\n')
            fp.write('1. Extraia o pacote em uma pasta fixa.\n2. Execute INSTALAR_BIBLIOTECAS.bat uma vez.\n3. Use INICIAR_PROGRAMA.bat para abrir.\n4. Para gerar .exe, instale o PyInstaller e execute: pyinstaller --onefile --windowed app.py\n5. Preserve a pasta dados para manter o banco SQLite.\n')
        messagebox.showinfo('Instalação', 'Arquivo gerado:\n'+path)

    def build_pdf(self):
        f=self.tab_pdf
        ttk.Label(f,text='Gerar PDF para Impressão',style='Title.TLabel').grid(row=0,column=0,columnspan=4,sticky='w',padx=16,pady=16)
        ttk.Label(f,text='Mês').grid(row=1,column=0,sticky='w',padx=16,pady=6)
        self.mes_var=tk.IntVar(value=datetime.now().month)
        ttk.Combobox(f,textvariable=self.mes_var,values=list(range(1,13)),width=8,state='readonly').grid(row=1,column=1,sticky='w',padx=8,pady=6)
        ttk.Label(f,text='Ano').grid(row=1,column=2,sticky='w',padx=16,pady=6)
        self.ano_var=tk.IntVar(value=datetime.now().year)
        ttk.Entry(f,textvariable=self.ano_var,width=12).grid(row=1,column=3,sticky='w',padx=8,pady=6)
        ttk.Label(f,text='Funcionário').grid(row=2,column=0,sticky='w',padx=16,pady=6)
        self.func_pdf_var=tk.StringVar(value='TODOS')
        self.combo_func=ttk.Combobox(f,textvariable=self.func_pdf_var,width=65,state='readonly')
        self.combo_func.grid(row=2,column=1,columnspan=3,sticky='ew',padx=8,pady=6)
        ttk.Label(f,text='Setor').grid(row=3,column=0,sticky='w',padx=16,pady=6)
        self.setor_pdf_var=tk.StringVar(value='TODOS')
        self.combo_setor=ttk.Combobox(f,textvariable=self.setor_pdf_var,width=35,state='readonly')
        self.combo_setor.grid(row=3,column=1,columnspan=3,sticky='ew',padx=8,pady=6)
        ttk.Button(f,text='Visualizar PDF antes de gerar',command=self.visualizar_pdf).grid(row=4,column=0,sticky='ew',padx=8,pady=12)
        ttk.Button(f,text='Gerar PDF de todos',command=lambda:self.gerar_pdf(True)).grid(row=4,column=1,sticky='ew',padx=8,pady=12)
        ttk.Button(f,text='Gerar PDF selecionado',command=lambda:self.gerar_pdf(False)).grid(row=4,column=2,sticky='ew',padx=8,pady=12)
        ttk.Button(f,text='Gerar PDF por setor',command=self.gerar_pdf_setor).grid(row=4,column=3,sticky='ew',padx=8,pady=12)
        ttk.Button(f,text='Gerar PDFs separados por setor',command=self.gerar_pdfs_todos_setores).grid(row=5,column=0,sticky='ew',padx=8,pady=8)
        ttk.Button(f,text='Abrir último PDF',command=self.abrir_ultimo_pdf).grid(row=5,column=1,sticky='ew',padx=8,pady=8)
        ttk.Button(f,text='Abrir pasta PDFs',command=self.open_pdfs).grid(row=5,column=2,sticky='ew',padx=8,pady=8)
        ttk.Button(f,text='Gerar PDFs individuais de todos',command=self.gerar_pdfs_individuais_todos).grid(row=5,column=3,sticky='ew',padx=8,pady=8)
        ttk.Button(f,text='Atualizar histórico',command=self.carregar_historico_pdf).grid(row=6,column=3,sticky='ew',padx=8,pady=8)
        self.pdf_info=ttk.Label(f,text='Os arquivos serão salvos automaticamente em PDFs/ANO/MÊS. Use Visualizar PDF para conferir antes de gerar.',font=('Arial',11))
        self.pdf_info.grid(row=7,column=0,columnspan=4,sticky='w',padx=16)
        self.hist_pdf=ttk.Treeview(f, columns=('data','tipo','setor','mes','ano','qtd','arquivo'), show='headings', height=10)
        for col,txt_h,w in [('data','Data/Hora',140),('tipo','Tipo',140),('setor','Setor',110),('mes','Mês',50),('ano','Ano',55),('qtd','Qtd.',45),('arquivo','Arquivo',430)]:
            self.hist_pdf.heading(col,text=txt_h); self.hist_pdf.column(col,width=w)
        self.hist_pdf.grid(row=8,column=0,columnspan=4,sticky='nsew',padx=12,pady=12)
        f.rowconfigure(8,weight=1)
        for c in range(5): f.columnconfigure(c,weight=1)


    def _funcionario_options(self):
        return [f"{r[0]} - {r[1]}" for r in con().execute("SELECT id,nome FROM funcionarios ORDER BY nome").fetchall()]

    def _func_id_from_option(self, txt):
        try: return int(str(txt).split(' - ')[0])
        except Exception: return None

    def build_documentos(self):
        f=self.tab_documentos
        ttk.Label(f,text='Central de Documentos RH - Gerais',style='Title.TLabel').grid(row=0,column=0,columnspan=6,sticky='w',padx=16,pady=16)
        ttk.Label(f,text='Selecione o funcionário e o modelo. Documentos de férias são gerados exclusivamente no módulo Férias.',wraplength=980).grid(row=1,column=0,columnspan=6,sticky='w',padx=16,pady=(0,8))
        self.doc_func=tk.StringVar(); self.doc_tipo=tk.StringVar(value='Declaração'); self.doc_data=tk.StringVar(value=date.today().strftime('%d/%m/%Y')); self.doc_titulo=tk.StringVar()
        self.doc_extra_vars={}
        ttk.Label(f,text='Funcionário').grid(row=2,column=0,sticky='w',padx=8); self.doc_combo=ttk.Combobox(f,textvariable=self.doc_func,values=self._funcionario_options(),width=34); self.doc_combo.grid(row=3,column=0,sticky='ew',padx=8); self.doc_combo.bind('<<ComboboxSelected>>', lambda e: self.preencher_campos_automaticos_documento())
        ttk.Label(f,text='Modelo').grid(row=2,column=1,sticky='w',padx=8); self.doc_tipo_combo=ttk.Combobox(f,textvariable=self.doc_tipo,values=['Contrato','Advertência','Suspensão','Declaração','Ficha de Registro','Termo','Recibo','EPI','Uniforme','Banco de Horas'],width=18,state='readonly'); self.doc_tipo_combo.grid(row=3,column=1,sticky='ew',padx=8); self.doc_tipo_combo.bind('<<ComboboxSelected>>', lambda e: (self.atualizar_campos_documento(), self.preencher_campos_automaticos_documento()))
        ttk.Label(f,text='Data do documento').grid(row=2,column=2,sticky='w',padx=8); ttk.Entry(f,textvariable=self.doc_data,width=12).grid(row=3,column=2,sticky='ew',padx=8)
        ttk.Label(f,text='Título/assunto').grid(row=2,column=3,sticky='w',padx=8); ttk.Entry(f,textvariable=self.doc_titulo,width=28).grid(row=3,column=3,sticky='ew',padx=8)
        ttk.Button(f,text='Salvar registro',command=self.salvar_documento).grid(row=3,column=4,sticky='ew',padx=8)
        ttk.Button(f,text='Gerar preenchido',command=self.gerar_copia_modelo_documento).grid(row=3,column=5,sticky='ew',padx=8)

        self.doc_campos_frame=ttk.LabelFrame(f,text='Campos do modelo selecionado')
        self.doc_campos_frame.grid(row=4,column=0,columnspan=6,sticky='ew',padx=8,pady=8)
        ttk.Button(f,text='Abrir modelos',command=self.abrir_modelos_documentos).grid(row=5,column=0,sticky='ew',padx=8,pady=4)
        ttk.Button(f,text='Abrir documentos gerados',command=self.abrir_documentos_gerados).grid(row=5,column=1,sticky='ew',padx=8,pady=4)
        ttk.Button(f,text='Limpar campos',command=self.limpar_campos_documento).grid(row=5,column=2,sticky='ew',padx=8,pady=4)
        ttk.Button(f,text='Ver placeholders',command=self.mostrar_placeholders_documento).grid(row=5,column=3,sticky='ew',padx=8,pady=4)
        ttk.Button(f,text='Importar modelo Word',command=self.importar_modelo_word_documento).grid(row=5,column=4,sticky='ew',padx=8,pady=4)
        ttk.Button(f,text='Importar modelo mestre',command=self.importar_modelo_mestre_documento).grid(row=5,column=5,sticky='ew',padx=8,pady=4)
        ttk.Label(f,text='Observações / bloco de escrita geral').grid(row=6,column=0,columnspan=6,sticky='w',padx=8)
        self.doc_obs=tk.Text(f,height=4); self.doc_obs.grid(row=7,column=0,columnspan=6,sticky='ew',padx=8,pady=8)
        self.doc_tree=ttk.Treeview(f,columns=('data','func','tipo','titulo','obs'),show='headings',height=12)
        for col,txt,w in [('data','Data',90),('func','Funcionário',240),('tipo','Tipo',130),('titulo','Título',220),('obs','Observação',380)]: self.doc_tree.heading(col,text=txt); self.doc_tree.column(col,width=w)
        self.doc_tree.grid(row=8,column=0,columnspan=6,sticky='nsew',padx=8,pady=8); f.rowconfigure(8,weight=1)
        self.atualizar_campos_documento()

    def atualizar_campos_documento(self):
        if not hasattr(self, 'doc_campos_frame'): return
        for w in self.doc_campos_frame.winfo_children():
            w.destroy()
        self.doc_extra_vars={}
        tipo = self.doc_tipo.get()
        campos = list(DOC_CAMPOS_MODELO.get(tipo, []))
        # Também lê automaticamente os placeholders existentes no Word do modelo.
        # Campos já conhecidos pelo sistema são preenchidos automaticamente;
        # campos desconhecidos viram campos editáveis nesta tela.
        modelo = os.path.join(MODELOS_DIR, self._modelo_por_tipo(tipo))
        chaves_atuais = {normalizar_placeholder(chave) for chave, _ in campos}
        for ph in self._placeholders_no_docx(modelo):
            if ph not in CAMPOS_AUTOMATICOS_DOCUMENTOS and ph not in chaves_atuais:
                campos.append((ph, ph.replace('_', ' ').title()))
                chaves_atuais.add(ph)
        if not campos:
            ttk.Label(self.doc_campos_frame,text='Este modelo não possui campos específicos. Use o bloco de observações.').grid(row=0,column=0,sticky='w',padx=8,pady=8)
            return
        for i,(chave,label) in enumerate(campos):
            chave_norm = normalizar_placeholder(chave)
            var=tk.StringVar()
            self.doc_extra_vars[chave_norm]=var
            r=(i//3)*2; c=i%3
            ttk.Label(self.doc_campos_frame,text=label).grid(row=r,column=c,sticky='w',padx=8,pady=(6,0))
            ent = ttk.Entry(self.doc_campos_frame,textvariable=var,width=34)
            ent.grid(row=r+1,column=c,sticky='ew',padx=8,pady=(0,6))
            if chave_norm in ('DATA_INICIO','DATA_FIM','INICIO','TERMINO'):
                ent.bind('<FocusOut>', lambda e: self.preencher_campos_automaticos_documento())
                ent.bind('<Return>', lambda e: self.preencher_campos_automaticos_documento())
        for c in range(3): self.doc_campos_frame.columnconfigure(c,weight=1)
        self.preencher_campos_automaticos_documento()

    def preencher_campos_automaticos_documento(self):
        """Preenche campos automáticos dos documentos ao selecionar o funcionário.
        Especialmente em modelos de Férias, calcula período aquisitivo pela admissão.
        """
        if getattr(self, '_doc_auto_fill_running', False):
            return
        self._doc_auto_fill_running = True
        try:
            extras = getattr(self, 'doc_extra_vars', {})
            fid = self._func_id_from_option(self.doc_func.get()) if hasattr(self, 'doc_func') else None
            admissao = None
            if fid:
                with con() as db:
                    row = db.execute('SELECT admissao FROM funcionarios WHERE id=?', (fid,)).fetchone()
                    admissao = row[0] if row else None
            if fid and hasattr(self, 'doc_tipo') and self.doc_tipo.get() == 'Férias':
                contexto = self._ferias_contexto_mais_recente(fid)
                for chave, valor in contexto.items():
                    if chave in extras and not extras[chave].get().strip():
                        extras[chave].set(valor)
            data_inicio = ''
            if 'DATA_INICIO' in extras:
                data_inicio = extras['DATA_INICIO'].get().strip()
            elif 'INICIO' in extras:
                data_inicio = extras['INICIO'].get().strip()
            data_fim = ''
            if 'DATA_FIM' in extras:
                data_fim = extras['DATA_FIM'].get().strip()
            elif 'TERMINO' in extras:
                data_fim = extras['TERMINO'].get().strip()
            ref = data_inicio or (self.doc_data.get().strip() if hasattr(self, 'doc_data') else None) or date.today()
            aq_ini, aq_fim, aq_completo = calcular_periodo_aquisitivo(admissao, ref)
            auto_vals = {
                'PERIODO_AQUISITIVO': aq_completo,
                'PERIODO_AQUISITIVO_INICIO': aq_ini,
                'PERIODO_AQUISITIVO_FIM': aq_fim,
                'PERIODO_AQUISITIVO_COMPLETO': aq_completo,
            }
            if data_inicio and data_fim:
                auto_vals['PERIODO'] = f'{data_inicio} a {data_fim}'
                try:
                    ini_f = parse_data_br(data_inicio)
                    fim_f = parse_data_br(data_fim)
                    auto_vals['DIAS_FERIAS'] = str((fim_f - ini_f).days + 1)
                    auto_vals['DATA_RETORNO'] = (fim_f + timedelta(days=1)).strftime('%d/%m/%Y')
                except Exception:
                    pass
            for chave, valor in auto_vals.items():
                if chave in extras and not extras[chave].get().strip():
                    extras[chave].set(valor)
        finally:
            self._doc_auto_fill_running = False

    def limpar_campos_documento(self):
        if hasattr(self,'doc_titulo'): self.doc_titulo.set('')
        if hasattr(self,'doc_obs'): self.doc_obs.delete('1.0','end')
        for var in getattr(self,'doc_extra_vars',{}).values(): var.set('')

    def mostrar_placeholders_documento(self):
        tipo = self.doc_tipo.get()
        campos = DOC_CAMPOS_MODELO.get(tipo, [])
        base = ['{{EMPRESA}}','{{CNPJ}}','{{ENDERECO}}','{{CIDADE}}','{{UF}}','{{FUNCIONARIO}}','{{NOME}}','{{CPF}}','{{CTPS}}','{{ADMISSAO}}','{{FUNCAO}}','{{CARGO}}','{{SETOR}}','{{CBO}}','{{SALARIO}}','{{JORNADA}}','{{HORARIO_TRABALHO}}','{{DATA_DOCUMENTO}}','{{TITULO}}','{{OBSERVACAO}}','{{LOCAL_DATA}}','{{PERIODO_AQUISITIVO}}','{{PERIODO_AQUISITIVO_INICIO}}','{{PERIODO_AQUISITIVO_FIM}}','{{PERIODO_AQUISITIVO_COMPLETO}}']
        extras = ['{{'+normalizar_placeholder(chave)+'}}' for chave,_ in campos]
        modelo = os.path.join(MODELOS_DIR, self._modelo_por_tipo(tipo))
        dinamicos = ['{{'+x+'}}' for x in self._placeholders_no_docx(modelo) if x not in CAMPOS_AUTOMATICOS_DOCUMENTOS]
        messagebox.showinfo('Placeholders do modelo', 'Campos disponíveis para usar no Word:\n\n' + '\n'.join(sorted(set(base+extras+dinamicos))))

    def salvar_documento(self):
        fid=self._func_id_from_option(self.doc_func.get())
        if not fid: messagebox.showwarning('Documentos','Selecione um funcionário.'); return
        extras = []
        for chave,var in getattr(self,'doc_extra_vars',{}).items():
            val = var.get().strip()
            if val:
                extras.append(f'{chave}: {val}')
        obs = self.doc_obs.get('1.0','end').strip()
        obs_completa = (obs + ('\n' if obs and extras else '') + '\n'.join(extras)).strip()
        with con() as db: db.execute('INSERT INTO documentos_rh(funcionario_id,tipo,data,titulo,observacao,ativo) VALUES(?,?,?,?,?,1)',(fid,self.doc_tipo.get(),parse_data_br(self.doc_data.get()).isoformat(),self.doc_titulo.get(),obs_completa))
        log_action(self.usuario,'DOCUMENTOS','Documento RH cadastrado'); self.carregar_documentos(); messagebox.showinfo('Documentos','Documento salvo.')

    def carregar_documentos(self):
        if not hasattr(self,'doc_tree'): return
        self.doc_tree.delete(*self.doc_tree.get_children())
        for r in con().execute("SELECT d.data,f.nome,d.tipo,d.titulo,d.observacao FROM documentos_rh d LEFT JOIN funcionarios f ON f.id=d.funcionario_id WHERE d.ativo=1 ORDER BY d.data DESC LIMIT 300"):
            self.doc_tree.insert('', 'end', values=(fmt_data(r[0]),r[1],r[2],r[3],r[4]))

    def abrir_modelos_documentos(self):
        self._open(MODELOS_DIR)

    def abrir_documentos_gerados(self):
        self._open(DOCS_GERADOS_DIR)

    def _placeholders_no_docx(self, caminho_docx):
        """Lê placeholders {{CAMPO}} no Word e retorna nomes normalizados."""
        encontrados = set()
        if not os.path.exists(caminho_docx):
            return []
        try:
            with zipfile.ZipFile(caminho_docx, 'r') as z:
                for nome in z.namelist():
                    if nome.startswith('word/') and nome.endswith('.xml'):
                        try:
                            xml = z.read(nome).decode('utf-8', errors='ignore')
                        except Exception:
                            continue
                        for bruto in re.findall(r'\{\{\s*([^{}]+?)\s*\}\}', xml):
                            chave = normalizar_placeholder(bruto)
                            if chave:
                                encontrados.add(chave)
        except Exception:
            pass
        return sorted(encontrados)

    def importar_modelo_mestre_documento(self):
        origem = filedialog.askopenfilename(title='Selecionar modelo mestre Word (.docx)', filetypes=[('Documento Word','*.docx')])
        if not origem:
            return
        try:
            os.makedirs(MODELOS_DIR, exist_ok=True)
            if os.path.exists(MODELO_MESTRE_WORD):
                backup = MODELO_MESTRE_WORD.replace('.docx', '_' + datetime.now().strftime('%Y%m%d_%H%M%S') + '_backup.docx')
                shutil.copy2(MODELO_MESTRE_WORD, backup)
            shutil.copy2(origem, MODELO_MESTRE_WORD)
            log_action(self.usuario,'DOCUMENTOS','Modelo mestre Word importado')
            messagebox.showinfo('Modelo mestre', 'Modelo mestre importado. Ele será usado como identidade visual/base dos novos modelos Word.')
        except Exception as exc:
            messagebox.showerror('Modelo mestre', str(exc))

    def _modelo_por_tipo(self, tipo):
        mapa = {
            'Contrato':'01_Contrato_de_Trabalho.docx',
            'Advertência':'02_Advertencia_Disciplinar.docx',
            'Suspensão':'03_Suspensao_Disciplinar.docx',
            'Declaração':'04_Declaracao_de_Vinculo.docx',
            'Ficha de Registro':'05_Ficha_de_Registro.docx',
            'Termo':'12_Termo_de_Confidencialidade.docx',
            'Recibo':'11_Recibo_de_Pagamento.docx',
            'EPI':'06_Termo_de_Entrega_de_EPI.docx',
            'Uniforme':'07_Termo_de_Entrega_de_Uniforme.docx',
            'Banco de Horas':'10_Termo_Banco_de_Horas.docx'
        }
        return mapa.get(tipo, '04_Declaracao_de_Vinculo.docx')

    def importar_modelo_word_documento(self):
        tipo = self.doc_tipo.get() or 'Declaração'
        origem = filedialog.askopenfilename(title='Selecionar modelo Word (.docx)', filetypes=[('Documento Word','*.docx')])
        if not origem:
            return
        try:
            destino = os.path.join(MODELOS_DIR, self._modelo_por_tipo(tipo))
            os.makedirs(MODELOS_DIR, exist_ok=True)
            if os.path.exists(destino):
                backup = destino.replace('.docx', '_' + datetime.now().strftime('%Y%m%d_%H%M%S') + '_backup.docx')
                shutil.copy2(destino, backup)
            shutil.copy2(origem, destino)
            log_action(self.usuario,'DOCUMENTOS',f'Modelo Word importado para {tipo}')
            messagebox.showinfo('Modelo Word', f'Modelo importado para o tipo {tipo}.\nUse placeholders como {{FUNCIONARIO}}, {{CPF}}, {{PERIODO_AQUISITIVO}} etc.')
            self.atualizar_campos_documento()
        except Exception as exc:
            messagebox.showerror('Modelo Word', str(exc))

    def _safe_filename(self, texto, limite=80):
        texto = texto or 'arquivo'
        seguro = ''.join(ch for ch in texto if ch.isalnum() or ch in (' ', '_', '-')).strip()
        seguro = seguro.replace(' ', '_')
        return (seguro[:limite] or 'arquivo')

    def _dados_documento_funcionario(self, fid):
        with con() as db:
            f = db.execute("""SELECT nome, cpf, ctps, admissao, funcao, setor, cbo, salario, jornada, horario_trabalho, descanso, sabado
                              FROM funcionarios WHERE id=?""", (fid,)).fetchone()
            e = db.execute('SELECT nome, cnpj, endereco, numero, bairro, cidade, uf FROM empresa WHERE id=1').fetchone()
        if not f:
            return None
        nome_empresa = e[0] if e and e[0] else DEFAULT_EMPRESA['nome']
        endereco = ''
        if e:
            partes = [e[2] or '', e[3] or '', e[4] or '', e[5] or '', e[6] or '']
            endereco = ' - '.join([x for x in partes if x])
        hoje = date.today()
        data_doc = self.doc_data.get().strip() if hasattr(self, 'doc_data') else hoje.strftime('%d/%m/%Y')
        obs = self.doc_obs.get('1.0', 'end').strip() if hasattr(self, 'doc_obs') else ''
        titulo = self.doc_titulo.get().strip() if hasattr(self, 'doc_titulo') else ''
        salario = f[7] or 0
        try:
            salario_fmt = f'R$ {float(salario):,.2f}'.replace(',', 'X').replace('.', ',').replace('X', '.')
        except Exception:
            salario_fmt = str(salario or '')
        adm = fmt_data(f[3]) if f[3] else ''
        extras_documento = {}
        if hasattr(self, 'doc_tipo') and self.doc_tipo.get() == 'Férias':
            extras_documento.update(self._ferias_contexto_mais_recente(fid))
        for chave, var in getattr(self, 'doc_extra_vars', {}).items():
            valor = var.get().strip() if hasattr(var, 'get') else ''
            if valor:
                extras_documento[normalizar_placeholder(chave)] = valor

        # Cálculo automático do período aquisitivo de férias pela admissão.
        # Se o usuário informar DATA_INICIO das férias, usa essa data como referência; senão usa a data atual.
        ref_aquisitivo = extras_documento.get('DATA_INICIO') or data_doc or date.today()
        aq_ini, aq_fim, aq_completo = calcular_periodo_aquisitivo(f[3], ref_aquisitivo)
        if not extras_documento.get('PERIODO_AQUISITIVO'):
            extras_documento['PERIODO_AQUISITIVO'] = aq_completo
        if not extras_documento.get('PERIODO_AQUISITIVO_INICIO'):
            extras_documento['PERIODO_AQUISITIVO_INICIO'] = aq_ini
        if not extras_documento.get('PERIODO_AQUISITIVO_FIM'):
            extras_documento['PERIODO_AQUISITIVO_FIM'] = aq_fim
        if not extras_documento.get('PERIODO_AQUISITIVO_COMPLETO'):
            extras_documento['PERIODO_AQUISITIVO_COMPLETO'] = aq_completo
        if not extras_documento.get('DATA_RETORNO') and extras_documento.get('DATA_FIM'):
            try:
                extras_documento['DATA_RETORNO'] = (parse_data_br(extras_documento.get('DATA_FIM')) + timedelta(days=1)).strftime('%d/%m/%Y')
            except Exception:
                pass
        if not extras_documento.get('DIAS_FERIAS') and extras_documento.get('DATA_INICIO') and extras_documento.get('DATA_FIM'):
            try:
                ini_f = parse_data_br(extras_documento.get('DATA_INICIO'))
                fim_f = parse_data_br(extras_documento.get('DATA_FIM'))
                extras_documento['DIAS_FERIAS'] = str((fim_f - ini_f).days + 1)
            except Exception:
                pass
        if not extras_documento.get('PERIODO') and extras_documento.get('DATA_INICIO') and extras_documento.get('DATA_FIM'):
            extras_documento['PERIODO'] = f"{extras_documento.get('DATA_INICIO')} a {extras_documento.get('DATA_FIM')}"

        valores = {
            'EMPRESA': nome_empresa,
            'CNPJ': e[1] if e and e[1] else DEFAULT_EMPRESA['cnpj'],
            'ENDERECO': endereco,
            'CIDADE': (e[5] if e and e[5] else DEFAULT_EMPRESA['cidade']),
            'UF': (e[6] if e and e[6] else DEFAULT_EMPRESA['uf']),
            'FUNCIONARIO': f[0] or '',
            'NOME': f[0] or '',
            'CPF': f[1] or '',
            'CTPS': f[2] or '',
            'ADMISSAO': adm,
            'DATA_ADMISSAO': adm,
            'FUNCAO': f[4] or '',
            'CARGO': f[4] or '',
            'SETOR': f[5] or '',
            'CBO': f[6] or '',
            'SALARIO': salario_fmt,
            'JORNADA': f[8] or '',
            'HORARIO_TRABALHO': f[9] or '',
            'DESCANSO': f[10] or '',
            'SABADO': f[11] or '',
            'DATA': data_doc,
            'DATA_DOCUMENTO': data_doc,
            'TITULO': titulo,
            'OBSERVACAO': obs,
            'OBSERVAÇÃO': obs,
            'MOTIVO': obs or titulo,
            'PERIODO': obs or titulo,
            'RESPONSAVEL': nome_empresa,
            'LOCAL_DATA': f"{(e[5] if e and e[5] else DEFAULT_EMPRESA['cidade']).title()}/{(e[6] if e and e[6] else DEFAULT_EMPRESA['uf'])}, {data_doc}",
        }
        valores.update(extras_documento)
        # Sinônimos úteis para modelos antigos ou personalizados.
        if 'DATA_INICIO' in valores: valores.setdefault('INICIO', valores.get('DATA_INICIO'))
        if 'DATA_FIM' in valores: valores.setdefault('TERMINO', valores.get('DATA_FIM'))
        if 'DESCRICAO_FATO' in valores: valores.setdefault('DESCRICAO', valores.get('DESCRICAO_FATO'))
        if 'VALOR' in valores: valores.setdefault('VALOR_RECEBIDO', valores.get('VALOR'))
        if 'EPI_ITEM' in valores: valores.setdefault('ITEM', valores.get('EPI_ITEM'))
        if 'UNIFORME_ITEM' in valores: valores.setdefault('ITEM_UNIFORME', valores.get('UNIFORME_ITEM'))
        return valores

    def _campo_docx_destacado(self, chave, valor):
        """Define quais campos preenchidos pelo sistema devem sair em negrito no Word."""
        chave = normalizar_placeholder(chave)
        destaque_chaves = {
            'FUNCIONARIO','NOME','MATRICULA','CPF','CTPS','RG','CARGO','FUNCAO','SETOR','JORNADA','EMPRESA',
            'ADMISSAO','DATA_ADMISSAO','DATA','DATA_DOCUMENTO','DATA_INICIO','DATA_FIM','INICIO','TERMINO','DATA_RETORNO',
            'PERIODO','PERIODO_AQUISITIVO','PERIODO_AQUISITIVO_INICIO','PERIODO_AQUISITIVO_FIM','PERIODO_AQUISITIVO_COMPLETO',
            'PERIODO_CONCESSIVO','PERIODO_CONCESSIVO_FIM','DIAS_FERIAS','DIAS_GOZADOS','DIAS_ABONO','DIAS_RESTANTES',
            'SALARIO','SALARIO_BASE','MEDIA_VARIAVEIS','VALOR','VALOR_FERIAS','VALOR_UM_TERCO','VALOR_ABONO',
            'ADIANTAMENTO_13','VALOR_13','TOTAL_BRUTO','INSS_ESTIMADO','IRRF_ESTIMADO','LIQUIDO_ESTIMADO'
        }
        if chave in destaque_chaves:
            return True
        texto = str(valor or '')
        # Datas, períodos e valores monetários também são destacados, mesmo em campos personalizados.
        if re.search(r'\b\d{2}/\d{2}/\d{4}\b', texto):
            return True
        if 'R$' in texto:
            return True
        return False

    def _substituir_placeholders_docx(self, caminho_docx, valores):
        try:
            from docx import Document
        except Exception as exc:
            raise RuntimeError('A biblioteca python-docx não está instalada. Execute INSTALAR_BIBLIOTECAS.bat novamente.') from exc

        # Aceita {{CAMPO}} e {CAMPO}. O valor preenchido pode receber negrito automaticamente.
        pattern = re.compile(r'(\{\{([A-Za-z0-9_ÁÉÍÓÚÀÈÌÒÙÃÕÂÊÎÔÛÇáéíóúàèìòùãõâêîôûç ]+)\}\}|\{([A-Za-z0-9_ÁÉÍÓÚÀÈÌÒÙÃÕÂÊÎÔÛÇáéíóúàèìòùãõâêîôûç ]+)\})')

        def tratar_paragrafo(paragraph):
            original = paragraph.text
            if not original or ('{' not in original):
                return
            partes = []
            pos = 0
            mudou = False
            for m in pattern.finditer(original):
                if m.start() > pos:
                    partes.append(('texto', original[pos:m.start()], False))
                chave = normalizar_placeholder(m.group(2) or m.group(3) or '')
                if chave in valores:
                    valor = str(valores.get(chave) or '')
                    partes.append(('valor', valor, self._campo_docx_destacado(chave, valor)))
                    mudou = True
                else:
                    partes.append(('texto', m.group(0), False))
                pos = m.end()
            if pos < len(original):
                partes.append(('texto', original[pos:], False))
            if not mudou:
                return

            # Mantém o estilo do parágrafo e recria os runs para destacar só os valores preenchidos.
            for run in paragraph.runs:
                run.text = ''
            for _, texto, negrito in partes:
                if texto == '':
                    continue
                r = paragraph.add_run(texto)
                r.bold = bool(negrito)

        doc = Document(caminho_docx)
        for paragraph in doc.paragraphs:
            tratar_paragrafo(paragraph)
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        tratar_paragrafo(paragraph)
        for section in doc.sections:
            for container in (section.header, section.footer):
                for paragraph in container.paragraphs:
                    tratar_paragrafo(paragraph)
                for table in container.tables:
                    for row in table.rows:
                        for cell in row.cells:
                            for paragraph in cell.paragraphs:
                                tratar_paragrafo(paragraph)
        doc.save(caminho_docx)

    def gerar_copia_modelo_documento(self):
        fid = self._func_id_from_option(self.doc_func.get())
        if not fid:
            messagebox.showwarning('Documentos','Selecione um funcionário.'); return
        dados = self._dados_documento_funcionario(fid)
        if not dados:
            messagebox.showwarning('Documentos','Funcionário não encontrado.'); return
        tipo = self.doc_tipo.get() or 'Declaração'
        modelo = os.path.join(MODELOS_DIR, self._modelo_por_tipo(tipo))
        if not os.path.exists(modelo):
            messagebox.showerror('Documentos','Modelo não encontrado na pasta modelos_documentos.'); return
        seguro = self._safe_filename(dados.get('FUNCIONARIO'), 60)
        tipo_seguro = self._safe_filename(tipo, 30)
        data_arq = datetime.now().strftime('%Y%m%d_%H%M%S')
        destino = os.path.join(DOCS_GERADOS_DIR, f'{tipo_seguro}_{seguro}_{data_arq}.docx')
        shutil.copy2(modelo, destino)
        try:
            self._substituir_placeholders_docx(destino, dados)
            extras = []
            for chave,var in getattr(self,'doc_extra_vars',{}).items():
                val = var.get().strip()
                if val:
                    extras.append(f'{chave}: {val}')
            obs = self.doc_obs.get('1.0','end').strip()
            obs_completa = (obs + ('\n' if obs and extras else '') + '\n'.join(extras)).strip()
            with con() as db:
                db.execute("""INSERT INTO documentos_rh(funcionario_id,tipo,data,titulo,observacao,arquivo,ativo)
                              VALUES(?,?,?,?,?,?,1)""", (fid, tipo, parse_data_br(self.doc_data.get()).isoformat(), self.doc_titulo.get(), obs_completa, destino))
            log_action(self.usuario,'DOCUMENTOS',f'Documento preenchido gerado: {os.path.basename(destino)}')
            self.carregar_documentos()
            messagebox.showinfo('Documentos', f'Documento preenchido gerado em documentos_gerados:\n{os.path.basename(destino)}')
            try:
                if os.name=='nt': os.startfile(destino)
            except Exception:
                pass
        except Exception as exc:
            messagebox.showerror('Erro ao gerar documento', str(exc))
            try:
                if os.path.exists(destino): os.remove(destino)
            except Exception:
                pass

    def build_epis(self):
        f=self.tab_epis; ttk.Label(f,text='Controle de EPIs',style='Title.TLabel').grid(row=0,column=0,columnspan=6,sticky='w',padx=16,pady=16)
        self.epi_func=tk.StringVar(); self.epi_item=tk.StringVar(); self.epi_ca=tk.StringVar(); self.epi_entrega=tk.StringVar(value=date.today().strftime('%d/%m/%Y')); self.epi_validade=tk.StringVar()
        vals=[('Funcionário',self.epi_func,self._funcionario_options(),0),('Item',self.epi_item,None,1),('CA',self.epi_ca,None,2),('Entrega',self.epi_entrega,None,3),('Validade',self.epi_validade,None,4)]
        for lab,var,values,col in vals:
            ttk.Label(f,text=lab).grid(row=1,column=col,sticky='w',padx=8); (ttk.Combobox(f,textvariable=var,values=values,width=30) if values else ttk.Entry(f,textvariable=var)).grid(row=2,column=col,sticky='ew',padx=8)
        ttk.Button(f,text='Salvar EPI',command=self.salvar_epi).grid(row=2,column=5,sticky='ew',padx=8)
        self.epi_obs=tk.Text(f,height=3); self.epi_obs.grid(row=3,column=0,columnspan=6,sticky='ew',padx=8,pady=8)
        self.epi_tree=ttk.Treeview(f,columns=('func','item','ca','entrega','validade','obs'),show='headings',height=13)
        for col,txt,w in [('func','Funcionário',240),('item','Item',150),('ca','CA',80),('entrega','Entrega',90),('validade','Validade',90),('obs','Observação',260)]: self.epi_tree.heading(col,text=txt); self.epi_tree.column(col,width=w)
        self.epi_tree.grid(row=4,column=0,columnspan=6,sticky='nsew',padx=8,pady=8); f.rowconfigure(4,weight=1)

    def salvar_epi(self):
        fid=self._func_id_from_option(self.epi_func.get())
        if not fid or not self.epi_item.get().strip(): messagebox.showwarning('EPIs','Informe funcionário e item.'); return
        validade=parse_data_br(self.epi_validade.get()).isoformat() if self.epi_validade.get().strip() else ''
        with con() as db: db.execute('INSERT INTO epis(funcionario_id,item,ca,data_entrega,validade,observacao,ativo) VALUES(?,?,?,?,?,?,1)',(fid,self.epi_item.get().upper(),self.epi_ca.get(),parse_data_br(self.epi_entrega.get()).isoformat(),validade,self.epi_obs.get('1.0','end').strip()))
        log_action(self.usuario,'EPIS','EPI cadastrado'); self.carregar_epis(); messagebox.showinfo('EPIs','EPI salvo.')

    def carregar_epis(self):
        if not hasattr(self,'epi_tree'): return
        self.epi_tree.delete(*self.epi_tree.get_children())
        for r in con().execute("SELECT f.nome,e.item,e.ca,e.data_entrega,e.validade,e.observacao FROM epis e LEFT JOIN funcionarios f ON f.id=e.funcionario_id WHERE e.ativo=1 ORDER BY e.data_entrega DESC LIMIT 300"):
            self.epi_tree.insert('', 'end', values=(r[0],r[1],r[2],fmt_data(r[3]),fmt_data(r[4]),r[5]))

    def build_exames(self):
        f=self.tab_exames; ttk.Label(f,text='Controle de Exames / ASO',style='Title.TLabel').grid(row=0,column=0,columnspan=6,sticky='w',padx=16,pady=16)
        self.ex_func=tk.StringVar(); self.ex_tipo=tk.StringVar(value='Periódico'); self.ex_data=tk.StringVar(value=date.today().strftime('%d/%m/%Y')); self.ex_validade=tk.StringVar(); self.ex_clinica=tk.StringVar()
        for col,(lab,var,vals) in enumerate([('Funcionário',self.ex_func,self._funcionario_options()),('Tipo',self.ex_tipo,['Admissional','Periódico','Demissional','Retorno ao Trabalho','Mudança de Função']),('Data',self.ex_data,None),('Validade',self.ex_validade,None),('Clínica',self.ex_clinica,None)]):
            ttk.Label(f,text=lab).grid(row=1,column=col,sticky='w',padx=8); (ttk.Combobox(f,textvariable=var,values=vals,width=28) if vals else ttk.Entry(f,textvariable=var)).grid(row=2,column=col,sticky='ew',padx=8)
        ttk.Button(f,text='Salvar exame',command=self.salvar_exame).grid(row=2,column=5,sticky='ew',padx=8)
        self.ex_obs=tk.Text(f,height=3); self.ex_obs.grid(row=3,column=0,columnspan=6,sticky='ew',padx=8,pady=8)
        self.ex_tree=ttk.Treeview(f,columns=('func','tipo','data','validade','clinica','obs'),show='headings',height=13)
        for col,txt,w in [('func','Funcionário',230),('tipo','Tipo',150),('data','Data',90),('validade','Validade',90),('clinica','Clínica',150),('obs','Observação',250)]: self.ex_tree.heading(col,text=txt); self.ex_tree.column(col,width=w)
        self.ex_tree.grid(row=4,column=0,columnspan=6,sticky='nsew',padx=8,pady=8); f.rowconfigure(4,weight=1)

    def salvar_exame(self):
        fid=self._func_id_from_option(self.ex_func.get())
        if not fid: messagebox.showwarning('Exames','Selecione funcionário.'); return
        validade=parse_data_br(self.ex_validade.get()).isoformat() if self.ex_validade.get().strip() else ''
        with con() as db: db.execute('INSERT INTO exames(funcionario_id,tipo,data_exame,validade,clinica,observacao,ativo) VALUES(?,?,?,?,?,?,1)',(fid,self.ex_tipo.get(),parse_data_br(self.ex_data.get()).isoformat(),validade,self.ex_clinica.get(),self.ex_obs.get('1.0','end').strip()))
        log_action(self.usuario,'EXAMES','Exame cadastrado'); self.carregar_exames(); messagebox.showinfo('Exames','Exame salvo.')

    def carregar_exames(self):
        if not hasattr(self,'ex_tree'): return
        self.ex_tree.delete(*self.ex_tree.get_children())
        for r in con().execute("SELECT f.nome,e.tipo,e.data_exame,e.validade,e.clinica,e.observacao FROM exames e LEFT JOIN funcionarios f ON f.id=e.funcionario_id WHERE e.ativo=1 ORDER BY e.data_exame DESC LIMIT 300"):
            self.ex_tree.insert('', 'end', values=(r[0],r[1],fmt_data(r[2]),fmt_data(r[3]),r[4],r[5]))

    def build_agenda(self):
        f=self.tab_agenda; ttk.Label(f,text='Agenda RH e Alertas',style='Title.TLabel').grid(row=0,column=0,columnspan=5,sticky='w',padx=16,pady=16)
        self.ag_data=tk.StringVar(value=date.today().strftime('%d/%m/%Y')); self.ag_titulo=tk.StringVar(); self.ag_tipo=tk.StringVar(value='Lembrete'); self.ag_func=tk.StringVar()
        for col,(lab,var,vals) in enumerate([('Data',self.ag_data,None),('Título',self.ag_titulo,None),('Tipo',self.ag_tipo,['Lembrete','Férias','Exame','Contrato','EPI','Retorno']),('Funcionário',self.ag_func,self._funcionario_options())]):
            ttk.Label(f,text=lab).grid(row=1,column=col,sticky='w',padx=8); (ttk.Combobox(f,textvariable=var,values=vals,width=28) if vals else ttk.Entry(f,textvariable=var)).grid(row=2,column=col,sticky='ew',padx=8)
        ttk.Button(f,text='Salvar lembrete',command=self.salvar_agenda).grid(row=2,column=4,sticky='ew',padx=8)
        self.ag_obs=tk.Text(f,height=3); self.ag_obs.grid(row=3,column=0,columnspan=5,sticky='ew',padx=8,pady=8)
        self.ag_tree=ttk.Treeview(f,columns=('data','tipo','titulo','func','obs','status'),show='headings',height=14)
        for col,txt,w in [('data','Data',90),('tipo','Tipo',100),('titulo','Título',220),('func','Funcionário',220),('obs','Observação',260),('status','Status',90)]: self.ag_tree.heading(col,text=txt); self.ag_tree.column(col,width=w)
        self.ag_tree.grid(row=4,column=0,columnspan=5,sticky='nsew',padx=8,pady=8); f.rowconfigure(4,weight=1)

    def salvar_agenda(self):
        fid=self._func_id_from_option(self.ag_func.get())
        with con() as db: db.execute('INSERT INTO agenda_rh(data,titulo,tipo,funcionario_id,observacao,concluido) VALUES(?,?,?,?,?,0)',(parse_data_br(self.ag_data.get()).isoformat(),self.ag_titulo.get(),self.ag_tipo.get(),fid,self.ag_obs.get('1.0','end').strip()))
        log_action(self.usuario,'AGENDA','Lembrete cadastrado'); self.carregar_agenda(); messagebox.showinfo('Agenda','Lembrete salvo.')

    def carregar_agenda(self):
        if not hasattr(self,'ag_tree'): return
        self.ag_tree.delete(*self.ag_tree.get_children())
        for r in con().execute("SELECT a.data,a.tipo,a.titulo,f.nome,a.observacao,a.concluido FROM agenda_rh a LEFT JOIN funcionarios f ON f.id=a.funcionario_id ORDER BY a.data ASC LIMIT 300"):
            self.ag_tree.insert('', 'end', values=(fmt_data(r[0]),r[1],r[2],r[3] or '',r[4],'Concluído' if r[5] else 'Pendente'))

    def build_central_pdfs(self):
        f=self.tab_central_pdfs; ttk.Label(f,text='Central de PDFs Gerados',style='Title.TLabel').grid(row=0,column=0,columnspan=4,sticky='w',padx=16,pady=16)
        ttk.Button(f,text='Atualizar lista',command=self.carregar_central_pdfs).grid(row=1,column=0,sticky='ew',padx=8,pady=8)
        ttk.Button(f,text='Abrir PDF selecionado',command=self.abrir_pdf_central).grid(row=1,column=1,sticky='ew',padx=8,pady=8)
        ttk.Button(f,text='Abrir pasta PDFs',command=self.open_pdfs).grid(row=1,column=2,sticky='ew',padx=8,pady=8)
        self.central_tree=ttk.Treeview(f,columns=('data','tipo','setor','mes','ano','qtd','arquivo'),show='headings',height=16)
        for col,txt,w in [('data','Data/Hora',140),('tipo','Tipo',130),('setor','Setor',110),('mes','Mês',50),('ano','Ano',60),('qtd','Qtd.',50),('arquivo','Arquivo',520)]: self.central_tree.heading(col,text=txt); self.central_tree.column(col,width=w)
        self.central_tree.grid(row=2,column=0,columnspan=4,sticky='nsew',padx=8,pady=8); f.rowconfigure(2,weight=1)

    def carregar_central_pdfs(self):
        if not hasattr(self,'central_tree'): return
        self.central_tree.delete(*self.central_tree.get_children())
        for r in con().execute('SELECT data_hora,tipo,setor,mes,ano,quantidade,arquivo FROM pdf_historico ORDER BY id DESC LIMIT 500'):
            self.central_tree.insert('', 'end', values=r)

    def abrir_pdf_central(self):
        sel=self.central_tree.selection()
        if not sel: return
        arq=self.central_tree.item(sel[0])['values'][6]
        if arq and os.path.exists(arq): webbrowser.open(arq)
        else: messagebox.showwarning('Central PDFs','Arquivo não localizado.')

    def build_assistente(self):
        f=self.tab_assistente; ttk.Label(f,text='Assistente RH Izzant',style='Title.TLabel').grid(row=0,column=0,columnspan=3,sticky='w',padx=16,pady=16)
        ttk.Label(f,text='Perguntas prontas para consulta rápida:').grid(row=1,column=0,sticky='w',padx=8)
        self.ai_combo=ttk.Combobox(f,values=['Quem está de férias hoje?','Funcionários ativos por setor','ASO/exames vencendo em 30 dias','EPIs vencendo em 30 dias','Ocorrências do mês','PDFs gerados no mês'],width=45)
        self.ai_combo.grid(row=2,column=0,sticky='ew',padx=8,pady=8); self.ai_combo.set('Funcionários ativos por setor')
        ttk.Button(f,text='Consultar',command=self.executar_assistente).grid(row=2,column=1,sticky='ew',padx=8)
        self.ai_text=tk.Text(f,height=22); self.ai_text.grid(row=3,column=0,columnspan=3,sticky='nsew',padx=8,pady=8); f.rowconfigure(3,weight=1); f.columnconfigure(0,weight=1)

    def executar_assistente(self):
        q=self.ai_combo.get(); hoje=date.today().isoformat(); out=[]
        with con() as db:
            if 'férias hoje' in q.lower():
                rows=db.execute('SELECT f.nome,fc.inicio,fc.fim FROM ferias_controle fc JOIN funcionarios f ON f.id=fc.funcionario_id WHERE fc.ativo=1 AND fc.inicio<=? AND fc.fim>=? ORDER BY f.nome',(hoje,hoje)).fetchall(); out=[f'{r[0]} - {fmt_data(r[1])} a {fmt_data(r[2])}' for r in rows] or ['Nenhum funcionário em férias hoje.']
            elif 'ativos por setor' in q.lower():
                rows=db.execute('SELECT setor,COUNT(*) FROM funcionarios WHERE ativo=1 GROUP BY setor ORDER BY setor').fetchall(); out=[f'{r[0]}: {r[1]}' for r in rows]
            elif 'exames' in q.lower():
                limite=(date.today()+timedelta(days=30)).isoformat(); rows=db.execute("SELECT f.nome,e.tipo,e.validade FROM exames e JOIN funcionarios f ON f.id=e.funcionario_id WHERE e.ativo=1 AND e.validade<>'' AND e.validade<=? ORDER BY e.validade",(limite,)).fetchall(); out=[f'{fmt_data(r[2])} - {r[0]} - {r[1]}' for r in rows] or ['Nenhum exame vencendo em 30 dias.']
            elif 'epis' in q.lower():
                limite=(date.today()+timedelta(days=30)).isoformat(); rows=db.execute("SELECT f.nome,e.item,e.validade FROM epis e JOIN funcionarios f ON f.id=e.funcionario_id WHERE e.ativo=1 AND e.validade<>'' AND e.validade<=? ORDER BY e.validade",(limite,)).fetchall(); out=[f'{fmt_data(r[2])} - {r[0]} - {r[1]}' for r in rows] or ['Nenhum EPI vencendo em 30 dias.']
            elif 'ocorrências' in q.lower():
                ano=date.today().year; mes=date.today().month; ini=date(ano,mes,1).isoformat(); fim=date(ano,mes,calendar.monthrange(ano,mes)[1]).isoformat(); rows=db.execute('SELECT o.tipo,COUNT(*) FROM ocorrencias o WHERE o.ativo=1 AND o.data_inicio<=? AND o.data_fim>=? GROUP BY o.tipo',(fim,ini)).fetchall(); out=[f'{r[0]}: {r[1]}' for r in rows] or ['Sem ocorrências no mês.']
            else:
                rows=db.execute('SELECT tipo,COUNT(*) FROM pdf_historico WHERE mes=? AND ano=? GROUP BY tipo',(date.today().month,date.today().year)).fetchall(); out=[f'{r[0]}: {r[1]}' for r in rows] or ['Nenhum PDF registrado no mês.']
        self.ai_text.delete('1.0','end'); self.ai_text.insert('end','\n'.join(out))

    def build_backup(self):
        f=self.tab_backup
        ttk.Label(f,text='Backup e Manutenção',style='Title.TLabel').grid(row=0,column=0,columnspan=4,sticky='w',padx=16,pady=16)
        ttk.Button(f,text='Fazer backup agora',command=self.backup_now).grid(row=1,column=0,sticky='ew',padx=12,pady=8)
        ttk.Button(f,text='Abrir pasta backups',command=self.open_backups).grid(row=1,column=1,sticky='ew',padx=12,pady=8)
        ttk.Button(f,text='Abrir pasta dados',command=self.open_data).grid(row=1,column=2,sticky='ew',padx=12,pady=8)
        ttk.Button(f,text='Restaurar backup selecionado',command=self.restaurar_backup).grid(row=1,column=4,sticky='ew',padx=12,pady=8)
        ttk.Button(f,text='Atualizar painel',command=self.refresh_dashboard).grid(row=1,column=3,sticky='ew',padx=12,pady=8)
        self.backup_list=tk.Listbox(f, height=18)
        self.backup_list.grid(row=2,column=0,columnspan=4,sticky='nsew',padx=12,pady=12)
        f.rowconfigure(2,weight=1)
        for c in range(4): f.columnconfigure(c,weight=1)

    def build_relatorios(self):
        f=self.tab_rel
        ttk.Label(f,text='Relatórios',style='Title.TLabel').grid(row=0,column=0,columnspan=4,sticky='w',padx=16,pady=16)
        ttk.Button(f,text='Funcionários ativos por setor',command=self.rel_funcionarios_setor).grid(row=1,column=0,sticky='ew',padx=12,pady=8)
        ttk.Button(f,text='Funcionários inativos',command=self.rel_inativos).grid(row=1,column=1,sticky='ew',padx=12,pady=8)
        ttk.Button(f,text='Abrir pasta relatórios',command=self.open_relatorios).grid(row=1,column=2,sticky='ew',padx=12,pady=8)
        self.rel_info=ttk.Label(f,text='Os relatórios são salvos em CSV, compatíveis com Excel.',font=('Arial',11))
        self.rel_info.grid(row=2,column=0,columnspan=4,sticky='w',padx=16,pady=8)
        for c in range(4): f.columnconfigure(c,weight=1)

    def build_logs(self):
        f=self.tab_logs
        ttk.Label(f,text='Auditoria do Sistema',style='Title.TLabel').grid(row=0,column=0,columnspan=4,sticky='w',padx=16,pady=16)
        ttk.Button(f,text='Atualizar auditoria',command=self.carregar_logs).grid(row=1,column=0,sticky='ew',padx=12,pady=8)
        ttk.Button(f,text='Exportar auditoria CSV',command=self.exportar_logs).grid(row=1,column=1,sticky='ew',padx=12,pady=8)
        self.log_tree=ttk.Treeview(f, columns=('data','usuario','acao','detalhes'), show='headings')
        for col,txt_h,w in [('data','Data/Hora',150),('usuario','Usuário',110),('acao','Ação',150),('detalhes','Detalhes',520)]:
            self.log_tree.heading(col,text=txt_h); self.log_tree.column(col,width=w)
        self.log_tree.grid(row=2,column=0,columnspan=4,sticky='nsew',padx=12,pady=8)
        f.rowconfigure(2,weight=1)
        for c in range(4): f.columnconfigure(c,weight=1)

    def build_usuarios(self):
        f=self.tab_usuarios
        ttk.Label(f,text='Usuários e Segurança',style='Title.TLabel').grid(row=0,column=0,columnspan=4,sticky='w',padx=16,pady=16)
        ttk.Label(f,text='Usuário').grid(row=1,column=0,sticky='w',padx=12,pady=6)
        self.u_usuario=tk.StringVar(); ttk.Entry(f,textvariable=self.u_usuario,width=30).grid(row=1,column=1,sticky='ew',padx=8,pady=6)
        ttk.Label(f,text='Senha').grid(row=2,column=0,sticky='w',padx=12,pady=6)
        self.u_senha=tk.StringVar(); ttk.Entry(f,textvariable=self.u_senha,width=30,show='*').grid(row=2,column=1,sticky='ew',padx=8,pady=6)
        ttk.Label(f,text='Perfil').grid(row=3,column=0,sticky='w',padx=12,pady=6)
        self.u_perfil=tk.StringVar(value='RH'); ttk.Combobox(f,textvariable=self.u_perfil,values=['Administrador','RH','Consulta'],state='readonly').grid(row=3,column=1,sticky='ew',padx=8,pady=6)
        ttk.Button(f,text='Criar/Atualizar usuário',command=self.salvar_usuario).grid(row=4,column=1,sticky='ew',padx=8,pady=10)
        ttk.Button(f,text='Atualizar lista',command=self.carregar_usuarios).grid(row=4,column=2,sticky='ew',padx=8,pady=10)
        self.user_tree=ttk.Treeview(f, columns=('usuario','perfil','ativo'), show='headings')
        for col,txt_h,w in [('usuario','Usuário',180),('perfil','Perfil',150),('ativo','Ativo',80)]:
            self.user_tree.heading(col,text=txt_h); self.user_tree.column(col,width=w)
        self.user_tree.grid(row=5,column=0,columnspan=4,sticky='nsew',padx=12,pady=8)
        f.rowconfigure(5,weight=1)
        for c in range(4): f.columnconfigure(c,weight=1)

    def rel_funcionarios_setor(self):
        os.makedirs(RELATORIO_DIR, exist_ok=True)
        path=os.path.join(RELATORIO_DIR, 'funcionarios_ativos_por_setor.csv')
        funcs=get_funcionarios(True)
        with open(path,'w',newline='',encoding='utf-8-sig') as fp:
            wr=csv.writer(fp, delimiter=';')
            wr.writerow(['Setor','Nome','Função','Admissão','Salário'])
            for f in funcs:
                wr.writerow([f.get('setor') or 'GERAL', f.get('nome'), f.get('funcao'), fmt_data(f.get('admissao')), money(f.get('salario'))])
        log_action(self.usuario,'RELATÓRIO','Funcionários ativos por setor')
        self.rel_info.config(text='Relatório salvo em: '+path)
        messagebox.showinfo('Relatório', 'Relatório salvo em:\n'+path)

    def rel_inativos(self):
        os.makedirs(RELATORIO_DIR, exist_ok=True)
        path=os.path.join(RELATORIO_DIR, 'funcionarios_inativos.csv')
        funcs=[f for f in get_funcionarios(False) if not f.get('ativo')]
        with open(path,'w',newline='',encoding='utf-8-sig') as fp:
            wr=csv.writer(fp, delimiter=';')
            wr.writerow(['Setor','Nome','Função','Admissão','Salário'])
            for f in funcs:
                wr.writerow([f.get('setor') or 'GERAL', f.get('nome'), f.get('funcao'), fmt_data(f.get('admissao')), money(f.get('salario'))])
        log_action(self.usuario,'RELATÓRIO','Funcionários inativos')
        self.rel_info.config(text='Relatório salvo em: '+path)
        messagebox.showinfo('Relatório', 'Relatório salvo em:\n'+path)

    def carregar_logs(self):
        if not hasattr(self,'log_tree'): return
        for i in self.log_tree.get_children(): self.log_tree.delete(i)
        with con() as db:
            rows=db.execute('SELECT data_hora,usuario,acao,detalhes FROM logs ORDER BY id DESC LIMIT 300').fetchall()
        for r in rows: self.log_tree.insert('', 'end', values=r)

    def exportar_logs(self):
        os.makedirs(RELATORIO_DIR, exist_ok=True)
        path=os.path.join(RELATORIO_DIR, 'auditoria.csv')
        with con() as db:
            rows=db.execute('SELECT data_hora,usuario,acao,detalhes FROM logs ORDER BY id DESC').fetchall()
        with open(path,'w',newline='',encoding='utf-8-sig') as fp:
            wr=csv.writer(fp,delimiter=';'); wr.writerow(['Data/Hora','Usuário','Ação','Detalhes']); wr.writerows(rows)
        messagebox.showinfo('Auditoria','Auditoria exportada em:\n'+path)

    def carregar_usuarios(self):
        if not hasattr(self,'user_tree'): return
        for i in self.user_tree.get_children(): self.user_tree.delete(i)
        with con() as db:
            rows=db.execute('SELECT usuario,perfil,ativo FROM usuarios ORDER BY usuario').fetchall()
        for usuario,perfil,ativo in rows:
            self.user_tree.insert('', 'end', values=(usuario,perfil,'Sim' if ativo else 'Não'))

    def salvar_usuario(self):
        usuario=self.u_usuario.get().strip()
        senha=self.u_senha.get().strip()
        perfil=self.u_perfil.get().strip() or 'RH'
        if not usuario or not senha:
            messagebox.showwarning('Atenção','Informe usuário e senha.'); return
        with con() as db:
            db.execute('INSERT INTO usuarios(usuario,senha,perfil,ativo) VALUES(?,?,?,1) ON CONFLICT(usuario) DO UPDATE SET senha=excluded.senha, perfil=excluded.perfil, ativo=1', (usuario,senha,perfil))
        log_action(self.usuario,'USUÁRIO',f'Criado/atualizado: {usuario}')
        self.u_usuario.set(''); self.u_senha.set('')
        self.carregar_usuarios()
        messagebox.showinfo('Usuário','Usuário salvo.')

    def load_all(self):
        emp=get_empresa()
        for k,v in self.emp_vars.items(): v.set(emp.get(k,''))
        self.populate_setores_tree(); self.populate_jornadas_tree(); self.populate_escalas_tree(); self.populate_feriados_tree(); self.populate_tree(); self.populate_combo(); self.populate_setor_combo(); self.populate_jornada_combo(); self.populate_ocorrencias_combo(); self.refresh_dashboard(); self.carregar_logs(); self.carregar_usuarios(); self.carregar_historico_pdf(); self.carregar_ferias(); self.carregar_banco_horas(); self.carregar_documentos(); self.carregar_epis(); self.carregar_exames(); self.carregar_agenda(); self.carregar_central_pdfs()

    def refresh_dashboard(self):
        try:
            self.card_total.config(text=str(len(get_funcionarios(True))))
            self.card_mes.config(text=MESES[datetime.now().month-1])
            b=[x for x in os.listdir(BACKUP_DIR) if x.endswith('.db')] if os.path.exists(BACKUP_DIR) else []
            self.card_backup.config(text=str(len(b)))
            with con() as db:
                self.card_setores.config(text=str(db.execute('SELECT COUNT(*) FROM setores WHERE ativo=1').fetchone()[0]))
                self.card_jornadas.config(text=str(db.execute('SELECT COUNT(*) FROM jornadas WHERE ativo=1').fetchone()[0]))
                self.card_feriados.config(text=str(db.execute('SELECT COUNT(*) FROM feriados WHERE ativo=1').fetchone()[0]))
                ym=datetime.now().strftime('%Y-%m')
                self.card_ocorrencias.config(text=str(db.execute("SELECT COUNT(*) FROM ocorrencias WHERE ativo=1 AND substr(data_inicio,1,7)<=? AND substr(data_fim,1,7)>=?",(ym,ym)).fetchone()[0]))
            self.backup_list.delete(0,'end')
            for name in sorted(b, reverse=True)[:200]: self.backup_list.insert('end', name)
        except Exception:
            pass

    def save_empresa(self):
        vals={k:v.get().strip().upper() if k!='cnpj' else v.get().strip() for k,v in self.emp_vars.items()}
        with con() as db:
            db.execute("""UPDATE empresa SET nome=?,cnpj=?,endereco=?,numero=?,bairro=?,cidade=?,uf=? WHERE id=1""",
                       (vals['nome'],vals['cnpj'],vals['endereco'],vals['numero'],vals['bairro'],vals['cidade'],vals['uf']))
        self.refresh_dashboard()
        log_action(self.usuario,'EMPRESA','Dados da empresa atualizados')
        messagebox.showinfo('Salvo','Dados da empresa salvos automaticamente.')

    def clear_func(self):
        self.selected_id=None
        for v in self.fvars.values(): v.set('')
        self.fvars['setor'].set('GERAL')
        self.fvars['jornada'].set((get_jornadas(True)[0]['nome'] if get_jornadas(True) else 'HORÁRIO DE TRABALHO DE SEGUNDA A SEXTA-FEIRA'))
        self.fvars['horario_trabalho'].set('')
        self.fvars['descanso'].set('REMUNERADO')
        self.fvars['sabado'].set('COMPENSADO')
        self.aplicar_jornada_funcionario()

    def aplicar_jornada_funcionario(self):
        """Atualiza automaticamente os campos Sábado e Descanso conforme a jornada selecionada.
        Especialmente para 12x36, sábado e DSR devem seguir a escala, não o padrão compensado.
        """
        if not hasattr(self, 'fvars') or 'jornada' not in self.fvars:
            return
        nome_jornada = (self.fvars['jornada'].get() or '').strip()
        j = normalizar_jornada_12x36(get_jornada_por_nome(nome_jornada))
        if not j:
            if '12X36' in nome_jornada.upper().replace(' ', ''):
                self.fvars['sabado'].set('CONFORME ESCALA')
                self.fvars['descanso'].set('CONFORME ESCALA 12X36')
            return
        horario = texto_horario_jornada(j)
        if 'horario_trabalho' in self.fvars:
            self.fvars['horario_trabalho'].set(horario)
        if hasattr(self, 'combo_horario_func'):
            horarios = sorted(set([horario] + [texto_horario_jornada(x) for x in get_jornadas(True) if texto_horario_jornada(x)]))
            self.combo_horario_func['values'] = horarios
        self.fvars['sabado'].set((j.get('sabado_tratamento') or 'COMPENSADO').upper())
        self.fvars['descanso'].set((j.get('descanso_semanal') or 'DOMINGO').upper())

    def valores_jornada_para_funcionario(self, nome_jornada, descanso_digitado='', sabado_digitado=''):
        """Retorna descanso e sábado coerentes com a jornada cadastrada.
        Garante que funcionários vinculados a jornada 12x36 recebam CONFORME ESCALA.
        """
        j = normalizar_jornada_12x36(get_jornada_por_nome(nome_jornada))
        if j:
            return ((j.get('descanso_semanal') or descanso_digitado or 'DOMINGO').upper(),
                    (j.get('sabado_tratamento') or sabado_digitado or 'COMPENSADO').upper())
        if '12X36' in (nome_jornada or '').upper().replace(' ', ''):
            return ('CONFORME ESCALA 12X36', 'CONFORME ESCALA')
        return ((descanso_digitado or 'REMUNERADO').upper(), (sabado_digitado or 'COMPENSADO').upper())

    def save_func(self):
        vals={k:v.get().strip() for k,v in self.fvars.items()}
        if not vals['nome']:
            messagebox.showwarning('Atenção','Informe o nome do funcionário.'); return
        try: sal=float(vals['salario'].replace('.','').replace(',','.')) if vals['salario'] else 0
        except Exception: sal=0
        descanso_calc, sabado_calc = self.valores_jornada_para_funcionario(vals.get('jornada'), vals.get('descanso'), vals.get('sabado'))
        vals['descanso'] = descanso_calc
        vals['sabado'] = sabado_calc
        if not vals.get('horario_trabalho'):
            j_tmp = normalizar_jornada_12x36(get_jornada_por_nome(vals.get('jornada')))
            vals['horario_trabalho'] = texto_horario_jornada(j_tmp) if j_tmp else ''
        with con() as db:
            db.execute("INSERT OR IGNORE INTO setores(nome,descricao,ativo) VALUES(?,?,1)", ((vals.get('setor') or 'GERAL').upper(), 'Criado pelo cadastro de funcionários'))
            db.execute("UPDATE setores SET ativo=1 WHERE nome=?", ((vals.get('setor') or 'GERAL').upper(),))
            if self.selected_id:
                db.execute("""UPDATE funcionarios SET nome=?,cpf=?,ctps=?,admissao=?,funcao=?,setor=?,cbo=?,salario=?,jornada=?,horario_trabalho=?,descanso=?,sabado=?,ativo=1 WHERE id=?""",
                           (vals['nome'].upper(),vals['cpf'],vals['ctps'],vals['admissao'],vals['funcao'].upper(),(vals.get('setor') or 'GERAL').upper(),vals['cbo'],sal,vals['jornada'].upper(),vals.get('horario_trabalho',''),vals['descanso'].upper(),vals['sabado'].upper(),self.selected_id))
            else:
                db.execute("""INSERT INTO funcionarios(nome,cpf,ctps,admissao,funcao,setor,cbo,salario,jornada,horario_trabalho,descanso,sabado,ativo) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,1)""",
                           (vals['nome'].upper(),vals['cpf'],vals['ctps'],vals['admissao'],vals['funcao'].upper(),(vals.get('setor') or 'GERAL').upper(),vals['cbo'],sal,vals['jornada'].upper(),vals.get('horario_trabalho',''),vals['descanso'].upper(),vals['sabado'].upper()))
        self.clear_func(); self.populate_setores_tree(); self.populate_jornadas_tree(); self.populate_escalas_tree(); self.populate_feriados_tree(); self.populate_tree(); self.populate_combo(); self.populate_setor_combo(); self.populate_jornada_combo(); self.populate_ocorrencias_combo(); self.refresh_dashboard()
        log_action(self.usuario,'FUNCIONÁRIO','Cadastro salvo/atualizado: '+vals['nome'].upper())
        messagebox.showinfo('Salvo','Funcionário salvo automaticamente no banco local.')

    def atualizar_funcionarios_conforme_jornada(self):
        """Recalcula campos Sábado e Descanso de todos os funcionários ativos conforme a jornada vinculada.
        Corrige cadastros antigos, inclusive 12x36, sem alterar o layout da folha.
        """
        corrigidos = 0
        with con() as db:
            rows = db.execute("SELECT id,jornada,COALESCE(horario_trabalho,''),descanso,sabado FROM funcionarios WHERE ativo=1").fetchall()
            for fid, jornada, horario_atual, descanso, sabado in rows:
                novo_descanso, novo_sabado = self.valores_jornada_para_funcionario(jornada, descanso, sabado)
                j_tmp = normalizar_jornada_12x36(get_jornada_por_nome(jornada))
                novo_horario = texto_horario_jornada(j_tmp) if j_tmp else (horario_atual or '')
                if (novo_descanso != (descanso or '').upper()) or (novo_sabado != (sabado or '').upper()) or (novo_horario != (horario_atual or '')):
                    db.execute('UPDATE funcionarios SET horario_trabalho=?, descanso=?, sabado=? WHERE id=?', (novo_horario, novo_descanso, novo_sabado, fid))
                    corrigidos += 1
        self.populate_tree(); self.populate_combo(); self.refresh_dashboard()
        log_action(self.usuario,'FUNCIONÁRIO',f'Atualizadas jornadas de funcionários: {corrigidos}')
        messagebox.showinfo('Jornadas atualizadas', f'Funcionários atualizados conforme a jornada cadastrada: {corrigidos}')

    def populate_tree(self):
        if not hasattr(self,'tree'): return
        for i in self.tree.get_children(): self.tree.delete(i)
        termo=self.search_var.get().lower() if hasattr(self,'search_var') else ''
        funcs=get_funcionarios(False if getattr(self,'show_inativos',None) and self.show_inativos.get() else True)
        for f in funcs:
            if termo and termo not in f['nome'].lower() and termo not in (f.get('funcao') or '').lower() and termo not in (f.get('setor') or '').lower() and termo not in (f.get('cpf') or '').lower(): continue
            status='Ativo' if f.get('ativo') else 'Inativo'
            self.tree.insert('', 'end', iid=str(f['id']), values=(f['nome'],f.get('setor') or 'GERAL',f.get('funcao') or '',fmt_data(f.get('admissao')),money(f.get('salario')),status))

    def on_select(self, _=None):
        sel=self.tree.selection()
        if not sel: return
        fid=int(sel[0]); self.selected_id=fid
        f=next((x for x in get_funcionarios(False) if x['id']==fid), None)
        if not f: return
        for k,v in self.fvars.items(): v.set(str(f.get(k) or ''))
        if f.get('salario') is not None: self.fvars['salario'].set(money(f['salario']))
        self.aplicar_jornada_funcionario()

    def duplicar_func(self):
        if not self.selected_id:
            messagebox.showwarning('Atenção','Selecione um funcionário para duplicar.'); return
        self.selected_id=None
        nome=self.fvars['nome'].get()
        self.fvars['nome'].set(nome + ' - CÓPIA')
        messagebox.showinfo('Duplicado','Revise o nome e clique em Salvar / Atualizar para criar o novo cadastro.')

    def delete_func(self):
        if not self.selected_id:
            messagebox.showwarning('Atenção','Selecione um funcionário.'); return
        if not messagebox.askyesno('Confirmar','Inativar este funcionário? Ele não aparecerá mais na geração de PDFs, mas o cadastro ficará salvo.'):
            return
        with con() as db: db.execute('UPDATE funcionarios SET ativo=0 WHERE id=?',(self.selected_id,))
        self.clear_func(); self.populate_tree(); self.populate_combo(); self.populate_setor_combo(); self.populate_ocorrencias_combo(); self.refresh_dashboard()
        log_action(self.usuario,'FUNCIONÁRIO','Funcionário desativado')
        messagebox.showinfo('Inativado','Funcionário inativado. Para ver ou reativar, marque Mostrar inativos.')

    def reativar_func(self):
        if not self.selected_id:
            messagebox.showwarning('Atenção','Selecione um funcionário.'); return
        with con() as db: db.execute('UPDATE funcionarios SET ativo=1 WHERE id=?',(self.selected_id,))
        self.populate_tree(); self.populate_combo(); self.populate_setor_combo(); self.populate_ocorrencias_combo(); self.refresh_dashboard()
        log_action(self.usuario,'FUNCIONÁRIO','Funcionário reativado')
        messagebox.showinfo('Reativado','Funcionário reativado.')

    def populate_combo(self):
        if not hasattr(self,'combo_func'): return
        funcs=get_funcionarios()
        vals=['TODOS']+[f"{f['id']} - {f['nome']}" for f in funcs]
        self.combo_func['values']=vals
        if self.func_pdf_var.get() not in vals: self.func_pdf_var.set('TODOS')
        self.populate_setor_combo()

    def populate_setor_combo(self):
        setores=get_setores(True)
        if hasattr(self, 'combo_setor_func'):
            self.combo_setor_func['values']=setores
            if self.fvars.get('setor') and self.fvars['setor'].get() not in setores:
                self.fvars['setor'].set('GERAL' if 'GERAL' in setores else (setores[0] if setores else 'GERAL'))
        if not hasattr(self,'combo_setor'):
            return
        vals=['TODOS'] + setores
        self.combo_setor['values']=vals
        if self.setor_pdf_var.get() not in vals: self.setor_pdf_var.set('TODOS')


    def populate_jornada_combo(self):
        jornadas_ativas = get_jornadas(True)
        vals=[j['nome'] for j in jornadas_ativas]
        horarios=sorted(set([texto_horario_jornada(normalizar_jornada_12x36(j)) for j in jornadas_ativas if texto_horario_jornada(normalizar_jornada_12x36(j))]))
        if hasattr(self, 'combo_jornada_func'):
            self.combo_jornada_func['values']=vals
            if vals and self.fvars.get('jornada') and self.fvars['jornada'].get() not in vals:
                self.fvars['jornada'].set(vals[0])
        if hasattr(self, 'combo_horario_func'):
            self.combo_horario_func['values']=horarios

    def _selecionar_funcionarios_para_pdf(self):
        """Usado pela pré-visualização: respeita funcionário selecionado e setor selecionado."""
        funcs = get_funcionarios()
        setor = self.setor_pdf_var.get().strip().upper() if hasattr(self, 'setor_pdf_var') else 'TODOS'
        val = self.func_pdf_var.get() if hasattr(self, 'func_pdf_var') else 'TODOS'
        if val and val != 'TODOS':
            try:
                fid = int(val.split(' - ')[0])
                funcs = [f for f in funcs if f['id'] == fid]
            except Exception:
                pass
        elif setor and setor != 'TODOS':
            funcs = [f for f in funcs if (f.get('setor') or 'GERAL').strip().upper() == setor]
        return funcs

    def visualizar_pdf(self):
        """Gera um PDF temporário e mostra dentro do programa, sem salvar no histórico."""
        mes = int(self.mes_var.get())
        ano = int(self.ano_var.get())
        funcs = self._selecionar_funcionarios_para_pdf()
        if not funcs:
            messagebox.showwarning('Atenção', 'Não há funcionários para visualizar com os filtros selecionados.')
            return
        preview_dir = os.path.join(BASE_DIR, 'preview')
        os.makedirs(preview_dir, exist_ok=True)
        destino = os.path.join(preview_dir, 'PREVIEW_Folha_de_Ponto.pdf')
        gerar_pdf_funcionarios(funcs, mes, ano, destino)
        self.pdf_info.config(text=f'Pré-visualização gerada temporariamente: {destino}')
        self.abrir_visualizador_pdf(destino)

    def abrir_visualizador_pdf(self, pdf_path):
        """Visualizador interno simples. Usa PyMuPDF + Pillow; se faltar biblioteca, abre no leitor padrão."""
        try:
            import fitz  # PyMuPDF
            from PIL import Image, ImageTk
        except Exception:
            messagebox.showwarning(
                'Biblioteca ausente',
                'Para visualizar dentro do programa, execute INSTALAR_BIBLIOTECAS.bat.\n\n'
                'O PDF será aberto no visualizador padrão do Windows.'
            )
            try:
                if os.name == 'nt': os.startfile(pdf_path)
                else: webbrowser.open(pdf_path)
            except Exception:
                pass
            return

        win = tk.Toplevel(self)
        win.title('Visualização da Folha de Ponto')
        win.geometry('950x720')
        win.configure(bg='#f3f4f6')
        doc = fitz.open(pdf_path)
        state = {'page': 0, 'img': None, 'zoom': 1.35}

        toolbar = ttk.Frame(win)
        toolbar.pack(fill='x', padx=8, pady=6)
        lbl = ttk.Label(toolbar, text='Página 1 de %d' % len(doc), font=('Arial', 10, 'bold'))
        lbl.pack(side='left', padx=8)
        frame_canvas = ttk.Frame(win)
        frame_canvas.pack(fill='both', expand=True, padx=8, pady=6)
        canvas_pdf = tk.Canvas(frame_canvas, bg='white')
        vsb = ttk.Scrollbar(frame_canvas, orient='vertical', command=canvas_pdf.yview)
        hsb = ttk.Scrollbar(frame_canvas, orient='horizontal', command=canvas_pdf.xview)
        canvas_pdf.configure(yscrollcommand=vsb.set, xscrollcommand=hsb.set)
        canvas_pdf.grid(row=0, column=0, sticky='nsew')
        vsb.grid(row=0, column=1, sticky='ns')
        hsb.grid(row=1, column=0, sticky='ew')
        frame_canvas.rowconfigure(0, weight=1)
        frame_canvas.columnconfigure(0, weight=1)

        def render():
            page = doc.load_page(state['page'])
            mat = fitz.Matrix(state['zoom'], state['zoom'])
            pix = page.get_pixmap(matrix=mat, alpha=False)
            img = Image.frombytes('RGB', [pix.width, pix.height], pix.samples)
            state['img'] = ImageTk.PhotoImage(img)
            canvas_pdf.delete('all')
            canvas_pdf.create_image(0, 0, image=state['img'], anchor='nw')
            canvas_pdf.config(scrollregion=(0, 0, pix.width, pix.height))
            lbl.config(text='Página %d de %d' % (state['page'] + 1, len(doc)))

        def prev_page():
            if state['page'] > 0:
                state['page'] -= 1
                render()

        def next_page():
            if state['page'] < len(doc) - 1:
                state['page'] += 1
                render()

        def zoom_in():
            state['zoom'] = min(2.4, state['zoom'] + 0.15)
            render()

        def zoom_out():
            state['zoom'] = max(0.7, state['zoom'] - 0.15)
            render()

        ttk.Button(toolbar, text='Anterior', command=prev_page).pack(side='left', padx=4)
        ttk.Button(toolbar, text='Próxima', command=next_page).pack(side='left', padx=4)
        ttk.Button(toolbar, text='Zoom -', command=zoom_out).pack(side='left', padx=(16,4))
        ttk.Button(toolbar, text='Zoom +', command=zoom_in).pack(side='left', padx=4)
        ttk.Button(toolbar, text='Abrir no leitor padrão', command=lambda: os.startfile(pdf_path) if os.name=='nt' else webbrowser.open(pdf_path)).pack(side='right', padx=4)
        ttk.Button(toolbar, text='Fechar', command=win.destroy).pack(side='right', padx=4)
        render()

    def gerar_pdf_setor(self):
        setor = self.setor_pdf_var.get().strip().upper() if hasattr(self, 'setor_pdf_var') else 'TODOS'
        if not setor or setor == 'TODOS':
            messagebox.showwarning('Atenção','Selecione um setor para gerar o PDF por setor.')
            return
        mes=int(self.mes_var.get()); ano=int(self.ano_var.get())
        funcs=[f for f in get_funcionarios() if (f.get('setor') or 'GERAL').strip().upper()==setor]
        if not funcs:
            messagebox.showwarning('Atenção','Não há funcionários ativos neste setor.')
            return
        pasta=os.path.join(PDF_DIR, str(ano), f'{mes:02d}_{MESES[mes-1]}', 'Setores')
        os.makedirs(pasta, exist_ok=True)
        destino=os.path.join(pasta, f'Folhas_Ponto_Setor_{safe_name(setor)}_{MESES[mes-1]}_{ano}.pdf')
        gerar_pdf_funcionarios(funcs, mes, ano, destino)
        self.last_pdf=destino
        try: self.card_pdf.config(text=os.path.basename(destino)[:22]+'...')
        except Exception: pass
        self.pdf_info.config(text=f'PDF do setor salvo em: {destino}')
        log_action(self.usuario,'PDF',f'Gerado setor {setor}: {destino}')
        registrar_pdf(self.usuario, 'PDF por setor', setor, mes, ano, destino, len(funcs))
        self.carregar_historico_pdf(); self.carregar_ferias(); self.carregar_banco_horas()
        messagebox.showinfo('PDF por setor gerado', f'Setor: {setor}\nFuncionários: {len(funcs)}\n\nSalvo em:\n{destino}')
        self.abrir_ultimo_pdf()

    def gerar_pdf(self, todos=True):
        mes=int(self.mes_var.get()); ano=int(self.ano_var.get())
        funcs=get_funcionarios()
        if not todos:
            val=self.func_pdf_var.get()
            if val=='TODOS':
                messagebox.showwarning('Atenção','Selecione um funcionário ou use Gerar PDF de todos.'); return
            fid=int(val.split(' - ')[0]); funcs=[f for f in funcs if f['id']==fid]
        pasta=os.path.join(PDF_DIR, str(ano), f'{mes:02d}_{MESES[mes-1]}')
        os.makedirs(pasta, exist_ok=True)
        if todos:
            destino=os.path.join(pasta, f'Folhas_Ponto_Todos_{MESES[mes-1]}_{ano}.pdf')
        else:
            destino=os.path.join(pasta, f'Folha_Ponto_{safe_name(funcs[0]["nome"])}_{MESES[mes-1]}_{ano}.pdf')
        gerar_pdf_funcionarios(funcs, mes, ano, destino)
        self.last_pdf=destino
        try: self.card_pdf.config(text=os.path.basename(destino)[:22]+'...')
        except Exception: pass
        self.pdf_info.config(text=f'Último PDF salvo em: {destino}')
        log_action(self.usuario,'PDF','Gerado: '+destino)
        registrar_pdf(self.usuario, 'PDF todos' if todos else 'PDF individual', 'TODOS' if todos else (funcs[0].get('setor') or 'GERAL'), mes, ano, destino, len(funcs))
        self.carregar_historico_pdf(); self.carregar_ferias(); self.carregar_banco_horas()
        messagebox.showinfo('PDF gerado', f'PDF salvo automaticamente em:\n{destino}')
        self.abrir_ultimo_pdf()

    def gerar_pdfs_todos_setores(self):
        mes=int(self.mes_var.get()); ano=int(self.ano_var.get())
        setores=get_setores(True)
        pasta_base=os.path.join(PDF_DIR, str(ano), f'{mes:02d}_{MESES[mes-1]}', 'POR_SETOR')
        os.makedirs(pasta_base, exist_ok=True)
        gerados=0; ultimo=None
        for setor in setores:
            funcs=[f for f in get_funcionarios() if (f.get('setor') or 'GERAL').strip().upper()==setor]
            if not funcs:
                continue
            destino=os.path.join(pasta_base, f'Folhas_Ponto_Setor_{safe_name(setor)}_{MESES[mes-1]}_{ano}.pdf')
            gerar_pdf_funcionarios(funcs, mes, ano, destino)
            ultimo=destino
            gerados+=1
        self.last_pdf=ultimo
        try:
            if ultimo: self.card_pdf.config(text=os.path.basename(ultimo)[:22]+'...')
        except Exception:
            pass
        self.pdf_info.config(text=f'PDFs por setor salvos em: {pasta_base}')
        log_action(self.usuario,'PDF',f'Gerados PDFs separados por setor: {gerados}')
        registrar_pdf(self.usuario, 'PDFs separados por setor', 'TODOS', mes, ano, pasta_base, gerados)
        self.carregar_historico_pdf(); self.carregar_ferias(); self.carregar_banco_horas()
        messagebox.showinfo('PDFs por setor', f'Foram gerados {gerados} arquivo(s) por setor em:\n{pasta_base}')
        self.open_pdfs()

    def abrir_ultimo_pdf(self):
        if self.last_pdf and os.path.exists(self.last_pdf):
            try:
                if os.name=='nt': os.startfile(self.last_pdf)
                else: os.system(f'xdg-open "{self.last_pdf}" >/dev/null 2>&1 &')
            except Exception: pass
        else:
            messagebox.showinfo('PDF','Nenhum PDF foi gerado nesta sessão.')


    def gerar_pdfs_individuais_todos(self):
        mes=int(self.mes_var.get()); ano=int(self.ano_var.get())
        funcs=get_funcionarios()
        if not funcs:
            messagebox.showwarning('Atenção','Não há funcionários ativos.'); return
        pasta_base=os.path.join(PDF_DIR, str(ano), f'{mes:02d}_{MESES[mes-1]}', 'INDIVIDUAIS')
        gerados=0; ultimo=None
        for f in funcs:
            setor=safe_name((f.get('setor') or 'GERAL').strip().upper())
            pasta=os.path.join(pasta_base, setor)
            os.makedirs(pasta, exist_ok=True)
            destino=os.path.join(pasta, f'Folha_Ponto_{safe_name(f["nome"])}_{MESES[mes-1]}_{ano}.pdf')
            gerar_pdf_funcionarios([f], mes, ano, destino)
            ultimo=destino; gerados += 1
        self.last_pdf=ultimo
        self.pdf_info.config(text=f'PDFs individuais salvos em: {pasta_base}')
        registrar_pdf(self.usuario, 'PDFs individuais', 'TODOS', mes, ano, pasta_base, gerados)
        log_action(self.usuario,'PDF',f'Gerados PDFs individuais: {gerados}')
        self.carregar_historico_pdf(); self.carregar_ferias(); self.carregar_banco_horas()
        messagebox.showinfo('PDFs individuais', f'Foram gerados {gerados} PDF(s) individuais em:\n{pasta_base}')
        self.open_pdfs()

    def carregar_historico_pdf(self):
        if not hasattr(self, 'hist_pdf'): return
        for i in self.hist_pdf.get_children(): self.hist_pdf.delete(i)
        try:
            with con() as db:
                rows=db.execute('SELECT data_hora,tipo,setor,mes,ano,quantidade,arquivo FROM pdf_historico ORDER BY id DESC LIMIT 200').fetchall()
            for r in rows: self.hist_pdf.insert('', 'end', values=r)
        except Exception:
            pass


    def criar_modelo_importacao_excel(self, path):
        """Cria a planilha padrão de importação do sistema."""
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
            from openpyxl.worksheet.datavalidation import DataValidation
            from openpyxl.comments import Comment
        except Exception:
            raise RuntimeError('Instale a biblioteca openpyxl. Use o arquivo INSTALAR_BIBLIOTECAS.bat.')

        wb = Workbook()
        ws = wb.active
        ws.title = 'Funcionarios'
        headers = [
            'Matrícula','Nome Completo','CPF','RG/CTPS','Data de Nascimento','Data de Admissão',
            'Cargo/Função','Setor','CBO','Jornada','Horário de Trabalho','Salário','Empresa','Telefone','E-mail','Status'
        ]
        ws.append(headers)
        exemplo = [
            '0001','JOÃO DA SILVA','000.000.000-00','0.000.000','01/01/1990','01/02/2024',
            'MOTORISTA','GERAL','7823-05','COMERCIAL 44H','07:42 - 12:00 / 13:30 - 18:00',1800.00,
            get_empresa().get('nome',''), '(47) 99999-9999','joao@empresa.com','Ativo'
        ]
        ws.append(exemplo)
        ws.freeze_panes = 'A2'
        header_fill = PatternFill('solid', fgColor='1F2937')
        thin = Side(style='thin', color='D1D5DB')
        border = Border(left=thin, right=thin, top=thin, bottom=thin)
        for cell in ws[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            cell.border = border
        for row in ws.iter_rows(min_row=2, max_row=2):
            for cell in row:
                cell.border = border
                cell.alignment = Alignment(vertical='center')
        widths = [12,34,18,18,18,18,24,20,12,24,34,14,28,18,28,12]
        for i,w in enumerate(widths,1):
            ws.column_dimensions[chr(64+i)].width = w
        for row in range(2, 502):
            ws.cell(row=row, column=5).number_format = 'DD/MM/YYYY'
            ws.cell(row=row, column=6).number_format = 'DD/MM/YYYY'
            ws.cell(row=row, column=12).number_format = '#,##0.00'
        dv_status = DataValidation(type='list', formula1='"Ativo,Inativo"', allow_blank=False)
        ws.add_data_validation(dv_status)
        dv_status.add('P2:P501')
        setores = get_setores(False)
        if setores:
            ws2 = wb.create_sheet('Setores')
            ws2.append(['Setores cadastrados'])
            for s_nome in setores:
                ws2.append([s_nome])
            ws2.column_dimensions['A'].width = 28
            for cell in ws2[1]:
                cell.font = Font(bold=True, color='FFFFFF')
                cell.fill = header_fill
            formula = f'=Setores!$A$2:$A${max(2, len(setores)+1)}'
            dv_setor = DataValidation(type='list', formula1=formula, allow_blank=True)
            ws.add_data_validation(dv_setor)
            dv_setor.add('H2:H501')
        ws['A1'].comment = Comment('Matrícula é opcional. Se informada, o sistema usa esse campo para atualizar o cadastro existente.', 'Ponto Saúde')
        ws['B1'].comment = Comment('Campo obrigatório.', 'Ponto Saúde')
        ws['F1'].comment = Comment('Digite no padrão DD/MM/AAAA.', 'Ponto Saúde')
        ws['H1'].comment = Comment('Se o setor não existir, o sistema cria automaticamente.', 'Ponto Saúde')
        ws['P1'].comment = Comment('Use Ativo ou Inativo.', 'Ponto Saúde')
        wb.save(path)

    def baixar_modelo_importacao(self):
        path = filedialog.asksaveasfilename(
            title='Salvar modelo de importação de funcionários',
            defaultextension='.xlsx',
            initialfile='modelo_importacao_funcionarios_ponto_saude.xlsx',
            filetypes=[('Excel','*.xlsx')]
        )
        if not path:
            return
        try:
            self.criar_modelo_importacao_excel(path)
            log_action(self.usuario, 'MODELO', 'Modelo de importação gerado')
            messagebox.showinfo('Modelo gerado', 'Modelo de importação salvo em:\n' + path)
        except Exception as e:
            messagebox.showerror('Erro ao gerar modelo', str(e))

    def normalizar_header(self, texto):
        import unicodedata
        texto = str(texto or '').strip().lower()
        texto = ''.join(ch for ch in unicodedata.normalize('NFD', texto) if unicodedata.category(ch) != 'Mn')
        for ch in ['/', '.', '-', '_', '(', ')']:
            texto = texto.replace(ch, ' ')
        return ' '.join(texto.split())

    def valor_data_planilha_para_iso(self, valor):
        if valor in (None, ''):
            return ''
        if isinstance(valor, datetime):
            return valor.date().isoformat()
        if isinstance(valor, date):
            return valor.isoformat()
        try:
            return data_para_iso(str(valor).strip())
        except Exception:
            return ''

    def valor_moeda_para_float(self, valor):
        if valor in (None, ''):
            return 0.0
        if isinstance(valor, (int, float)):
            return float(valor)
        s = str(valor).strip().replace('R$', '').replace(' ', '')
        if ',' in s:
            s = s.replace('.', '').replace(',', '.')
        try:
            return float(s)
        except Exception:
            return 0.0

    def exportar_excel(self):
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill, Alignment
        except Exception:
            messagebox.showerror('Biblioteca ausente','Instale a biblioteca openpyxl. Use o arquivo INSTALAR_BIBLIOTECAS.bat.')
            return
        path=filedialog.asksaveasfilename(title='Exportar funcionários para Excel', defaultextension='.xlsx', filetypes=[('Excel','*.xlsx')])
        if not path: return
        funcs=get_funcionarios(False)
        wb=Workbook(); ws=wb.active; ws.title='Funcionarios'
        headers=['Nome','CPF','CTPS','Admissão','Função','Setor','CBO','Salário','Jornada','Horário de Trabalho','Descanso','Sábado','Ativo']
        ws.append(headers)
        for cell in ws[1]:
            cell.font=Font(bold=True, color='FFFFFF'); cell.fill=PatternFill('solid', fgColor='1F2937'); cell.alignment=Alignment(horizontal='center')
        for f in funcs:
            ws.append([f.get('nome'),f.get('cpf'),f.get('ctps'),f.get('admissao'),f.get('funcao'),f.get('setor') or 'GERAL',f.get('cbo'),f.get('salario'),f.get('jornada'),f.get('horario_trabalho'),f.get('descanso'),f.get('sabado'),'Sim' if f.get('ativo') else 'Não'])
        widths=[34,16,18,14,18,18,10,12,30,34,18,16,10]
        for idx,w in enumerate(widths,1): ws.column_dimensions[chr(64+idx)].width=w
        wb.save(path)
        log_action(self.usuario,'EXPORTAÇÃO','Funcionários Excel')
        messagebox.showinfo('Exportado', f'Arquivo Excel salvo em:\n{path}')

    def restaurar_backup(self):
        sel=self.backup_list.curselection() if hasattr(self,'backup_list') else []
        if not sel:
            messagebox.showwarning('Atenção','Selecione um backup na lista.'); return
        name=self.backup_list.get(sel[0]); src=os.path.join(BACKUP_DIR, name)
        if not os.path.exists(src): return
        if not messagebox.askyesno('Confirmar restauração','Isso substituirá o banco atual pelo backup selecionado. Deseja continuar?'):
            return
        backup_db()
        shutil.copy2(src, DB_PATH)
        log_action(self.usuario,'BACKUP','Backup restaurado: '+name)
        self.load_all()
        messagebox.showinfo('Restaurado','Backup restaurado com sucesso.')

    def on_close(self):
        try:
            backup_db()
            log_action(self.usuario,'BACKUP','Backup automático ao fechar')
        except Exception:
            pass
        self.destroy()

    def importar_excel(self):
        path=filedialog.askopenfilename(title='Importar funcionários do Excel', filetypes=[('Excel','*.xlsx *.xlsm'),('Todos','*.*')])
        if not path: return
        try:
            from openpyxl import load_workbook
        except Exception:
            messagebox.showerror('Biblioteca ausente','Instale a biblioteca openpyxl. Use o arquivo INSTALAR_BIBLIOTECAS.bat.')
            return
        try:
            wb=load_workbook(path, data_only=True, read_only=True)
            ws=wb['Funcionarios'] if 'Funcionarios' in wb.sheetnames else wb.active
            rows=list(ws.iter_rows(values_only=True))
            if not rows:
                messagebox.showwarning('Importação','A planilha está vazia.'); return
            header=[self.normalizar_header(c) for c in rows[0]]
            aliases={
                'matricula':['matricula','matrícula','codigo','código','registro'],
                'nome':['nome completo','nome','funcionario','funcionário','empregado'],
                'cpf':['cpf'],
                'ctps':['rg ctps','ctps','rg','rg/ctps','ci','c i'],
                'nascimento':['data de nascimento','nascimento'],
                'admissao':['data de admissao','data de admissão','admissao','admissão'],
                'funcao':['cargo funcao','cargo função','cargo','funcao','função'],
                'setor':['setor','departamento'],
                'cbo':['cbo'],
                'jornada':['jornada','horario','horário'],
                'salario':['salario','salário','salario base','salário base'],
                'empresa':['empresa'],
                'telefone':['telefone','fone','celular'],
                'email':['e mail','email','e-mail'],
                'status':['status','ativo']
            }
            idx={}
            for key,names in aliases.items():
                for n in names:
                    n_norm=self.normalizar_header(n)
                    if n_norm in header:
                        idx[key]=header.index(n_norm); break
            if 'nome' not in idx:
                messagebox.showerror('Importação','Não encontrei a coluna Nome Completo/Nome na primeira linha da planilha.')
                return
            inseridos=atualizados=erros=0
            mensagens=[]
            with con() as db:
                for linha_num,row in enumerate(rows[1:], start=2):
                    def val(k):
                        pos=idx.get(k)
                        return row[pos] if pos is not None and pos < len(row) else ''
                    nome=str(val('nome') or '').strip().upper()
                    if not nome:
                        continue
                    if nome in ('NOME COMPLETO','NOME','FUNCIONÁRIO','FUNCIONARIO'):
                        continue
                    cpf=str(val('cpf') or '').strip()
                    ctps=str(val('ctps') or val('matricula') or '').strip()
                    admissao=self.valor_data_planilha_para_iso(val('admissao'))
                    funcao=str(val('funcao') or '').strip().upper()
                    setor=str(val('setor') or 'GERAL').strip().upper() or 'GERAL'
                    cbo=str(val('cbo') or '').strip()
                    salario=self.valor_moeda_para_float(val('salario'))
                    jornada=str(val('jornada') or 'HORÁRIO DE TRABALHO DE SEGUNDA A SEXTA-FEIRA').strip()
                    horario_trabalho=str(val('horario trabalho') or val('horario de trabalho') or '').strip()
                    if not horario_trabalho:
                        j_tmp=normalizar_jornada_12x36(get_jornada_por_nome(jornada))
                        horario_trabalho=texto_horario_jornada(j_tmp) if j_tmp else ''
                    status=str(val('status') or 'Ativo').strip().upper()
                    ativo=0 if status in ('INATIVO','NÃO','NAO','0','FALSE','FALSO') else 1
                    if len(nome) < 3:
                        erros += 1; mensagens.append(f'Linha {linha_num}: nome inválido.'); continue
                    db.execute("INSERT OR IGNORE INTO setores(nome,descricao,ativo) VALUES(?,?,1)", (setor, 'Criado automaticamente na importação'))
                    existe=None
                    if ctps:
                        existe=db.execute('SELECT id FROM funcionarios WHERE ctps=?', (ctps,)).fetchone()
                    if not existe and cpf:
                        existe=db.execute('SELECT id FROM funcionarios WHERE cpf=?', (cpf,)).fetchone()
                    if not existe:
                        existe=db.execute('SELECT id FROM funcionarios WHERE nome=?', (nome,)).fetchone()
                    if existe:
                        db.execute("""UPDATE funcionarios SET nome=?, cpf=?, ctps=?, admissao=?, funcao=?, setor=?, cbo=?, salario=?, jornada=?, horario_trabalho=?, descanso=?, sabado=?, ativo=? WHERE id=?""",
                                   (nome, cpf, ctps, admissao, funcao, setor, cbo, salario, jornada, horario_trabalho, self.valores_jornada_para_funcionario(jornada)[0], self.valores_jornada_para_funcionario(jornada)[1], ativo, existe[0]))
                        atualizados += 1
                    else:
                        db.execute("""INSERT INTO funcionarios(nome,cpf,ctps,admissao,funcao,setor,cbo,salario,jornada,horario_trabalho,descanso,sabado,ativo) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                                   (nome, cpf, ctps, admissao, funcao, setor, cbo, salario, jornada, horario_trabalho, self.valores_jornada_para_funcionario(jornada)[0], self.valores_jornada_para_funcionario(jornada)[1], ativo))
                        inseridos += 1
            self.populate_setores_tree(); self.populate_jornadas_tree(); self.populate_escalas_tree(); self.populate_feriados_tree(); self.populate_tree(); self.populate_combo(); self.populate_setor_combo(); self.populate_jornada_combo(); self.populate_ocorrencias_combo(); self.refresh_dashboard()
            log_action(self.usuario,'IMPORTAÇÃO',f'Excel: {inseridos} inseridos, {atualizados} atualizados, {erros} erros')
            msg=f'Importação concluída.\n\n✔ {inseridos} funcionários inseridos\n✔ {atualizados} funcionários atualizados\n⚠ {erros} registros com erro'
            if mensagens:
                msg += '\n\nPrimeiros erros:\n' + '\n'.join(mensagens[:8])
            messagebox.showinfo('Importação concluída', msg)
        except Exception as e:
            messagebox.showerror('Erro ao importar', str(e))

    def exportar_csv(self):
        path=filedialog.asksaveasfilename(title='Exportar funcionários', defaultextension='.csv', filetypes=[('CSV','*.csv')])
        if not path: return
        funcs=get_funcionarios(False)
        with open(path,'w',newline='',encoding='utf-8-sig') as fp:
            wr=csv.writer(fp, delimiter=';')
            wr.writerow(['Nome','CPF','CTPS','Admissão','Função','Setor','CBO','Salário','Jornada','Horário de Trabalho','Descanso','Sábado','Ativo'])
            for f in funcs:
                wr.writerow([f.get('nome'),f.get('cpf'),f.get('ctps'),fmt_data(f.get('admissao')),f.get('funcao'),f.get('setor') or 'GERAL',f.get('cbo'),money(f.get('salario')),f.get('jornada'),f.get('horario_trabalho'),f.get('descanso'),f.get('sabado'),'Sim' if f.get('ativo') else 'Não'])
        log_action(self.usuario,'EXPORTAÇÃO','Funcionários CSV')
        messagebox.showinfo('Exportado', f'Arquivo salvo em:\n{path}')

    def backup_now(self):
        backup_db(); log_action(self.usuario,'BACKUP','Backup manual criado'); self.refresh_dashboard(); messagebox.showinfo('Backup','Backup criado na pasta backups.')
    def open_base(self): self._open(BASE_DIR)
    def open_pdfs(self): self._open(PDF_DIR)
    def open_data(self): self._open(DATA_DIR)
    def open_backups(self): self._open(BACKUP_DIR)
    def open_relatorios(self): self._open(RELATORIO_DIR)
    def _open(self, path):
        os.makedirs(path, exist_ok=True)
        try:
            if os.name=='nt': os.startfile(path)
            else: os.system(f'xdg-open "{path}" >/dev/null 2>&1 &')
        except Exception: pass

if __name__ == '__main__':
    init_db()
    login=LoginDialog()
    login.mainloop()
    if login.usuario:
        app=App(login.usuario, login.perfil)
        app.mainloop()
